"""Payroll reports: the catalogue of 3.8 §5.13, previews, exports, and the overview KPIs of 3.8 §5.1.

Every report is a generator over the caller's data scope (EC-SEC-05): run-level data through
``access.run_scope_clause`` on ``payroll:reports:read`` (legal entity / pay group / company), ledgers
through their entity and pay group, audit reports only with ``payroll:audit:read``. Filters (period
range, legal entity, pay group, department, cost center, branch, currency, status) narrow the same scoped
query, so counts, totals and rows never include out-of-scope records.

* Money is never summed across currencies (EC-MON-04): totals are per currency; a converted total is
  shown only on request (``convert_to``) and always with its FX source, rate and date, from the tenant's
  usable rates (never 1:1; a missing rate leaves the conversion unavailable with the reason).
* Run-based reports read FINAL runs (finalized or later) unless the ``status`` filter asks for DRAFT,
  APPROVED, ALL or one run status; each row carries its run stage (Draft / Approved / Final).
* Salary analytics (employer cost and CTC vs actual by group) hide groups smaller than
  ``privacy.small_group_suppression_size``; when a single group of a currency is hidden, the next
  smallest is hidden too so it cannot be derived from the total (complementary suppression).
* Preview: the first 50 rows with the full row count and totals. Export: CSV, XLSX or PDF through a
  background job (``payroll.report_export``) that rebuilds the report with the requester's live grants,
  watermarks it "Confidential — exported by {user} at {time}", escapes spreadsheet formulas (EC-DOC-17)
  and masks sensitive columns unless the exporter has full access (``payroll:reports:export`` at COMPANY
  scope). Only the requester downloads the file, for 24 hours; requests and downloads are audited.
* Payment reports read the payments milestone's tables (``app.models.payment_flows``); bank account
  numbers never appear in them, only the masked beneficiary snapshot.

``overview_kpis`` / ``overview_list`` serve GET /api/v1/payroll/overview from ONE population query, so a
KPI count always equals the length of its drill-down list (EC-UX-03).
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any, Callable, Dict, Iterable, List, Mapping, Optional, Sequence, Set, Tuple

from sqlalchemy import and_, false, func, or_
from sqlalchemy.orm import Session

from app.core import clock, storage
from app.core.errors import AppError, bad_request, conflict, forbidden, not_found
from app.core.principal import Principal
from app.models.audit_log import AuditLog
from app.models.compensation import CompensationAssignment, SalaryTemplate, SalaryTemplateVersion
from app.models.employee import Department, Employee
from app.models.jobs import Job
from app.models.org import Branch, EmployeeAssignment, LegalEntity
from app.models.pay_org import CostCenter, Grade, PayGroup
from app.models.pay_runs import (
    YtdBalance,
    PayPeriod,
    PayrollException,
    PayrollLine,
    PayrollRun,
    PayrollRunEmployee,
    StatutoryLiability,
)
from app.models.reports_compliance import StatutoryReport
from app.models.variable_inputs import Loan, ReimbursementClaim
from app.models.work_items import Task
from app.payroll import audit as payroll_audit
from app.payroll import currencies
from app.payroll import fx as fxlib
from app.payroll import state_machines as sm
from app.payroll.banking import masking
from app.payroll.banking.files import escape_csv_cell
from app.payroll.money import Money, amount_str
from app.payroll.packs.base import FIXTURE_LABEL
from app.payroll.services import access, jobs
from app.payroll.services import fx as fx_service
from app.payroll.services import settings as settings_engine
from app.services.audit_log_service import AuditLogService

READ = "payroll:reports:read"
EXPORT = "payroll:reports:export"
AUDIT_READ = payroll_audit.PAYROLL_AUDIT_PERMISSION
COMPENSATION_READ = "payroll:compensation:read"
RUN_READ = "payroll:run:read"
OVERVIEW_PERMISSIONS = (RUN_READ, READ)
AUDIT_MODULE = payroll_audit.PAYROLL_RUNS
PREVIEW_ROWS = 50
MAX_ROWS = 100_000
PDF_MAX_ROWS = 5_000
EXPORT_JOB = "payroll.report_export"
EXPORT_CATEGORY = "report-exports"
EXPORT_TTL = timedelta(hours=24)
FORMATS = ("CSV", "XLSX", "PDF")
SUPPRESSION_KEY = "privacy.small_group_suppression_size"
CONFIDENTIAL = "Confidential — exported by {user} at {time}"
_ZERO = Decimal("0")

FINAL_STATUSES = ("FINALIZED", "PAYMENT_READY", "SUBMITTED", "PARTIALLY_PAID", "PAID", "RECONCILED", "CLOSED")
APPROVED_STATUSES = ("APPROVED",)
DRAFT_STATUSES = ("DRAFT", "INPUTS_PENDING", "READY", "CALCULATING", "CALCULATED", "VALIDATION_FAILED",
                  "REVIEW_REQUIRED", "APPROVAL_PENDING", "REJECTED", "REOPENED_FOR_CORRECTION")
STAGES: Dict[str, Tuple[str, ...]] = {
    "FINAL": FINAL_STATUSES, "APPROVED": APPROVED_STATUSES, "DRAFT": DRAFT_STATUSES,
    "ALL": FINAL_STATUSES + APPROVED_STATUSES + DRAFT_STATUSES,
}
GROUP_BY = ("legal_entity", "department", "cost_center", "grade")
PERIODS_PER_YEAR = {"MONTHLY": 12, "SEMI_MONTHLY": 24, "BI_WEEKLY": 26, "WEEKLY": 52}


def run_stage(status: str) -> str:
    """Draft / Approved / Final label of a run state (3.8 §5.1)."""
    if status in FINAL_STATUSES:
        return "Final"
    if status in APPROVED_STATUSES:
        return "Approved"
    return "Draft"


# =========================================================================== filters and context


@dataclass(frozen=True)
class Filters:
    period_from: Optional[date] = None
    period_to: Optional[date] = None
    legal_entity_id: Optional[int] = None
    pay_group_id: Optional[int] = None
    department_id: Optional[int] = None
    cost_center_id: Optional[int] = None
    branch_id: Optional[int] = None
    currency: Optional[str] = None
    status: Optional[str] = None
    group_by: Optional[str] = None
    convert_to: Optional[str] = None
    fx_date: Optional[date] = None
    tax_year: Optional[str] = None

    def to_payload(self) -> dict:
        return {k: (v.isoformat() if isinstance(v, date) else v) for k, v in asdict(self).items() if v is not None}

    @classmethod
    def from_payload(cls, data: Mapping[str, Any]) -> "Filters":
        values: Dict[str, Any] = {}
        for key, value in (data or {}).items():
            if key not in cls.__dataclass_fields__ or value is None:
                continue
            values[key] = date.fromisoformat(value) if key in ("period_from", "period_to", "fx_date") else value
        return cls(**values)

    def validate(self) -> "Filters":
        if self.period_from and self.period_to and self.period_from > self.period_to:
            raise bad_request("The period range ends before it starts.", "INVALID_PERIOD_RANGE")
        for code in (self.currency, self.convert_to):
            if code is not None and currencies.find(code) is None:
                raise bad_request(f"Unknown currency {code}.", "UNKNOWN_CURRENCY")
        if self.group_by is not None and self.group_by not in GROUP_BY:
            raise bad_request(f"group_by must be one of {', '.join(GROUP_BY)}.", "INVALID_GROUP_BY")
        return self


@dataclass
class Ctx:
    db: Session
    principal: Principal
    filters: Filters
    today: date
    report: "ReportDef"
    names: Dict[str, Dict[int, str]] = field(default_factory=dict)
    notes: List[str] = field(default_factory=list)
    default_group: str = "legal_entity"  # salary analytics without a group_by filter

    @property
    def cid(self) -> int:
        return self.principal.company_id

    def name(self, kind: str, object_id: Optional[int]) -> Optional[str]:
        if object_id is None:
            return None
        if kind not in self.names:
            self.names[kind] = _load_names(self.db, self.cid, kind)
        return self.names[kind].get(object_id)


def _load_names(db: Session, company_id: int, kind: str) -> Dict[int, str]:
    model, label = {
        "legal_entity": (LegalEntity, LegalEntity.code), "pay_group": (PayGroup, PayGroup.code),
        "department": (Department, Department.name), "cost_center": (CostCenter, CostCenter.code),
        "grade": (Grade, Grade.code), "branch": (Branch, Branch.code),
    }[kind]
    return {i: n for i, n in db.query(model.id, label).filter(model.company_id == company_id)}


@dataclass(frozen=True)
class Col:
    key: str
    label: str
    type: str = "text"  # text | int | decimal | money | date | datetime | bool
    sensitive: bool = False
    total: bool = True


@dataclass(frozen=True)
class ReportDef:
    code: str
    group: str
    name: str
    description: str
    columns: Tuple[Col, ...]
    generator: Callable[[Ctx], List[dict]]
    extra_permissions: Tuple[str, ...] = ()
    analytics: bool = False
    run_based: bool = False
    filters: Tuple[str, ...] = ("period", "legal_entity", "pay_group", "department", "cost_center", "currency",
                                "status")
    requires: Optional[str] = None  # an optional model module (payments) the report reads

    @property
    def money_columns(self) -> List[Col]:
        return [c for c in self.columns if c.type == "money" and c.total]


# =========================================================================== shared query pieces


def _held(principal: Principal, permissions: Iterable[str]) -> List[str]:
    return [p for p in permissions if principal.has(p)]


def scope_any(principal: Principal, permissions: Sequence[str], legal_entity_col, pay_group_col=None):
    """Union of the run-level scopes of the permissions held (false when none)."""
    held = _held(principal, permissions)
    if not held:
        return false()
    return or_(*(access.run_scope_clause(principal, p, legal_entity_col, pay_group_col) for p in held))


def _run_statuses(status: Optional[str]) -> Tuple[str, ...]:
    if status is None:
        return FINAL_STATUSES
    if status in STAGES:
        return STAGES[status]
    if status in STAGES["ALL"]:
        return (status,)
    raise bad_request("For run reports, status is FINAL, APPROVED, DRAFT, ALL or a run status.", "INVALID_STATUS")


def _employee_name(first: Optional[str], last: Optional[str]) -> str:
    return " ".join(p for p in (first, last) if p)


def run_employee_query(ctx: Ctx, *extra, statuses: Optional[Sequence[str]] = None):
    """Run employees of in-scope, filtered runs, with their run, period, employee and assignment."""
    f = ctx.filters
    query = (
        ctx.db.query(PayrollRunEmployee, PayrollRun, PayPeriod, Employee, EmployeeAssignment, *extra)
        .join(PayrollRun, and_(PayrollRun.company_id == PayrollRunEmployee.company_id,
                               PayrollRun.id == PayrollRunEmployee.run_id))
        .join(PayPeriod, and_(PayPeriod.company_id == PayrollRun.company_id, PayPeriod.id == PayrollRun.pay_period_id))
        .join(Employee, and_(Employee.company_id == PayrollRunEmployee.company_id,
                             Employee.id == PayrollRunEmployee.employee_id))
        .outerjoin(EmployeeAssignment, and_(EmployeeAssignment.company_id == PayrollRunEmployee.company_id,
                                    EmployeeAssignment.id == PayrollRunEmployee.employee_assignment_id))
        .filter(PayrollRunEmployee.company_id == ctx.cid, PayrollRun.status != "CANCELLED",
                PayrollRun.status.in_(_run_statuses(f.status)),
                access.run_scope_clause(ctx.principal, READ, PayrollRun.legal_entity_id, PayrollRun.pay_group_id))
    )
    if statuses is not None:
        query = query.filter(PayrollRunEmployee.status.in_(tuple(statuses)))
    if f.period_from is not None:
        query = query.filter(PayPeriod.period_end >= f.period_from)
    if f.period_to is not None:
        query = query.filter(PayPeriod.period_start <= f.period_to)
    if f.legal_entity_id is not None:
        query = query.filter(PayrollRun.legal_entity_id == f.legal_entity_id)
    if f.pay_group_id is not None:
        query = query.filter(PayrollRun.pay_group_id == f.pay_group_id)
    if f.department_id is not None:
        query = query.filter(EmployeeAssignment.department_id == f.department_id)
    if f.cost_center_id is not None:
        query = query.filter(EmployeeAssignment.cost_center_id == f.cost_center_id)
    if f.branch_id is not None:
        query = query.filter(EmployeeAssignment.branch_id == f.branch_id)
    if f.currency is not None:
        query = query.filter(PayrollRunEmployee.currency == f.currency)
    return query


def _base_row(ctx: Ctx, rem: PayrollRunEmployee, run: PayrollRun, period: PayPeriod, employee: Employee,
              assignment: Optional[EmployeeAssignment]) -> dict:
    return {
        "run_number": run.number, "run_type": run.run_type, "run_stage": run_stage(run.status),
        "period_start": period.period_start, "period_end": period.period_end,
        "employee_code": employee.employee_code, "employee_name": _employee_name(employee.first_name,
                                                                                employee.last_name),
        "legal_entity": ctx.name("legal_entity", run.legal_entity_id),
        "pay_group": ctx.name("pay_group", run.pay_group_id),
        "department": ctx.name("department", assignment.department_id if assignment else None),
        "cost_center": ctx.name("cost_center", assignment.cost_center_id if assignment else None),
        "currency": rem.currency,
    }


_EMPLOYEE_COLS = (
    Col("run_number", "Run"), Col("period_end", "Period end", "date"), Col("employee_code", "Employee code"),
    Col("employee_name", "Employee"), Col("legal_entity", "Legal entity"), Col("pay_group", "Pay group"),
)


def _d(value: Any) -> Decimal:
    return _ZERO if value is None else Decimal(value)


# =========================================================================== payroll group


def gen_payroll_register(ctx: Ctx) -> List[dict]:
    rows = []
    query = run_employee_query(ctx, statuses=("CALCULATED", "ON_HOLD"))
    for rem, run, period, employee, assignment in query.order_by(PayPeriod.period_end, PayrollRun.number,
                                                                  Employee.employee_code):
        if rem.status == "ON_HOLD" and run.status in FINAL_STATUSES:
            # A FULL hold keeps the employee out of finalization: a final register (and its totals) holds
            # finalized pay only; the held employee is listed by the exclusions report.
            continue
        row = _base_row(ctx, rem, run, period, employee, assignment)
        row.update(status=rem.status, hold=rem.hold_type, gross=_d(rem.gross), deductions=_d(rem.deductions),
                   reimbursements=_d(rem.reimbursements), net=_d(rem.net), employer_cost=_d(rem.employer_cost))
        rows.append(row)
    return rows


def _line_query(ctx: Ctx, *, categories: Optional[Sequence[str]] = None, kinds: Optional[Sequence[str]] = None):
    query = run_employee_query(ctx, PayrollLine, statuses=("CALCULATED",)).join(
        PayrollLine, and_(PayrollLine.company_id == PayrollRunEmployee.company_id,
                          PayrollLine.run_employee_id == PayrollRunEmployee.id))
    if categories is not None:
        query = query.filter(PayrollLine.category.in_(tuple(categories)))
    if kinds is not None:
        query = query.filter(PayrollLine.kind.in_(tuple(kinds)))
    return query


KIND_ORDER = {"EARNING": 1, "REIMBURSEMENT": 2, "BENEFIT_IN_KIND": 3, "DEDUCTION": 4, "EMPLOYER_CONTRIBUTION": 5,
              "ROUNDING": 6, "INFORMATIONAL": 7}


def gen_gross_to_net(ctx: Ctx) -> List[dict]:
    components: Dict[Tuple[str, str, str], dict] = {}
    for rem, run, period, employee, assignment, line in _line_query(ctx):
        key = (line.currency, line.kind, line.component_code)
        item = components.setdefault(key, {"currency": line.currency, "kind": line.kind,
                                           "component_code": line.component_code, "employees": set(),
                                           "amount": _ZERO, "order": line.sort_order or 0, "statutory": False,
                                           "fixture": False})
        item["employees"].add(rem.id)
        item["amount"] += _d(line.amount)
        item["statutory"] = item["statutory"] or bool(line.statutory)
        item["fixture"] = item["fixture"] or bool(line.fixture)
    summary: Dict[str, Dict[str, Any]] = {}
    for rem, *_ in run_employee_query(ctx, statuses=("CALCULATED",)):
        bucket = summary.setdefault(rem.currency, {"employees": 0, "gross": _ZERO, "deductions": _ZERO,
                                                   "reimbursements": _ZERO, "net": _ZERO, "employer_cost": _ZERO})
        bucket["employees"] += 1
        for key in ("gross", "deductions", "reimbursements", "net", "employer_cost"):
            bucket[key] += _d(getattr(rem, key))
    rows = []
    for (currency, kind, code), item in sorted(components.items(),
                                               key=lambda kv: (kv[0][0], KIND_ORDER.get(kv[0][1], 9), kv[1]["order"],
                                                               kv[0][2])):
        rows.append({"currency": currency, "section": kind, "component_code": code,
                     "employees": len(item["employees"]), "amount": item["amount"],
                     "label": FIXTURE_LABEL if item["fixture"] else ("Statutory" if item["statutory"] else None)})
    for currency, bucket in sorted(summary.items()):
        for key, label in (("gross", "GROSS"), ("deductions", "DEDUCTIONS"), ("reimbursements", "REIMBURSEMENTS"),
                           ("net", "NET"), ("employer_cost", "EMPLOYER_COST")):
            rows.append({"currency": currency, "section": "TOTAL", "component_code": label,
                         "employees": bucket["employees"], "amount": bucket[key], "label": None})
    return rows


def gen_variance(ctx: Ctx) -> List[dict]:
    current = list(run_employee_query(ctx, statuses=("CALCULATED",)).filter(PayrollRun.run_type == "REGULAR"))
    if not current:
        return []
    # The previous period of each pay group: the latest period that ends before this one starts.
    keys = {(run.pay_group_id, period.period_start) for _, run, period, *_ in current}
    previous_period: Dict[Tuple[int, date], Optional[int]] = {}
    for pay_group_id, start in keys:
        found = (ctx.db.query(PayPeriod.id)
                 .filter(PayPeriod.company_id == ctx.cid, PayPeriod.pay_group_id == pay_group_id,
                         PayPeriod.period_end < start)
                 .order_by(PayPeriod.period_end.desc()).first())
        previous_period[(pay_group_id, start)] = found[0] if found else None
    period_ids = sorted({p for p in previous_period.values() if p})
    previous: Dict[Tuple[int, int], PayrollRunEmployee] = {}
    if period_ids:
        for rem, run in (ctx.db.query(PayrollRunEmployee, PayrollRun)
                         .join(PayrollRun, and_(PayrollRun.company_id == PayrollRunEmployee.company_id,
                                                PayrollRun.id == PayrollRunEmployee.run_id))
                         .filter(PayrollRunEmployee.company_id == ctx.cid, PayrollRun.pay_period_id.in_(period_ids),
                                 PayrollRun.run_type == "REGULAR", PayrollRun.status.in_(FINAL_STATUSES),
                                 PayrollRunEmployee.status == "CALCULATED",
                                 access.run_scope_clause(ctx.principal, READ, PayrollRun.legal_entity_id,
                                                         PayrollRun.pay_group_id))):
            previous[(run.pay_period_id, rem.employee_id)] = rem
    rows = []
    for rem, run, period, employee, assignment in sorted(current, key=lambda r: (r[2].period_end, r[3].employee_code)):
        prev_period = previous_period.get((run.pay_group_id, period.period_start))
        prev = previous.get((prev_period, rem.employee_id)) if prev_period else None
        row = _base_row(ctx, rem, run, period, employee, assignment)
        same_currency = prev is not None and prev.currency == rem.currency
        prev_gross = _d(prev.gross) if same_currency else None
        prev_net = _d(prev.net) if same_currency else None
        net_delta = _d(rem.net) - prev_net if prev_net is not None else None
        percent = None
        if net_delta is not None and prev_net:
            percent = (net_delta * 100 / prev_net).quantize(Decimal("0.01"))
        flag = "NEW" if prev is None else ("CURRENCY_CHANGED" if not same_currency else None)
        row.update(previous_gross=prev_gross, gross=_d(rem.gross),
                   gross_delta=_d(rem.gross) - prev_gross if prev_gross is not None else None,
                   previous_net=prev_net, net=_d(rem.net), net_delta=net_delta, net_delta_percent=percent, flag=flag)
        rows.append(row)
    return rows


def gen_ytd(ctx: Ctx) -> List[dict]:
    f = ctx.filters
    query = (ctx.db.query(YtdBalance, Employee)
             .join(Employee, and_(Employee.company_id == YtdBalance.company_id, Employee.id == YtdBalance.employee_id))
             .filter(YtdBalance.company_id == ctx.cid,
                     access.run_scope_clause(ctx.principal, READ, YtdBalance.legal_entity_id, None)))
    if f.legal_entity_id is not None:
        query = query.filter(YtdBalance.legal_entity_id == f.legal_entity_id)
    if f.currency is not None:
        query = query.filter(YtdBalance.currency == f.currency)
    if f.tax_year is not None:
        query = query.filter(YtdBalance.tax_year == f.tax_year)
    rows = []
    for balance, employee in query.order_by(YtdBalance.tax_year.desc(), Employee.employee_code, YtdBalance.bucket):
        rows.append({"tax_year": balance.tax_year, "employee_code": employee.employee_code,
                     "employee_name": _employee_name(employee.first_name, employee.last_name),
                     "legal_entity": ctx.name("legal_entity", balance.legal_entity_id), "bucket": balance.bucket,
                     "amount": _d(balance.amount), "currency": balance.currency})
    return rows


def gen_corrected_view(ctx: Ctx) -> List[dict]:
    query = _line_query(ctx).filter(or_(PayrollLine.period_ref.isnot(None),
                                        PayrollRun.run_type.in_(("RETRO_ADJUSTMENT", "REVERSAL"))))
    items = list(query)
    original_runs: Set[int] = set()
    for rem, run, period, employee, assignment, line in items:
        ref = line.period_ref or {}
        if ref.get("run_id"):
            original_runs.add(int(ref["run_id"]))
        elif run.reverses_run_id:
            original_runs.add(run.reverses_run_id)
    originals: Dict[Tuple[int, int, str], Decimal] = defaultdict(lambda: _ZERO)
    numbers: Dict[int, str] = {}
    if original_runs:
        scope = access.run_scope_clause(ctx.principal, READ, PayrollRun.legal_entity_id, PayrollRun.pay_group_id)
        for number, run_id in ctx.db.query(PayrollRun.number, PayrollRun.id).filter(
                PayrollRun.company_id == ctx.cid, PayrollRun.id.in_(sorted(original_runs)), scope):
            numbers[run_id] = number
        if numbers:
            for run_id, employee_id, code, amount in (
                    ctx.db.query(PayrollLine.run_id, PayrollRunEmployee.employee_id, PayrollLine.component_code,
                                 PayrollLine.amount)
                    .join(PayrollRunEmployee, and_(PayrollRunEmployee.company_id == PayrollLine.company_id,
                                                   PayrollRunEmployee.id == PayrollLine.run_employee_id))
                    .filter(PayrollLine.company_id == ctx.cid, PayrollLine.run_id.in_(sorted(numbers)))):
                originals[(run_id, employee_id, code)] += _d(amount)
    rows = []
    for rem, run, period, employee, assignment, line in sorted(items, key=lambda r: (r[2].period_end, r[3].employee_code,
                                                                                     r[5].sort_order)):
        ref = line.period_ref or {}
        original_run = int(ref["run_id"]) if ref.get("run_id") else run.reverses_run_id
        visible = original_run in numbers
        original = originals.get((original_run, rem.employee_id, line.component_code)) if visible else None
        delta = _d(line.amount)
        rows.append({
            "correction_run": run.number, "correction_type": run.run_type, "period_end": period.period_end,
            "employee_code": employee.employee_code,
            "employee_name": _employee_name(employee.first_name, employee.last_name),
            "component_code": line.component_code,
            # A REVERSAL run reverses its own period (its lines reference the reversed lines, not a period).
            "original_period": f"{ref['start']} – {ref.get('end', '')}" if ref.get("start") else (
                f"{period.period_start.isoformat()} – {period.period_end.isoformat()}"
                if run.run_type == "REVERSAL" else None),
            "original_run": numbers.get(original_run) if visible else None,
            "original_amount": original, "delta": delta,
            "corrected_amount": (original + delta) if original is not None else None, "currency": line.currency,
        })
    return rows


# =========================================================================== cost group (salary analytics)


def _group_key(ctx: Ctx, run: PayrollRun, assignment: Optional[EmployeeAssignment]) -> Tuple[str, Optional[int]]:
    group_by = ctx.filters.group_by or ctx.default_group
    if group_by == "legal_entity":
        return group_by, run.legal_entity_id
    if group_by == "department":
        return group_by, assignment.department_id if assignment else None
    if group_by == "cost_center":
        return group_by, assignment.cost_center_id if assignment else None
    return "grade", assignment.grade_id if assignment else None


def gen_employer_cost(ctx: Ctx) -> List[dict]:
    ctx.default_group = "legal_entity"
    contributions: Dict[int, Decimal] = defaultdict(lambda: _ZERO)
    for rem, run, period, employee, assignment, line in _line_query(ctx, kinds=("EMPLOYER_CONTRIBUTION",)):
        contributions[rem.id] += _d(line.amount)
    groups: Dict[Tuple[str, Optional[int]], dict] = {}
    for rem, run, period, employee, assignment in run_employee_query(ctx, statuses=("CALCULATED",)):
        kind, key = _group_key(ctx, run, assignment)
        item = groups.setdefault((rem.currency, key), {"group": ctx.name(kind, key) or "(none)", "group_by": kind,
                                                        "currency": rem.currency, "employees": set(),
                                                        "gross": _ZERO, "employer_contributions": _ZERO,
                                                        "employer_cost": _ZERO})
        item["employees"].add(rem.employee_id)
        item["gross"] += _d(rem.gross)
        item["employer_contributions"] += contributions.get(rem.id, _ZERO)
        item["employer_cost"] += _d(rem.employer_cost)
    return [{**item, "employees": len(item["employees"])}
            for _, item in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[1]["group"]))]


def gen_ctc_vs_actual(ctx: Ctx) -> List[dict]:
    ctx.default_group = "grade"
    items = list(run_employee_query(ctx, PayGroup.pay_frequency, statuses=("CALCULATED",))
                 .join(PayGroup, and_(PayGroup.company_id == PayrollRun.company_id, PayGroup.id == PayrollRun.pay_group_id))
                 .filter(PayrollRun.run_type == "REGULAR"))
    assignment_ids = sorted({rem.employee_assignment_id for rem, *_ in items})
    plans: Dict[int, List[Tuple[CompensationAssignment, str]]] = defaultdict(list)
    if assignment_ids:
        for comp, structure_type in (
                ctx.db.query(CompensationAssignment, SalaryTemplate.type)
                .join(SalaryTemplateVersion, and_(SalaryTemplateVersion.company_id == CompensationAssignment.company_id,
                                                  SalaryTemplateVersion.id == CompensationAssignment.template_version_id))
                .join(SalaryTemplate, and_(SalaryTemplate.company_id == SalaryTemplateVersion.company_id,
                                           SalaryTemplate.id == SalaryTemplateVersion.template_id))
                .filter(CompensationAssignment.company_id == ctx.cid,
                        CompensationAssignment.assignment_id.in_(assignment_ids),
                        CompensationAssignment.status.in_(("ACTIVE", "SCHEDULED", "SUPERSEDED")))):
            plans[comp.assignment_id].append((comp, structure_type))
    groups: Dict[Tuple[str, Optional[int]], dict] = {}
    skipped = 0
    for rem, run, period, employee, assignment, frequency in items:
        plan = next((c for c, t in plans.get(rem.employee_assignment_id, [])
                     if t == "CTC_BASED" and c.effective_from <= period.period_end
                     and (c.effective_to is None or c.effective_to >= period.period_end)
                     and c.currency == rem.currency), None)
        per_year = PERIODS_PER_YEAR.get(frequency)
        if plan is None or per_year is None:
            skipped += 1
            continue
        kind, key = _group_key(ctx, run, assignment)
        item = groups.setdefault((rem.currency, key), {"group": ctx.name(kind, key) or "(none)", "group_by": kind,
                                                        "currency": rem.currency, "employees": set(),
                                                        "planned_ctc": _ZERO, "actual_cost": _ZERO})
        item["employees"].add(rem.employee_id)
        item["planned_ctc"] += currencies_quantize(_d(plan.pay_basis_amount) / per_year, rem.currency)
        item["actual_cost"] += _d(rem.employer_cost)
    if skipped:
        ctx.notes.append(f"{skipped} run employee(s) without a CTC-based salary for the period are not included.")
    rows = []
    for _, item in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[1]["group"])):
        rows.append({**item, "employees": len(item["employees"]),
                     "variance": item["actual_cost"] - item["planned_ctc"]})
    return rows


def currencies_quantize(amount: Decimal, currency: str) -> Decimal:
    return Money.of(amount, currency).round().amount


# =========================================================================== time and variable groups


def _lines_by_category(ctx: Ctx, categories: Sequence[str], *, kinds: Optional[Sequence[str]] = None) -> List[dict]:
    rows = []
    query = _line_query(ctx, categories=categories, kinds=kinds)
    for rem, run, period, employee, assignment, line in query.order_by(PayPeriod.period_end, Employee.employee_code,
                                                                       PayrollLine.sort_order):
        row = _base_row(ctx, rem, run, period, employee, assignment)
        row.update(component_code=line.component_code, category=line.category,
                   units=_d(line.units) if line.units is not None else None, amount=_d(line.amount),
                   currency=line.currency, explanation=(line.explanation or "")[:300])
        rows.append(row)
    return rows


def gen_overtime(ctx: Ctx) -> List[dict]:
    return _lines_by_category(ctx, ("OVERTIME",))


def gen_holiday_premiums(ctx: Ctx) -> List[dict]:
    return _lines_by_category(ctx, ("HOLIDAY_PREMIUM",))


def gen_late_penalties(ctx: Ctx) -> List[dict]:
    return _lines_by_category(ctx, ("PENALTY_IF_LEGAL",))


def gen_bonus(ctx: Ctx) -> List[dict]:
    return _lines_by_category(ctx, ("BONUS", "COMMISSION", "VARIABLE"))


def gen_leave_encashments(ctx: Ctx) -> List[dict]:
    return _lines_by_category(ctx, ("LEAVE_ENCASHMENT",))


def gen_arrears(ctx: Ctx) -> List[dict]:
    """Arrears and recoveries in the employee's perspective: an arrears earning is positive, a recovery (a
    recovery deduction, or a negative earning) negative. A REVERSAL run's lines reference the reversed lines,
    not an earlier period's arrears: they belong to the corrected view, never to this report."""
    rows = []
    query = (_line_query(ctx)
             .filter(or_(PayrollLine.category == "RETRO_ADJUSTMENT", PayrollLine.period_ref.isnot(None)),
                     PayrollRun.run_type != "REVERSAL"))
    for rem, run, period, employee, assignment, line in query.order_by(PayPeriod.period_end, Employee.employee_code):
        ref = line.period_ref or {}
        signed = -_d(line.amount) if line.kind == "DEDUCTION" else _d(line.amount)
        row = _base_row(ctx, rem, run, period, employee, assignment)
        row.update(component_code=line.component_code, original_period_start=ref.get("start"),
                   original_period_end=ref.get("end"), amount=signed, currency=line.currency,
                   direction="ARREARS" if signed >= 0 else "RECOVERY")
        rows.append(row)
    return rows


def gen_lop(ctx: Ctx) -> List[dict]:
    per_employee: Dict[int, dict] = {}
    for rem, run, period, employee, assignment, line in _line_query(ctx, kinds=("EARNING",)):
        ledger = [s for s in (line.sources or []) if isinstance(s, dict) and s.get("type") == "day_ledger"]
        if not ledger:
            continue
        item = per_employee.setdefault(rem.id, {**_base_row(ctx, rem, run, period, employee, assignment),
                                                "segments": {}, "dates": set(), "reduction": _ZERO})
        for source in ledger:
            segment = str(source.get("segment"))
            days = _d(source.get("lop_days"))
            item["segments"][segment] = max(item["segments"].get(segment, _ZERO), days)
            item["dates"].update(source.get("dates") or [])
        rate = line.rate
        if rate is not None and line.units is None and _d(rate) > _d(line.amount):
            item["reduction"] += _d(rate) - _d(line.amount)
    rows = []
    for item in sorted(per_employee.values(), key=lambda i: (i["period_end"], i["employee_code"])):
        lop_days = sum(item.pop("segments").values(), _ZERO)
        dates = item.pop("dates")
        if lop_days <= 0:
            continue
        item.update(lop_days=lop_days, lop_dates=len(dates), reduction=item["reduction"])
        rows.append(item)
    return rows


def gen_task_pay(ctx: Ctx) -> List[dict]:
    from app.models.pay_runs import PayrollLineSource

    rows = []
    query = (_line_query(ctx, categories=("TASK_EARNING",))
             .outerjoin(PayrollLineSource, and_(PayrollLineSource.company_id == PayrollLine.company_id,
                                                PayrollLineSource.line_id == PayrollLine.id,
                                                PayrollLineSource.source_type == "TASK"))
             .outerjoin(Task, and_(Task.company_id == PayrollLineSource.company_id,
                                   Task.id == PayrollLineSource.source_id))
             .add_columns(Task.project, Task.title, Task.id))
    seen: Set[Tuple[int, Optional[int]]] = set()
    for rem, run, period, employee, assignment, line, project, title, task_id in query.order_by(
            PayPeriod.period_end, Employee.employee_code):
        if (line.id, task_id) in seen:
            continue
        seen.add((line.id, task_id))
        row = _base_row(ctx, rem, run, period, employee, assignment)
        row.update(project=project or "(no project)", task=title, component_code=line.component_code,
                   units=_d(line.units) if line.units is not None else None, amount=_d(line.amount),
                   currency=line.currency)
        rows.append(row)
    return rows


# =========================================================================== money out (payments milestone tables)


def _payments():
    try:
        from app.models import payment_flows  # noqa: WPS433 (optional milestone)
    except ImportError:  # pragma: no cover - the payments milestone ships the module
        return None
    return payment_flows if hasattr(payment_flows, "PaymentInstruction") else None


def _masked_account(snapshot: Any) -> Optional[str]:
    data = masking.redact(snapshot or {})
    if not isinstance(data, dict):
        return None
    for key in ("account_masked", "masked_account", "iban_masked", "account", "iban"):
        value = data.get(key)
        if isinstance(value, str) and value:
            return value if "•" in value else masking.mask_value(value)
    last4 = data.get("last4")
    return f"{masking.MASK}{last4}" if isinstance(last4, str) and last4 else None


def _instruction_query(ctx: Ctx, statuses: Optional[Sequence[str]] = None):
    pf = _payments()
    f = ctx.filters
    Instr = pf.PaymentInstruction
    query = (ctx.db.query(Instr, PayrollRun, Employee)
             .outerjoin(PayrollRun, and_(PayrollRun.company_id == Instr.company_id, PayrollRun.id == Instr.run_id))
             .join(Employee, and_(Employee.company_id == Instr.company_id, Employee.id == Instr.employee_id))
             .filter(Instr.company_id == ctx.cid,
                     access.run_scope_clause(ctx.principal, READ, Instr.legal_entity_id, PayrollRun.pay_group_id)))
    if statuses is not None:
        query = query.filter(Instr.status.in_(tuple(statuses)))
    if f.status is not None:
        query = query.filter(Instr.status == f.status)
    if f.legal_entity_id is not None:
        query = query.filter(Instr.legal_entity_id == f.legal_entity_id)
    if f.pay_group_id is not None:
        query = query.filter(PayrollRun.pay_group_id == f.pay_group_id)
    if f.currency is not None:
        query = query.filter(Instr.currency == f.currency)
    if f.period_from is not None:
        query = query.filter(func.date(Instr.created_at) >= f.period_from)
    if f.period_to is not None:
        query = query.filter(func.date(Instr.created_at) <= f.period_to)
    return query


def _instruction_row(ctx: Ctx, instr, run: Optional[PayrollRun], employee: Employee) -> dict:
    try:
        label = sm.employee_label(instr.status, instr.submission_channel, instr.evidence_status).text
    except sm.InvalidTransition:
        label = instr.status
    return {
        "reference": instr.reference, "run_number": run.number if run else None,
        "employee_code": employee.employee_code,
        "employee_name": _employee_name(employee.first_name, employee.last_name),
        "legal_entity": ctx.name("legal_entity", instr.legal_entity_id), "amount": _d(instr.amount),
        "currency": instr.currency, "status": instr.status, "label": label,
        "submission_channel": instr.submission_channel, "evidence_status": instr.evidence_status,
        "account": _masked_account(instr.beneficiary_snapshot), "submitted_at": instr.submitted_at,
        "settled_at": instr.settled_at,
    }


def gen_payment_register(ctx: Ctx) -> List[dict]:
    Instr = _payments().PaymentInstruction
    return [_instruction_row(ctx, *row) for row in _instruction_query(ctx).order_by(Instr.id)]


def gen_failures(ctx: Ctx) -> List[dict]:
    Instr = _payments().PaymentInstruction
    rows = []
    for instr, run, employee in _instruction_query(ctx, ("FAILED", "RETURNED", "REJECTED", "REVERSED")).order_by(
            Instr.id):
        row = _instruction_row(ctx, instr, run, employee)
        row.update(failure_code=instr.failure_code, failure_message=instr.failure_message,
                   returned_at=instr.returned_at)
        rows.append(row)
    return rows


def gen_reconciliation(ctx: Ctx) -> List[dict]:
    pf = _payments()
    from app.models import payments as bank_models

    Funding = getattr(bank_models, "FundingAccount", None)
    if Funding is None:
        ctx.notes.append("Funding accounts are not available yet.")
        return []
    f = ctx.filters
    Line, Match = pf.BankStatementLine, pf.ReconciliationMatch
    query = (ctx.db.query(Line, Match, Funding.legal_entity_id)
             .join(Funding, and_(Funding.company_id == Line.company_id, Funding.id == Line.funding_account_id))
             .outerjoin(Match, and_(Match.company_id == Line.company_id, Match.statement_line_id == Line.id))
             .filter(Line.company_id == ctx.cid,
                     access.run_scope_clause(ctx.principal, READ, Funding.legal_entity_id, None)))
    if f.legal_entity_id is not None:
        query = query.filter(Funding.legal_entity_id == f.legal_entity_id)
    if f.currency is not None:
        query = query.filter(Line.currency == f.currency)
    if f.period_from is not None:
        query = query.filter(Line.value_date >= f.period_from)
    if f.period_to is not None:
        query = query.filter(Line.value_date <= f.period_to)
    if f.status is not None:
        query = query.filter(Match.status == f.status)
    rows = []
    for line, match, entity_id in query.order_by(Line.value_date, Line.id):
        rows.append({"value_date": line.value_date, "legal_entity": ctx.name("legal_entity", entity_id),
                     "direction": line.direction, "amount": _d(line.amount), "currency": line.currency,
                     "reference": masking.scrub_text(line.reference or ""),
                     "counterparty_account": line.counterparty_account_masked,
                     "match_status": match.status if match else "UNMATCHED",
                     "reason": match.reason if match else None,
                     "instruction_id": match.instruction_id if match else None,
                     "variance": _d(match.variance_amount) if match and match.variance_amount is not None else None})
    return rows


def gen_obligations(ctx: Ctx) -> List[dict]:
    pf = _payments()
    f = ctx.filters
    Obl = pf.PaymentObligation
    query = (ctx.db.query(Obl, Employee)
             .join(Employee, and_(Employee.company_id == Obl.company_id, Employee.id == Obl.employee_id))
             .filter(Obl.company_id == ctx.cid, Obl.status.in_(("OPEN", "PARTIAL", "UNCLAIMED")),
                     access.run_scope_clause(ctx.principal, READ, Obl.legal_entity_id, Obl.pay_group_id)))
    if f.legal_entity_id is not None:
        query = query.filter(Obl.legal_entity_id == f.legal_entity_id)
    if f.pay_group_id is not None:
        query = query.filter(Obl.pay_group_id == f.pay_group_id)
    if f.currency is not None:
        query = query.filter(Obl.currency == f.currency)
    if f.status is not None:
        query = query.filter(Obl.status == f.status)
    if f.period_from is not None:
        query = query.filter(Obl.pay_date >= f.period_from)
    if f.period_to is not None:
        query = query.filter(Obl.pay_date <= f.period_to)
    rows = []
    for obligation, employee in query.order_by(Obl.pay_date, Obl.id):
        rows.append({"source_type": obligation.source_type, "employee_code": employee.employee_code,
                     "employee_name": _employee_name(employee.first_name, employee.last_name),
                     "legal_entity": ctx.name("legal_entity", obligation.legal_entity_id),
                     "pay_date": obligation.pay_date, "amount": _d(obligation.amount), "currency": obligation.currency,
                     "status": obligation.status, "carried_reason": obligation.carried_reason})
    return rows


def gen_exclusions(ctx: Ctx) -> List[dict]:
    rows = []
    query = run_employee_query(ctx, statuses=("EXCLUDED", "ON_HOLD"))
    if ctx.filters.status is None:
        query = query.filter(PayrollRun.status.in_(STAGES["ALL"]))
    for rem, run, period, employee, assignment in query.order_by(PayPeriod.period_end, PayrollRun.number,
                                                                  Employee.employee_code):
        row = _base_row(ctx, rem, run, period, employee, assignment)
        hold = rem.hold or {}
        row.update(status=rem.status, reason=rem.exclusion_reason or hold.get("reason_code"),
                   note=(rem.exclusion_note or hold.get("note") or "")[:300], hold_type=rem.hold_type)
        rows.append(row)
    return rows


# =========================================================================== ledgers


def gen_loans(ctx: Ctx) -> List[dict]:
    from app.payroll.services import loans as loans_service

    f = ctx.filters
    query = (ctx.db.query(Loan, Employee)
             .join(Employee, and_(Employee.company_id == Loan.company_id, Employee.id == Loan.employee_id))
             .filter(Loan.company_id == ctx.cid,
                     access.run_scope_clause(ctx.principal, READ, Loan.legal_entity_id, Loan.pay_group_id),
                     Loan.status.in_((f.status,) if f.status else ("ACTIVE", "PAUSED", "APPROVED"))))
    if f.legal_entity_id is not None:
        query = query.filter(Loan.legal_entity_id == f.legal_entity_id)
    if f.pay_group_id is not None:
        query = query.filter(Loan.pay_group_id == f.pay_group_id)
    if f.currency is not None:
        query = query.filter(Loan.currency == f.currency)
    rows = []
    for loan, employee in query.order_by(Employee.employee_code, Loan.id):
        rows.append({"employee_code": employee.employee_code,
                     "employee_name": _employee_name(employee.first_name, employee.last_name),
                     "legal_entity": ctx.name("legal_entity", loan.legal_entity_id), "kind": loan.kind,
                     "status": loan.status, "principal": _d(loan.principal), "balance": loans_service.balance(ctx.db, loan),
                     "currency": loan.currency, "tenure": loan.tenure, "disbursed_on": loan.disbursed_on})
    return rows


def gen_reimbursements(ctx: Ctx) -> List[dict]:
    f = ctx.filters
    Claim = ReimbursementClaim
    query = (ctx.db.query(Claim, Employee)
             .join(Employee, and_(Employee.company_id == Claim.company_id, Employee.id == Claim.employee_id))
             .filter(Claim.company_id == ctx.cid,
                     access.run_scope_clause(ctx.principal, READ, Claim.legal_entity_id, Claim.pay_group_id)))
    for attr, value in (("legal_entity_id", f.legal_entity_id), ("pay_group_id", f.pay_group_id),
                        ("payout_currency", f.currency), ("status", f.status)):
        if value is not None:
            query = query.filter(getattr(Claim, attr) == value)
    if f.period_from is not None:
        query = query.filter(Claim.expense_date >= f.period_from)
    if f.period_to is not None:
        query = query.filter(Claim.expense_date <= f.period_to)
    rows = []
    for claim, employee in query.order_by(Claim.expense_date, Claim.id):
        rows.append({"expense_date": claim.expense_date, "employee_code": employee.employee_code,
                     "employee_name": _employee_name(employee.first_name, employee.last_name),
                     "claim_type": claim.claim_type_name, "status": claim.status,
                     "claimed": _d(claim.payout_amount),
                     "approved": _d(claim.approved_amount) if claim.approved_amount is not None else None,
                     "currency": claim.payout_currency, "payout_route": claim.payout_route, "paid_ref": claim.paid_ref})
    return rows


# =========================================================================== compliance


def gen_liabilities(ctx: Ctx) -> List[dict]:
    f = ctx.filters
    query = (ctx.db.query(StatutoryLiability, PayrollRun.number, PayPeriod.period_end)
             .join(PayrollRun, and_(PayrollRun.company_id == StatutoryLiability.company_id,
                                    PayrollRun.id == StatutoryLiability.run_id))
             .join(PayPeriod, and_(PayPeriod.company_id == PayrollRun.company_id,
                                   PayPeriod.id == PayrollRun.pay_period_id))
             .filter(StatutoryLiability.company_id == ctx.cid,
                     access.run_scope_clause(ctx.principal, READ, PayrollRun.legal_entity_id, PayrollRun.pay_group_id)))
    for col, value in ((StatutoryLiability.legal_entity_id, f.legal_entity_id), (PayrollRun.pay_group_id, f.pay_group_id),
                       (StatutoryLiability.currency, f.currency), (StatutoryLiability.status, f.status)):
        if value is not None:
            query = query.filter(col == value)
    if f.period_from is not None:
        query = query.filter(PayPeriod.period_end >= f.period_from)
    if f.period_to is not None:
        query = query.filter(PayPeriod.period_start <= f.period_to)
    rows = []
    for liability, number, period_end in query.order_by(StatutoryLiability.due_date.asc().nullslast(),
                                                         StatutoryLiability.id):
        rows.append({"run_number": number, "period_end": period_end,
                     "legal_entity": ctx.name("legal_entity", liability.legal_entity_id), "scheme": liability.scheme,
                     "payee": liability.payee_authority, "employee_amount": _d(liability.employee_amount),
                     "employer_amount": _d(liability.employer_amount),
                     "total": _d(liability.employee_amount) + _d(liability.employer_amount),
                     "currency": liability.currency, "due_date": liability.due_date, "status": liability.status,
                     "remittance_ref": liability.remittance_ref,
                     "label": FIXTURE_LABEL if liability.fixture else None})
    return rows


def gen_statutory_filings(ctx: Ctx) -> List[dict]:
    f = ctx.filters
    query = ctx.db.query(StatutoryReport).filter(
        StatutoryReport.company_id == ctx.cid,
        access.run_scope_clause(ctx.principal, READ, StatutoryReport.legal_entity_id, None))
    for col, value in ((StatutoryReport.legal_entity_id, f.legal_entity_id), (StatutoryReport.status, f.status)):
        if value is not None:
            query = query.filter(col == value)
    if f.period_from is not None:
        query = query.filter(StatutoryReport.period_end >= f.period_from)
    if f.period_to is not None:
        query = query.filter(StatutoryReport.period_start <= f.period_to)
    return [{"legal_entity": ctx.name("legal_entity", r.legal_entity_id), "report_code": r.report_code,
             "report_name": r.report_name, "period_start": r.period_start, "period_end": r.period_end,
             "version": r.version, "status": r.status, "due_date": r.due_date, "filing_ref": r.filing_ref,
             "acknowledgement_ref": r.acknowledgement_ref, "pack_version": r.pack_version,
             "label": r.watermark}
            for r in query.order_by(StatutoryReport.period_start, StatutoryReport.report_code, StatutoryReport.version)]


# =========================================================================== audit group


def _audit_rows(ctx: Ctx, modules: Sequence[str], *, entity_types: Optional[Sequence[str]] = None) -> List[dict]:
    f = ctx.filters
    query = ctx.db.query(AuditLog).filter(AuditLog.company_id == ctx.cid, AuditLog.module.in_(tuple(modules)))
    if entity_types is not None:
        query = query.filter(AuditLog.entity_type.in_(tuple(entity_types)))
    if f.period_from is not None:
        query = query.filter(func.date(AuditLog.created_at) >= f.period_from)
    if f.period_to is not None:
        query = query.filter(func.date(AuditLog.created_at) <= f.period_to)
    rows = []
    for row in query.order_by(AuditLog.created_at, AuditLog.id).limit(MAX_ROWS + 1):
        rows.append({"at": row.created_at, "actor": row.actor_email or row.actor_type, "module": row.module,
                     "action": row.action, "entity_type": row.entity_type, "entity_id": row.entity_id,
                     "description": masking.scrub_text(row.description or ""), "ip_address": row.ip_address})
    return rows


def gen_settings_changes(ctx: Ctx) -> List[dict]:
    from app.payroll.services.settings import SETTINGS_MODULE

    return _audit_rows(ctx, (payroll_audit.PAYROLL_SETTINGS, payroll_audit.PAYROLL_POLICIES, SETTINGS_MODULE))


def gen_bank_events(ctx: Ctx) -> List[dict]:
    return _audit_rows(ctx, (payroll_audit.BANK_ACCOUNTS,))


def gen_compensation_changes(ctx: Ctx) -> List[dict]:
    f = ctx.filters
    query = (ctx.db.query(CompensationAssignment, EmployeeAssignment, Employee)
             .join(EmployeeAssignment, and_(EmployeeAssignment.company_id == CompensationAssignment.company_id,
                                         EmployeeAssignment.id == CompensationAssignment.assignment_id))
             .join(Employee, and_(Employee.company_id == CompensationAssignment.company_id,
                                  Employee.id == CompensationAssignment.employee_id))
             .filter(CompensationAssignment.company_id == ctx.cid,
                     access.run_scope_clause(ctx.principal, READ, EmployeeAssignment.legal_entity_id,
                                             EmployeeAssignment.pay_group_id),
                     # Salary data also needs compensation access at entity / pay group / company scope:
                     # a SELF grant (every employee has one) reaches no other person's salary here.
                     access.run_scope_clause(ctx.principal, COMPENSATION_READ, EmployeeAssignment.legal_entity_id,
                                             EmployeeAssignment.pay_group_id)))
    for col, value in ((EmployeeAssignment.legal_entity_id, f.legal_entity_id),
                       (EmployeeAssignment.pay_group_id, f.pay_group_id),
                       (EmployeeAssignment.department_id, f.department_id),
                       (EmployeeAssignment.cost_center_id, f.cost_center_id),
                       (CompensationAssignment.currency, f.currency), (CompensationAssignment.status, f.status)):
        if value is not None:
            query = query.filter(col == value)
    if f.period_from is not None:
        query = query.filter(CompensationAssignment.effective_from >= f.period_from)
    if f.period_to is not None:
        query = query.filter(CompensationAssignment.effective_from <= f.period_to)
    rows = []
    for comp, assignment, employee in query.order_by(CompensationAssignment.effective_from, Employee.employee_code):
        rows.append({"effective_from": comp.effective_from, "effective_to": comp.effective_to,
                     "employee_code": employee.employee_code,
                     "employee_name": _employee_name(employee.first_name, employee.last_name),
                     "legal_entity": ctx.name("legal_entity", assignment.legal_entity_id), "reason": comp.reason,
                     "status": comp.status, "pay_basis_amount": _d(comp.pay_basis_amount), "currency": comp.currency,
                     "approved_at": comp.approved_at})
    return rows


# =========================================================================== the catalogue

_M = "money"
_COMMON = (Col("run_number", "Run"), Col("run_stage", "Stage"), Col("period_end", "Period end", "date"),
           Col("employee_code", "Employee code"), Col("employee_name", "Employee"),
           Col("legal_entity", "Legal entity"), Col("pay_group", "Pay group"))
_LINE_COLS = _COMMON + (Col("component_code", "Component"), Col("units", "Units", "decimal"),
                        Col("amount", "Amount", _M), Col("currency", "Currency"))


def _def(code, group, name, description, columns, generator, **kw) -> ReportDef:
    return ReportDef(code, group, name, description, tuple(columns), generator, **kw)


CATALOGUE: Tuple[ReportDef, ...] = (
    _def("payroll_register", "Payroll", "Payroll register", "Every employee of the selected runs with gross, "
         "deductions, reimbursements, net and employer cost.",
         _COMMON + (Col("department", "Department"), Col("cost_center", "Cost center"), Col("status", "Status"),
                    Col("hold", "Hold"), Col("gross", "Gross", _M), Col("deductions", "Deductions", _M),
                    Col("reimbursements", "Reimbursements", _M), Col("net", "Net", _M),
                    Col("employer_cost", "Employer cost", _M), Col("currency", "Currency")),
         gen_payroll_register, run_based=True),
    _def("gross_to_net", "Payroll", "Gross-to-net summary", "Totals per component and the gross-to-net totals, "
         "per currency.",
         (Col("currency", "Currency"), Col("section", "Section"), Col("component_code", "Component"),
          Col("employees", "Employees", "int", total=False), Col("amount", "Amount", _M, total=False),
          Col("label", "Label")),
         gen_gross_to_net, run_based=True),
    _def("variance", "Payroll", "Variance vs previous period", "Gross and net of each employee against the "
         "previous finalized regular period of the pay group.",
         _COMMON + (Col("previous_gross", "Previous gross", _M), Col("gross", "Gross", _M),
                    Col("gross_delta", "Gross change", _M), Col("previous_net", "Previous net", _M),
                    Col("net", "Net", _M), Col("net_delta", "Net change", _M),
                    Col("net_delta_percent", "Net change %", "decimal", total=False), Col("flag", "Flag"),
                    Col("currency", "Currency")),
         gen_variance, run_based=True),
    _def("ytd_by_employee", "Payroll", "YTD by employee", "Year-to-date balances per employee, legal entity and tax "
         "year (EC-STA-08).",
         (Col("tax_year", "Tax year"), Col("employee_code", "Employee code"), Col("employee_name", "Employee"),
          Col("legal_entity", "Legal entity"), Col("bucket", "Bucket"), Col("amount", "Amount", _M),
          Col("currency", "Currency")),
         gen_ytd, filters=("legal_entity", "currency", "tax_year")),
    _def("corrected_view", "Payroll", "Corrected view (original + deltas)", "Retro, reversal and arrears lines with "
         "the original amount, the delta and the corrected amount.",
         (Col("correction_run", "Correction run"), Col("correction_type", "Type"),
          Col("period_end", "Period end", "date"), Col("employee_code", "Employee code"),
          Col("employee_name", "Employee"), Col("component_code", "Component"),
          Col("original_period", "Original period"), Col("original_run", "Original run"),
          Col("original_amount", "Original", _M), Col("delta", "Delta", _M),
          Col("corrected_amount", "Corrected", _M), Col("currency", "Currency")),
         gen_corrected_view, run_based=True),
    _def("employer_cost", "Cost", "Employer cost", "Employer cost by legal entity, department, cost center or grade "
         "(group_by). Small groups are hidden.",
         (Col("group_by", "Grouped by"), Col("group", "Group"), Col("employees", "Employees", "int"),
          Col("gross", "Gross", _M), Col("employer_contributions", "Employer contributions", _M),
          Col("employer_cost", "Employer cost", _M), Col("currency", "Currency")),
         gen_employer_cost, analytics=True, run_based=True,
         filters=("period", "legal_entity", "pay_group", "department", "cost_center", "currency", "status",
                  "group_by")),
    _def("ctc_vs_actual", "Cost", "CTC vs actual", "Planned CTC for the period (annual CTC ÷ periods per year) against "
         "the actual employer cost, by grade or another group. Small groups are hidden.",
         (Col("group_by", "Grouped by"), Col("group", "Group"), Col("employees", "Employees", "int"),
          Col("planned_ctc", "Planned CTC", _M), Col("actual_cost", "Actual cost", _M),
          Col("variance", "Variance", _M), Col("currency", "Currency")),
         gen_ctc_vs_actual, analytics=True, run_based=True,
         filters=("period", "legal_entity", "pay_group", "department", "cost_center", "currency", "status",
                  "group_by")),
    _def("overtime_by_bucket", "Time", "Overtime by bucket", "Overtime lines per employee and bucket with hours.",
         _LINE_COLS, gen_overtime, run_based=True),
    _def("lop_and_absence", "Time", "LOP and absence", "Loss-of-pay days per employee and the reduction on "
         "LOP-applicable lines (includes any other proration of those lines).",
         _COMMON + (Col("lop_days", "LOP days", "decimal"), Col("lop_dates", "LOP dates", "int"),
                    Col("reduction", "Prorated reduction", _M), Col("currency", "Currency")),
         gen_lop, run_based=True),
    _def("holiday_premiums", "Time", "Holiday-work premiums", "Holiday and weekly-off work premium lines.",
         _LINE_COLS, gen_holiday_premiums, run_based=True),
    _def("late_penalties", "Time", "Late penalties", "Late-arrival penalty deductions (only where lawful).",
         _LINE_COLS, gen_late_penalties, run_based=True),
    _def("bonus_incentives", "Variable", "Bonus & incentives", "Bonus, commission and variable pay lines.",
         _LINE_COLS, gen_bonus, run_based=True),
    _def("task_pay", "Variable", "Task pay by project and employee", "Task earnings with their project.",
         _COMMON + (Col("project", "Project"), Col("task", "Task"), Col("component_code", "Component"),
                    Col("units", "Units", "decimal"), Col("amount", "Amount", _M), Col("currency", "Currency")),
         gen_task_pay, run_based=True),
    _def("arrears", "Variable", "Arrears", "Arrears and recoveries with the original period they belong to.",
         _COMMON + (Col("component_code", "Component"), Col("original_period_start", "Original period start"),
                    Col("original_period_end", "Original period end"), Col("direction", "Direction"),
                    Col("amount", "Amount", _M), Col("currency", "Currency")),
         gen_arrears, run_based=True),
    _def("payment_register", "Money out", "Payment register", "Payment instructions with their status and evidence "
         "(paid only with verified evidence). Account numbers are masked.",
         (Col("reference", "Reference"), Col("run_number", "Run"), Col("employee_code", "Employee code"),
          Col("employee_name", "Employee"), Col("legal_entity", "Legal entity"), Col("amount", "Amount", _M),
          Col("currency", "Currency"), Col("status", "Status"), Col("label", "Shown to employee"),
          Col("submission_channel", "Channel"), Col("evidence_status", "Evidence"),
          Col("account", "Account", sensitive=True), Col("submitted_at", "Submitted", "datetime"),
          Col("settled_at", "Settled", "datetime")),
         gen_payment_register, requires="payments"),
    _def("failures_returns", "Money out", "Failures and returns", "Failed, rejected, returned and reversed payments.",
         (Col("reference", "Reference"), Col("run_number", "Run"), Col("employee_code", "Employee code"),
          Col("employee_name", "Employee"), Col("amount", "Amount", _M), Col("currency", "Currency"),
          Col("status", "Status"), Col("failure_code", "Failure code"), Col("failure_message", "Message"),
          Col("account", "Account", sensitive=True), Col("returned_at", "Returned", "datetime")),
         gen_failures, requires="payments"),
    _def("reconciliation_statement", "Money out", "Reconciliation statement", "Bank statement lines with their "
         "reconciliation decision.",
         (Col("value_date", "Value date", "date"), Col("legal_entity", "Legal entity"), Col("direction", "Direction"),
          Col("amount", "Amount", _M), Col("currency", "Currency"), Col("reference", "Reference"),
          Col("counterparty_account", "Counterparty account", sensitive=True), Col("match_status", "Match"),
          Col("reason", "Reason"), Col("instruction_id", "Instruction", "int", total=False),
          Col("variance", "Variance", _M)),
         gen_reconciliation, requires="payments"),
    _def("outstanding_obligations", "Money out", "Outstanding obligations", "Net pay and other obligations that are "
         "not fully settled (open, partial, unclaimed).",
         (Col("source_type", "Source"), Col("employee_code", "Employee code"), Col("employee_name", "Employee"),
          Col("legal_entity", "Legal entity"), Col("pay_date", "Pay date", "date"), Col("amount", "Amount", _M),
          Col("currency", "Currency"), Col("status", "Status"), Col("carried_reason", "Carried because")),
         gen_obligations, requires="payments"),
    _def("exclusions", "Money out", "Exclusions", "Employees excluded from runs or on hold, with the reason.",
         _COMMON + (Col("status", "Status"), Col("reason", "Reason"), Col("hold_type", "Hold"), Col("note", "Note")),
         gen_exclusions, run_based=True),
    _def("loans_outstanding", "Ledgers", "Loans outstanding", "Active loans and advances with their ledger balance.",
         (Col("employee_code", "Employee code"), Col("employee_name", "Employee"), Col("legal_entity", "Legal entity"),
          Col("kind", "Kind"), Col("status", "Status"), Col("principal", "Principal", _M),
          Col("balance", "Balance", _M), Col("currency", "Currency"), Col("tenure", "Tenure", "int", total=False),
          Col("disbursed_on", "Disbursed", "date")),
         gen_loans, filters=("legal_entity", "pay_group", "currency", "status")),
    _def("reimbursements", "Ledgers", "Reimbursements", "Expense claims with claimed and approved amounts.",
         (Col("expense_date", "Expense date", "date"), Col("employee_code", "Employee code"),
          Col("employee_name", "Employee"), Col("claim_type", "Type"), Col("status", "Status"),
          Col("claimed", "Claimed", _M), Col("approved", "Approved", _M), Col("currency", "Currency"),
          Col("payout_route", "Paid via"), Col("paid_ref", "Paid reference")),
         gen_reimbursements, filters=("period", "legal_entity", "pay_group", "currency", "status")),
    _def("leave_encashments", "Ledgers", "Leave encashments", "Leave encashment lines.",
         _LINE_COLS, gen_leave_encashments, run_based=True),
    _def("liabilities_remittances", "Compliance", "Liabilities and remittances", "Statutory liabilities (employee "
         "and employer portions) with due dates and remittance status.",
         (Col("run_number", "Run"), Col("period_end", "Period end", "date"), Col("legal_entity", "Legal entity"),
          Col("scheme", "Scheme"), Col("payee", "Payee"), Col("employee_amount", "Employee portion", _M),
          Col("employer_amount", "Employer portion", _M), Col("total", "Total", _M), Col("currency", "Currency"),
          Col("due_date", "Due", "date"), Col("status", "Status"), Col("remittance_ref", "Remittance reference"),
          Col("label", "Label")),
         gen_liabilities, filters=("period", "legal_entity", "pay_group", "currency", "status")),
    _def("statutory_filings", "Compliance", "Pack statutory reports and filing status", "Statutory reports with "
         "their lifecycle status, filing reference and acknowledgement.",
         (Col("legal_entity", "Legal entity"), Col("report_code", "Report"), Col("report_name", "Name"),
          Col("period_start", "Period start", "date"), Col("period_end", "Period end", "date"),
          Col("version", "Version", "int", total=False), Col("status", "Status"), Col("due_date", "Due", "date"),
          Col("filing_ref", "Filing reference"), Col("acknowledgement_ref", "Acknowledgement"),
          Col("pack_version", "Pack version"), Col("label", "Label")),
         gen_statutory_filings, filters=("period", "legal_entity", "status")),
    _def("settings_policy_changes", "Audit", "Settings and policy changes", "Payroll settings and policy events from "
         "the audit trail.",
         (Col("at", "At", "datetime"), Col("actor", "Actor", sensitive=True), Col("module", "Module"),
          Col("action", "Action"), Col("entity_type", "Object"), Col("entity_id", "Object id"),
          Col("description", "Description"), Col("ip_address", "IP address", sensitive=True)),
         gen_settings_changes, extra_permissions=(AUDIT_READ,), filters=("period",)),
    _def("compensation_changes", "Audit", "Compensation changes", "Salary assignments and revisions by effective date.",
         (Col("effective_from", "Effective from", "date"), Col("effective_to", "Effective to", "date"),
          Col("employee_code", "Employee code"), Col("employee_name", "Employee"),
          Col("legal_entity", "Legal entity"), Col("reason", "Reason"), Col("status", "Status"),
          Col("pay_basis_amount", "Pay basis", _M), Col("currency", "Currency"),
          Col("approved_at", "Approved", "datetime")),
         gen_compensation_changes, extra_permissions=(COMPENSATION_READ,)),
    _def("bank_detail_events", "Audit", "Bank-detail views and changes", "Bank account events from the audit trail: "
         "changes, approvals, verifications and every unmasked view.",
         (Col("at", "At", "datetime"), Col("actor", "Actor", sensitive=True), Col("action", "Action"),
          Col("entity_type", "Object"), Col("entity_id", "Object id"), Col("description", "Description"),
          Col("ip_address", "IP address", sensitive=True)),
         gen_bank_events, extra_permissions=(AUDIT_READ,), filters=("period",)),
)
BY_CODE: Dict[str, ReportDef] = {r.code: r for r in CATALOGUE}
GROUPS = ("Payroll", "Cost", "Time", "Variable", "Money out", "Ledgers", "Compliance", "Audit")


# =========================================================================== running a report


def _availability(report: ReportDef) -> Tuple[bool, Optional[str]]:
    if report.requires == "payments" and _payments() is None:
        return False, "Payments are not available yet."
    return True, None


def _can_run(principal: Principal, report: ReportDef) -> bool:
    return principal.has(READ) and all(principal.has(p) for p in report.extra_permissions)


def catalogue(db: Session, principal: Principal) -> dict:
    if not principal.has(READ):
        raise forbidden()
    groups: Dict[str, List[dict]] = {g: [] for g in GROUPS}
    for report in CATALOGUE:
        if not all(principal.has(p) for p in report.extra_permissions):
            continue
        available, reason = _availability(report)
        groups[report.group].append({
            "code": report.code, "name": report.name, "description": report.description,
            "filters": list(report.filters), "analytics": report.analytics, "available": available,
            "unavailable_reason": reason, "exportable": principal.has(EXPORT), "formats": list(FORMATS),
            "columns": [{"key": c.key, "label": c.label, "type": c.type, "sensitive": c.sensitive}
                        for c in report.columns],
        })
    return {"groups": [{"name": g, "reports": items} for g, items in groups.items() if items],
            "preview_rows": PREVIEW_ROWS}


def _report(principal: Principal, code: str, *, export: bool = False) -> ReportDef:
    report = BY_CODE.get(code)
    if report is None:
        raise not_found("Report")
    if not principal.has(READ) or (export and not principal.has(EXPORT)):
        raise forbidden()
    if not all(principal.has(p) for p in report.extra_permissions):
        raise forbidden()
    available, reason = _availability(report)
    if not available:
        raise conflict(reason or "This report is not available.", "REPORT_NOT_AVAILABLE")
    return report


def suppression_size(db: Session, company_id: int, on: date) -> int:
    try:
        value = settings_engine.resolve(db, company_id, SUPPRESSION_KEY, on=on).value
    except settings_engine.SettingConfigError:
        value = None
    return int(value) if isinstance(value, int) and value >= 1 else 5


def suppress(rows: List[dict], money_keys: Sequence[str], k: int) -> int:
    """Hide money of groups with fewer than ``k`` employees (complementary within a currency). Returns hidden."""
    by_currency: Dict[str, List[dict]] = defaultdict(list)
    for row in rows:
        by_currency[row.get("currency") or ""].append(row)
    hidden = 0
    for group in by_currency.values():
        small = [r for r in group if r.get("employees", 0) < k]
        if len(small) == 1 and len(group) > 1:
            others = sorted((r for r in group if r not in small), key=lambda r: r.get("employees", 0))
            small.append(others[0])
        for row in small:
            for key in money_keys:
                row[key] = None
            row["employees"] = None
            row["suppressed"] = True
            hidden += 1
    return hidden


def _totals(report: ReportDef, rows: Sequence[dict]) -> Dict[str, Dict[str, Decimal]]:
    totals: Dict[str, Dict[str, Decimal]] = {}
    for row in rows:
        currency = row.get("currency")
        if not currency:
            continue
        bucket = totals.setdefault(currency, {})
        for col in report.money_columns:
            value = row.get(col.key)
            if value is not None:
                bucket[col.key] = bucket.get(col.key, _ZERO) + Decimal(value)
    return totals


def _converted(ctx: Ctx, report: ReportDef, totals: Mapping[str, Mapping[str, Decimal]]) -> Optional[dict]:
    target = ctx.filters.convert_to
    if target is None or not totals:
        return None
    on = ctx.filters.fx_date or ctx.filters.period_to or ctx.today
    rates: List[fxlib.Rate] = []
    issues = []
    for currency in sorted(totals):
        if currency == target:
            continue
        found = fx_service.lookup(ctx.db, ctx.cid, currency, target, on)
        if not found.ok:
            issues.append({"currency": currency, "code": found.issue.code, "message": found.issue.message})
        else:
            rates.append(found.rate)
    if issues:
        return {"currency": target, "available": False, "issues": issues, "date": on,
                "message": "No usable FX rate for every currency; amounts stay per currency (never converted 1:1)."}
    result: Dict[str, Any] = {"currency": target, "available": True, "date": on, "totals": {}}
    label = None
    for col in report.money_columns:
        amounts = [Money.of(values[col.key], currency).round() for currency, values in totals.items()
                   if col.key in values]
        if not amounts:
            continue
        converted = fxlib.converted_total(amounts, target, rates, on=on, purpose="REPORTING")
        result["totals"][col.key] = converted.total.to_json()
        label = converted.label
    result["label"] = label
    result["rates"] = [{"base": r.base, "quote": r.quote, "rate": format(r.rate, "f"), "source": r.source,
                        "as_of": r.as_of.isoformat()} for r in rates]
    return result


@dataclass
class Result:
    report: ReportDef
    rows: List[dict]
    totals: Dict[str, Dict[str, Decimal]]
    converted: Optional[dict]
    suppressed: int
    notes: List[str]
    filters: Filters
    fixture: bool


def run_report(db: Session, principal: Principal, code: str, filters: Filters, *, export: bool = False) -> Result:
    report = _report(principal, code, export=export)
    filters = filters.validate()
    today = clock.company_today(db, principal.company_id)
    ctx = Ctx(db=db, principal=principal, filters=filters, today=today, report=report)
    rows = report.generator(ctx)
    if len(rows) > MAX_ROWS:
        raise conflict(f"The report has more than {MAX_ROWS} rows; narrow the filters.", "REPORT_TOO_LARGE")
    suppressed = 0
    if report.analytics:
        k = suppression_size(db, principal.company_id, today)
        suppressed = suppress(rows, [c.key for c in report.money_columns], k)
        if suppressed:
            ctx.notes.append(f"{suppressed} group(s) with fewer than {k} employees are hidden "
                             "(privacy.small_group_suppression_size).")
    totals = _totals(report, rows)
    if report.analytics:
        employees = sum(r.get("employees") or 0 for r in rows) + sum(1 for r in rows if r.get("suppressed"))
        k = suppression_size(db, principal.company_id, today)
        if rows and employees < k:
            totals = {}
            ctx.notes.append("Totals are hidden: the whole selection is smaller than the suppression size.")
    converted = _converted(ctx, report, totals)
    fixture = any(r.get("label") == FIXTURE_LABEL for r in rows)
    return Result(report, rows, totals, converted, suppressed, ctx.notes, filters, fixture)


def _cell_json(col: Col, value: Any, row: Mapping[str, Any]) -> Any:
    if value is None:
        return None
    if col.type == "money":
        currency = row.get("currency")
        return {"amount": amount_str(Decimal(value), currency), "currency": currency}
    if col.type == "decimal":
        return format(Decimal(value), "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _rows_json(report: ReportDef, rows: Sequence[dict]) -> List[dict]:
    result = []
    for row in rows:
        item = {c.key: _cell_json(c, row.get(c.key), row) for c in report.columns}
        if row.get("suppressed"):
            item["suppressed"] = True
        result.append(item)
    return result


def _totals_json(totals: Mapping[str, Mapping[str, Decimal]]) -> Dict[str, Dict[str, str]]:
    return {cur: {key: amount_str(value, cur) for key, value in sorted(values.items())}
            for cur, values in sorted(totals.items())}


def preview(db: Session, principal: Principal, code: str, filters: Filters) -> dict:
    result = run_report(db, principal, code, filters)
    report = result.report
    return {
        "report": {"code": report.code, "name": report.name, "group": report.group, "analytics": report.analytics},
        "filters": result.filters.to_payload(),
        "columns": [{"key": c.key, "label": c.label, "type": c.type, "sensitive": c.sensitive}
                    for c in report.columns],
        "rows": _rows_json(report, result.rows[:PREVIEW_ROWS]),
        "row_count": len(result.rows),
        "truncated": len(result.rows) > PREVIEW_ROWS,
        "totals": _totals_json(result.totals),
        "converted": result.converted,
        "suppressed_groups": result.suppressed,
        "notes": result.notes,
        "watermark": FIXTURE_LABEL if result.fixture else None,
    }


# =========================================================================== exports


def full_access(principal: Principal) -> bool:
    """Exports show sensitive columns unmasked only to exporters with company-wide export rights."""
    return principal.has_company_scope(EXPORT)


def _display(col: Col, value: Any, row: Mapping[str, Any], *, mask: bool) -> Any:
    if value is None:
        return "" if col.type != "money" else ("(hidden)" if row.get("suppressed") else "")
    if col.sensitive and mask:
        return masking.mask_value(str(value))
    if col.type == "money":
        return amount_str(Decimal(value), row.get("currency"))
    if col.type == "decimal":
        return format(Decimal(value), "f")
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return value


def _header_lines(result: Result, user_label: str, at: datetime) -> List[str]:
    lines = [CONFIDENTIAL.format(user=user_label, time=at.strftime("%Y-%m-%d %H:%M UTC"))]
    if result.fixture:
        lines.append(FIXTURE_LABEL)
    lines.append(f"{result.report.name} — filters: " + (", ".join(f"{k}={v}" for k, v in
                                                                 result.filters.to_payload().items()) or "none"))
    lines.extend(result.notes)
    return lines


def render_csv(result: Result, user_label: str, at: datetime, *, mask: bool) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")

    def write(values: Iterable[Any]) -> None:
        writer.writerow([escape_csv_cell("" if v is None else str(v)) for v in values])

    for line in _header_lines(result, user_label, at):
        write([line])
    write([c.label for c in result.report.columns])
    for row in result.rows:
        write([_display(c, row.get(c.key), row, mask=mask) for c in result.report.columns])
    for currency, values in sorted(result.totals.items()):
        write([f"Total {currency}"] + [amount_str(values[c.key], currency) if c.key in values else ""
                                       for c in result.report.columns[1:]])
    if result.converted and result.converted.get("available"):
        write([result.converted.get("label") or ""])
    return out.getvalue().encode("utf-8")


def render_xlsx(result: Result, user_label: str, at: datetime, *, mask: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = re.sub(r"[^A-Za-z0-9 ]", "", result.report.name)[:31] or "Report"

    def cell(value: Any) -> Any:
        if isinstance(value, str):
            return escape_csv_cell(value)
        return value

    for line in _header_lines(result, user_label, at):
        sheet.append([cell(line)])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.append([cell(c.label) for c in result.report.columns])
    for col_index in range(1, len(result.report.columns) + 1):
        sheet.cell(row=sheet.max_row, column=col_index).font = Font(bold=True)
    for row in result.rows:
        values = []
        for col in result.report.columns:
            value = row.get(col.key)
            if col.type in ("money", "decimal", "int") and value is not None and not (col.sensitive and mask):
                values.append(Decimal(value) if col.type != "int" else int(value))
            else:
                values.append(cell(str(_display(col, value, row, mask=mask))))
        sheet.append(values)
    for currency, values in sorted(result.totals.items()):
        sheet.append([cell(f"Total {currency}")] + [values.get(c.key) for c in result.report.columns[1:]])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(result: Result, user_label: str, at: datetime, *, mask: bool) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    from app.utils import pdf_fonts

    pdf_fonts.register()
    small = ParagraphStyle("cell", fontName=pdf_fonts.SANS, fontSize=6.5, leading=8)
    bold = ParagraphStyle("head", fontName=pdf_fonts.SANS_BOLD, fontSize=6.5, leading=8)
    title = ParagraphStyle("title", fontName=pdf_fonts.SANS_BOLD, fontSize=12, leading=15)
    note = ParagraphStyle("note", fontName=pdf_fonts.SANS, fontSize=7.5, leading=9, textColor=colors.HexColor("#B45309"))
    for style in (small, bold, title, note):
        style.shaping = 1
    buffer = io.BytesIO()
    watermark = CONFIDENTIAL.format(user=user_label, time=at.strftime("%Y-%m-%d %H:%M UTC"))

    def decorate(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(pdf_fonts.SANS, 7)
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.drawString(24, 14, watermark + (f" · {FIXTURE_LABEL}" if result.fixture else "")
                          + f" · page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24,
                            bottomMargin=30, title=result.report.name, author=user_label)
    story: List[Any] = [Paragraph(pdf_fonts.markup(result.report.name, True), title)]
    for line in _header_lines(result, user_label, at):
        story.append(Paragraph(pdf_fonts.markup(line), note))
    story.append(Spacer(1, 6))
    columns = result.report.columns
    data = [[Paragraph(pdf_fonts.markup(c.label, True), bold) for c in columns]]
    for row in result.rows[:PDF_MAX_ROWS]:
        data.append([Paragraph(pdf_fonts.markup(str(_display(c, row.get(c.key), row, mask=mask))), small)
                     for c in columns])
    for currency, values in sorted(result.totals.items()):
        data.append([Paragraph(pdf_fonts.markup(f"Total {currency}", True), bold)]
                    + [Paragraph(pdf_fonts.markup(amount_str(values[c.key], currency) if c.key in values else ""),
                                 bold) for c in columns[1:]])
    table = LongTable(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9CA3AF")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(table)
    if len(result.rows) > PDF_MAX_ROWS:
        story.append(Paragraph(pdf_fonts.markup(f"Only the first {PDF_MAX_ROWS} rows are in the PDF; export CSV or "
                                                "XLSX for all rows."), note))
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buffer.getvalue()


RENDERERS = {"CSV": (render_csv, "csv", "text/csv"),
             "XLSX": (render_xlsx, "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
             "PDF": (render_pdf, "pdf", "application/pdf")}


def request_export(db: Session, principal: Principal, code: str, filters: Filters, fmt: str) -> Job:
    """Queue an export of a report (checked now, rebuilt in the job with the requester's live grants)."""
    if fmt not in FORMATS:
        raise bad_request(f"format must be one of {', '.join(FORMATS)}.", "INVALID_FORMAT")
    report = _report(principal, code, export=True)
    filters = filters.validate()
    job = jobs.enqueue(db, company_id=principal.company_id, type=EXPORT_JOB,
                       payload={"report_code": report.code, "format_code": fmt, "filters": filters.to_payload()},
                       created_by_user_id=principal.user_id, max_attempts=3)
    AuditLogService.log(db, principal, AUDIT_MODULE, "REPORT_EXPORT_REQUESTED",
                        f"Requested {fmt} export of report {report.name}", "report_export", job.id,
                        company_id=principal.company_id,
                        details={"report_code": report.code, "format": fmt, "filters": filters.to_payload()})
    return job


@jobs.register(EXPORT_JOB)
def _build_export(ctx: jobs.JobContext) -> dict:
    from app.models.user import User
    from app.payroll.services.wage_inputs import principal_for_user

    db = ctx.db
    principal = principal_for_user(db, ctx.company_id, ctx.created_by_user_id)
    if principal is None:
        raise jobs.PermanentJobError("The requester is no longer active.")
    payload = ctx.payload
    fmt = payload["format_code"]
    try:
        result = run_report(db, principal, payload["report_code"], Filters.from_payload(payload.get("filters") or {}),
                            export=True)
    except Exception as exc:  # permission revoked, report gone, filters invalid: never retried
        raise jobs.PermanentJobError(getattr(exc, "detail", None) or str(exc)) from exc
    user = db.get(User, principal.user_id)
    label = getattr(user, "email", None) or f"user {principal.user_id}"
    at = clock.utcnow()
    mask = not full_access(principal)
    renderer, ext, media = RENDERERS[fmt]
    data = renderer(result, label, at, mask=mask)
    path = storage.save_bytes(ctx.company_id, EXPORT_CATEGORY, data, ext)
    digest = hashlib.sha256(data).hexdigest()
    file_name = f"{result.report.code}_{at.strftime('%Y%m%d_%H%M%S')}.{ext}"
    AuditLogService.log(db, principal, AUDIT_MODULE, "REPORT_EXPORTED", f"Exported report {result.report.name} "
                        f"({fmt}, {len(result.rows)} rows)", "report_export", ctx.job_id,
                        company_id=ctx.company_id,
                        details={"report_code": result.report.code, "format": fmt, "rows": len(result.rows),
                                 "sha256": digest, "masked": mask})
    return {"report_code": result.report.code, "format": fmt, "file_name": file_name, "media_type": media,
            "file_ref": path, "sha256": digest, "bytes": len(data), "rows": len(result.rows), "masked": mask,
            "expires_at": (at + EXPORT_TTL).isoformat(), "fixture": result.fixture,
            "scope": _scope_snapshot(principal, result.report)}


# The permissions whose run-level scope decides which rows a report export holds (rows are filtered by
# payroll:reports:read, compensation rows by payroll:compensation:read, plus the report's own extras).
SCOPED_PERMISSIONS = (READ, COMPENSATION_READ)


def _scope_snapshot(principal: Principal, report: ReportDef) -> Dict[str, dict]:
    """The requester's run-level scope when the file was built, so a download can be re-checked against it."""
    snapshot: Dict[str, dict] = {}
    for permission in dict.fromkeys((*SCOPED_PERMISSIONS, *report.extra_permissions)):
        target = access.target_scope(principal, permission)
        snapshot[permission] = {"all": target.all_company, "legal_entity_ids": sorted(target.legal_entity_ids),
                                "pay_group_ids": sorted(target.pay_group_ids)}
    return snapshot


def _still_covers(principal: Principal, snapshot: Mapping[str, Any]) -> bool:
    """The caller's scope today still reaches everything the export was built from (it may have grown)."""
    for permission, was in snapshot.items():
        now = access.target_scope(principal, permission)
        if now.all_company:
            continue
        if was.get("all"):
            return False
        if not set(was.get("legal_entity_ids") or ()) <= now.legal_entity_ids:
            return False
        if not set(was.get("pay_group_ids") or ()) <= now.pay_group_ids:
            return False
    return True


def _recheck_export_access(principal: Principal, job: Job) -> None:
    """EC-DOC-06 for a READY report export, at every link and download: the report's permissions are still
    held (403 otherwise), and the scope the rows were drawn from, and the company-wide export right an
    unmasked file needed, are still covered (404 otherwise, as for any object outside the caller's scope).
    A file without a recorded scope is served to company-wide readers only."""
    report = BY_CODE.get((job.payload or {}).get("report_code") or "")
    if report is None:
        raise not_found("Export")
    if not _can_run(principal, report):
        raise forbidden()
    result = job.result or {}
    if result.get("masked") is False and not full_access(principal):
        raise not_found("Export")
    snapshot = result.get("scope")
    if snapshot is None:
        if not principal.has_company_scope(READ):
            raise not_found("Export")
    elif not isinstance(snapshot, Mapping) or not _still_covers(principal, snapshot):
        raise not_found("Export")


def can_download(principal: Principal, job: Job) -> bool:
    """Whether a READY export may still be downloaded by the caller (the download center's link)."""
    if not principal.has(EXPORT):
        return False
    try:
        _recheck_export_access(principal, job)
    except AppError:
        return False
    return True


def _export_status(job: Job, now: datetime) -> str:
    if job.status in (jobs.QUEUED, jobs.RUNNING):
        return "GENERATING" if job.status == jobs.RUNNING else "QUEUED"
    if job.status == jobs.SUCCEEDED:
        expires = (job.result or {}).get("expires_at")
        if expires and datetime.fromisoformat(expires) <= now:
            return "EXPIRED"
        return "READY"
    return "FAILED"


def export_json(job: Job, now: Optional[datetime] = None) -> dict:
    now = now or clock.utcnow()
    result = job.result or {}
    payload = job.payload or {}
    return {"id": job.id, "report_code": payload.get("report_code"), "format": payload.get("format_code"),
            "filters": payload.get("filters") or {}, "status": _export_status(job, now),
            "file_name": result.get("file_name"), "rows": result.get("rows"), "bytes": result.get("bytes"),
            "masked": result.get("masked"), "expires_at": result.get("expires_at"), "created_at": job.created_at,
            "error": job.error if job.status in (jobs.FAILED, jobs.DEAD) else None}


def _my_export(db: Session, principal: Principal, job_id: int) -> Job:
    job = (db.query(Job).filter(Job.id == job_id, Job.company_id == principal.company_id, Job.type == EXPORT_JOB,
                                Job.created_by_user_id == principal.user_id).first())
    if job is None:
        raise not_found("Export")
    return job


def _my_exports_query(db: Session, principal: Principal):
    if not principal.has(READ):
        raise forbidden()
    return db.query(Job).filter(Job.company_id == principal.company_id, Job.type == EXPORT_JOB,
                                Job.created_by_user_id == principal.user_id)


def list_exports(db: Session, principal: Principal, *, skip: int = 0, limit: int = 50) -> List[dict]:
    rows = _my_exports_query(db, principal).order_by(Job.id.desc()).offset(skip).limit(limit).all()
    now = clock.utcnow()
    return [export_json(j, now) for j in rows]


def count_exports(db: Session, principal: Principal) -> int:
    return _my_exports_query(db, principal).count()


def get_export(db: Session, principal: Principal, job_id: int) -> dict:
    if not principal.has(READ):
        raise forbidden()
    return export_json(_my_export(db, principal, job_id))


def _ready_export(db: Session, principal: Principal, job_id: int) -> Job:
    if not principal.has(EXPORT):
        raise forbidden()
    job = _my_export(db, principal, job_id)
    status = _export_status(job, clock.utcnow())
    if status == "EXPIRED":
        raise conflict("This export has expired; export the report again.", "EXPORT_EXPIRED")
    if status != "READY":
        raise conflict("This export is not ready.", "EXPORT_NOT_READY")
    _recheck_export_access(principal, job)
    return job


def export_link(db: Session, principal: Principal, job_id: int) -> dict:
    """A signed, user- and session-bound, five-minute link to the caller's own READY export (EC-DOC-06),
    redeemed through ``GET /payroll/downloads/{token}`` (which re-checks access and audits the download)."""
    from app.payroll.services import downloads

    job = _ready_export(db, principal, job_id)
    return downloads.download_link(principal, downloads.REPORT_KIND, job.id)


def export_file(db: Session, principal: Principal, job_id: int) -> Tuple[bytes, str, str]:
    """(bytes, file name, media type) of the caller's own READY export; audited."""
    job = _ready_export(db, principal, job_id)
    result = job.result or {}
    data = storage.resolve(principal.company_id, result.get("file_ref")).read_bytes()
    if hashlib.sha256(data).hexdigest() != result.get("sha256"):
        raise conflict("The stored export does not match its recorded hash.", "FILE_HASH_MISMATCH")
    AuditLogService.log(db, principal, AUDIT_MODULE, "REPORT_EXPORT_DOWNLOADED",
                        f"Downloaded export #{job.id} ({result.get('file_name')})", "report_export", job.id,
                        company_id=principal.company_id, details={"sha256": result.get("sha256")})
    return data, result.get("file_name") or f"report.{job.payload.get('format_code', 'csv').lower()}", \
        result.get("media_type") or "application/octet-stream"


# =========================================================================== overview KPIs (3.8 §5.1)

OVERVIEW_LISTS = ("eligible", "with_salary", "without_salary", "processed", "paid", "on_hold")


@dataclass(frozen=True)
class OverviewFilters:
    period_from: date
    period_to: date
    legal_entity_id: Optional[int] = None
    pay_group_id: Optional[int] = None
    currency: Optional[str] = None
    status: Optional[str] = None
    branch_id: Optional[int] = None
    convert_to: Optional[str] = None
    fx_date: Optional[date] = None


def _month(today: date) -> Tuple[date, date]:
    start = today.replace(day=1)
    end = (start.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    return start, end


def overview_filters(db: Session, principal: Principal, **values: Any) -> OverviewFilters:
    start, end = values.pop("period_from", None), values.pop("period_to", None)
    if start is None or end is None:
        default_start, default_end = _month(clock.company_today(db, principal.company_id))
        start, end = start or default_start, end or default_end
    if start > end:
        raise bad_request("The period range ends before it starts.", "INVALID_PERIOD_RANGE")
    for code in (values.get("currency"), values.get("convert_to")):
        if code is not None and currencies.find(code) is None:
            raise bad_request(f"Unknown currency {code}.", "UNKNOWN_CURRENCY")
    return OverviewFilters(period_from=start, period_to=end, **values)


def _overview_scope(principal: Principal, legal_entity_col, pay_group_col):
    if not _held(principal, OVERVIEW_PERMISSIONS):
        raise forbidden()
    return scope_any(principal, OVERVIEW_PERMISSIONS, legal_entity_col, pay_group_col)


def _overview_runs(db: Session, principal: Principal, f: OverviewFilters) -> List[Tuple[PayrollRun, PayPeriod, PayGroup]]:
    """The latest non-cancelled regular run of each in-scope pay group period overlapping the range."""
    query = (db.query(PayrollRun, PayPeriod, PayGroup)
             .join(PayPeriod, and_(PayPeriod.company_id == PayrollRun.company_id, PayPeriod.id == PayrollRun.pay_period_id))
             .join(PayGroup, and_(PayGroup.company_id == PayrollRun.company_id, PayGroup.id == PayrollRun.pay_group_id))
             .filter(PayrollRun.company_id == principal.company_id, PayrollRun.run_type == "REGULAR",
                     PayrollRun.status != "CANCELLED", PayPeriod.period_end >= f.period_from,
                     PayPeriod.period_start <= f.period_to,
                     _overview_scope(principal, PayrollRun.legal_entity_id, PayrollRun.pay_group_id)))
    if f.legal_entity_id is not None:
        query = query.filter(PayrollRun.legal_entity_id == f.legal_entity_id)
    if f.pay_group_id is not None:
        query = query.filter(PayrollRun.pay_group_id == f.pay_group_id)
    if f.currency is not None:
        query = query.filter(PayGroup.currency == f.currency)
    latest: Dict[int, Tuple[PayrollRun, PayPeriod, PayGroup]] = {}
    for run, period, group in query.order_by(PayrollRun.regular_sequence, PayrollRun.id):
        latest[period.id] = (run, period, group)
    return list(latest.values())


@dataclass
class OverviewData:
    """One scoped query family behind every overview figure (EC-UX-03).

    ``rows`` is the population (one row per eligible assignment with its KPI flags); ``runs`` the latest
    regular run of each in-scope pay group period; ``members`` the run employees of exactly those runs AND
    those population assignments. Counts, drill-downs, money and statutory totals all read these three,
    so a filter or a scope narrows every figure the same way."""

    rows: List[dict]
    runs: List[Tuple[PayrollRun, PayPeriod, PayGroup]]
    members: List[PayrollRunEmployee]


def _population_query(db: Session, principal: Principal, f: OverviewFilters):
    query = (db.query(EmployeeAssignment, Employee, PayGroup)
             .join(Employee, and_(Employee.company_id == EmployeeAssignment.company_id,
                                  Employee.id == EmployeeAssignment.employee_id))
             .join(PayGroup, and_(PayGroup.company_id == EmployeeAssignment.company_id,
                                  PayGroup.id == EmployeeAssignment.pay_group_id))
             .filter(EmployeeAssignment.company_id == principal.company_id, EmployeeAssignment.pay_group_id.isnot(None),
                     EmployeeAssignment.effective_from <= f.period_to,
                     or_(EmployeeAssignment.effective_to.is_(None), EmployeeAssignment.effective_to >= f.period_from),
                     Employee.hire_date <= f.period_to,
                     or_(Employee.termination_date.is_(None), Employee.termination_date >= f.period_from),
                     _overview_scope(principal, EmployeeAssignment.legal_entity_id, EmployeeAssignment.pay_group_id)))
    if f.legal_entity_id is not None:
        query = query.filter(EmployeeAssignment.legal_entity_id == f.legal_entity_id)
    if f.pay_group_id is not None:
        query = query.filter(EmployeeAssignment.pay_group_id == f.pay_group_id)
    if f.branch_id is not None:
        query = query.filter(EmployeeAssignment.branch_id == f.branch_id)
    if f.currency is not None:
        query = query.filter(PayGroup.currency == f.currency)
    return query.order_by(Employee.employee_code, EmployeeAssignment.id)


def overview_data(db: Session, principal: Principal, f: OverviewFilters) -> OverviewData:
    """THE overview query family (see ``OverviewData``)."""
    cid = principal.company_id
    population = _population_query(db, principal, f).all()
    runs = _overview_runs(db, principal, f)
    if not population:
        return OverviewData([], runs, [])
    assignment_ids = [a.id for a, _, _ in population]
    with_salary = {aid for (aid,) in db.query(CompensationAssignment.assignment_id).filter(
        CompensationAssignment.company_id == cid, CompensationAssignment.assignment_id.in_(assignment_ids),
        CompensationAssignment.status.in_(("ACTIVE", "SCHEDULED", "SUPERSEDED")),
        CompensationAssignment.effective_from <= f.period_to,
        or_(CompensationAssignment.effective_to.is_(None), CompensationAssignment.effective_to >= f.period_from))}
    run_ids = [run.id for run, _, _ in runs]
    final_runs = {run.id for run, _, _ in runs if run.status in FINAL_STATUSES}
    members: List[PayrollRunEmployee] = []
    if run_ids:
        members = (db.query(PayrollRunEmployee)
                   .filter(PayrollRunEmployee.company_id == cid, PayrollRunEmployee.run_id.in_(run_ids),
                           PayrollRunEmployee.employee_assignment_id.in_(assignment_ids))
                   .order_by(PayrollRunEmployee.id).all())
    by_assignment: Dict[int, List[PayrollRunEmployee]] = defaultdict(list)
    for rem in members:
        by_assignment[rem.employee_assignment_id].append(rem)
    paid_members, held_by_instruction = _payment_flags(db, cid, [r.id for r in members])
    rows = []
    for assignment, employee, group in population:
        rems = by_assignment.get(assignment.id, [])
        processed = [r for r in rems if r.run_id in final_runs and r.status == "CALCULATED"]
        rows.append({
            "assignment_id": assignment.id, "employee_id": employee.id, "employee_code": employee.employee_code,
            "employee_name": _employee_name(employee.first_name, employee.last_name),
            "legal_entity_id": assignment.legal_entity_id, "pay_group_id": assignment.pay_group_id,
            "pay_group": group.code, "branch_id": assignment.branch_id, "currency": group.currency,
            "with_salary": assignment.id in with_salary,
            "processed": bool(processed),
            "paid": bool(processed) and all(r.id in paid_members for r in processed),
            "on_hold": any(r.hold_type is not None or r.id in held_by_instruction for r in rems),
            "run_ids": sorted({r.run_id for r in rems}),
        })
    return OverviewData(rows, runs, members)


def overview_population(db: Session, principal: Principal, f: OverviewFilters) -> List[dict]:
    """One row per eligible assignment with its KPI flags (the population of ``overview_data``)."""
    return overview_data(db, principal, f).rows


def _payment_flags(db: Session, company_id: int, run_employee_ids: Sequence[int]) -> Tuple[Set[int], Set[int]]:
    """(paid run employees, run employees with an ON_HOLD instruction). Paid needs every live instruction
    SETTLED or RECONCILED with VERIFIED evidence (contract §23.2) — never a submitted file alone."""
    pf = _payments()
    if pf is None or not run_employee_ids:
        return set(), set()
    Obl, Instr = pf.PaymentObligation, pf.PaymentInstruction
    rows = (db.query(Obl.run_employee_id, Instr.status, Instr.evidence_status)
            .join(Instr, and_(Instr.company_id == Obl.company_id, Instr.obligation_id == Obl.id))
            .filter(Obl.company_id == company_id, Obl.run_employee_id.in_(list(run_employee_ids))).all())
    state: Dict[int, List[Tuple[str, str]]] = defaultdict(list)
    for rem_id, status, evidence in rows:
        state[rem_id].append((status, evidence))
    paid, held = set(), set()
    for rem_id, items in state.items():
        live = [(s, e) for s, e in items if s not in ("CANCELLED", "REISSUED")]
        if any(s == "ON_HOLD" for s, _ in live):
            held.add(rem_id)
        if live and all(s in ("SETTLED", "RECONCILED") and e == "VERIFIED" for s, e in live):
            paid.add(rem_id)
    return paid, held


def _kpi_filter(name: str) -> Callable[[dict], bool]:
    return {
        "eligible": lambda r: True,
        "with_salary": lambda r: r["with_salary"],
        "without_salary": lambda r: not r["with_salary"],
        "processed": lambda r: r["processed"],
        "paid": lambda r: r["paid"],
        "on_hold": lambda r: r["on_hold"],
    }[name]


def overview_list(db: Session, principal: Principal, kpi: str, f: OverviewFilters, *, skip: int = 0,
                  limit: int = 100) -> dict:
    if kpi not in OVERVIEW_LISTS:
        raise not_found("KPI")
    rows = [r for r in overview_data(db, principal, f).rows if _kpi_filter(kpi)(r)]
    return {"kpi": kpi, "total": len(rows), "items": rows[skip:skip + limit]}


def _selected_runs(data: OverviewData, f: OverviewFilters) -> Dict[int, PayrollRun]:
    statuses = _run_statuses(f.status) if f.status else None
    return {run.id: run for run, _, _ in data.runs if statuses is None or run.status in statuses}


def _money_rows(data: OverviewData, f: OverviewFilters) -> List[dict]:
    """Gross / net / employer cost per currency and run stage of the population's calculated run employees."""
    selected = _selected_runs(data, f)
    totals: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for rem in data.members:
        run = selected.get(rem.run_id)
        if run is None or rem.status != "CALCULATED":
            continue
        bucket = totals.setdefault((rem.currency, run_stage(run.status)), {
            "gross": _ZERO, "net": _ZERO, "employer_cost": _ZERO, "employees": 0, "runs": set()})
        bucket["gross"] += _d(rem.gross)
        bucket["net"] += _d(rem.net)
        bucket["employer_cost"] += _d(rem.employer_cost)
        bucket["employees"] += 1
        bucket["runs"].add(rem.run_id)
    return [{"currency": cur, "stage": stage, "gross": amount_str(b["gross"], cur), "net": amount_str(b["net"], cur),
             "employer_cost": amount_str(b["employer_cost"], cur), "employees": b["employees"],
             "runs": len(b["runs"])}
            for (cur, stage), b in sorted(totals.items())]


def _statutory(db: Session, principal: Principal, f: OverviewFilters, data: OverviewData) -> List[dict]:
    """Statutory totals per entity, only where the entity's pack defines the components (EC-UX-07)."""
    from app.payroll.services import statutory_reports

    selected = _selected_runs(data, f)
    by_entity: Dict[int, List[PayrollRunEmployee]] = defaultdict(list)
    for rem in data.members:
        run = selected.get(rem.run_id)
        if run is not None and rem.status == "CALCULATED":
            by_entity[run.legal_entity_id].append(rem)
    result = []
    for entity_id, members in sorted(by_entity.items()):
        entity = db.query(LegalEntity).filter(LegalEntity.company_id == principal.company_id,
                                              LegalEntity.id == entity_id).first()
        if entity is None:
            continue
        found = statutory_reports.entity_pack(entity.country_code, f.period_from, f.period_to) or \
            statutory_reports.entity_pack(entity.country_code, f.period_to, f.period_to)
        if found is None or not found.usage.final:
            continue
        labels = {c.code: c.name for c in found.pack.components()}
        if not labels:
            continue
        stage_of = {rem.run_id: run_stage(selected[rem.run_id].status) for rem in members}
        sums: Dict[Tuple[str, str, str], Decimal] = defaultdict(lambda: _ZERO)
        for code, currency, run_id, amount in (
                db.query(PayrollLine.component_code, PayrollLine.currency, PayrollLine.run_id,
                         func.sum(PayrollLine.amount))
                .filter(PayrollLine.company_id == principal.company_id,
                        PayrollLine.run_employee_id.in_([rem.id for rem in members]),
                        PayrollLine.component_code.in_(list(labels)))
                .group_by(PayrollLine.component_code, PayrollLine.currency, PayrollLine.run_id)):
            sums[(code, currency, stage_of[run_id])] += _d(amount)
        if not sums:
            continue
        result.append({
            "legal_entity_id": entity_id, "pack_version": found.pack.meta.version, "watermark": found.watermark,
            "items": [{"code": code, "label": labels[code], "amount": amount_str(value, currency),
                       "currency": currency, "stage": stage}
                      for (code, currency, stage), value in sorted(sums.items())],
        })
    return result


def overview_kpis(db: Session, principal: Principal, f: OverviewFilters) -> dict:
    data = overview_data(db, principal, f)
    population, runs = data.rows, data.runs
    counts = {name: sum(1 for r in population if _kpi_filter(name)(r)) for name in OVERVIEW_LISTS}
    money = _money_rows(data, f)
    converted = None
    if f.convert_to and money:
        totals: Dict[str, Dict[str, Decimal]] = defaultdict(dict)
        for row in money:
            for key in ("gross", "net", "employer_cost"):
                totals[row["currency"]][key] = totals[row["currency"]].get(key, _ZERO) + Decimal(row[key])
        overview_def = ReportDef("overview", "Payroll", "Overview", "", (
            Col("gross", "Gross", "money"), Col("net", "Net", "money"), Col("employer_cost", "Employer cost", "money")),
            lambda ctx: [])
        ctx = Ctx(db=db, principal=principal, filters=Filters(convert_to=f.convert_to, fx_date=f.fx_date,
                                                              period_to=f.period_to),
                  today=clock.company_today(db, principal.company_id), report=overview_def)
        converted = _converted(ctx, overview_def, totals)
    run_ids = [run.id for run, _, _ in runs]
    exceptions: Dict[str, Any] = {"open": 0, "blocking": 0, "by_code": {}}
    if run_ids:
        for code, severity, count in (db.query(PayrollException.code, PayrollException.severity, func.count())
                                      .filter(PayrollException.company_id == principal.company_id,
                                              PayrollException.run_id.in_(run_ids), PayrollException.status == "OPEN")
                                      .group_by(PayrollException.code, PayrollException.severity)):
            exceptions["open"] += count
            exceptions["blocking"] += count if severity == "BLOCKING" else 0
            exceptions["by_code"][code] = exceptions["by_code"].get(code, 0) + count
    deadlines = [
        {"pay_group_id": group.id, "pay_group": group.code, "period_start": period.period_start,
         "period_end": period.period_end, "input_cutoff": period.input_cutoff_on,
         "approval_deadline": period.approval_deadline_on, "pay_date": period.pay_date,
         "run": {"id": run.id, "number": run.number, "status": run.status, "stage": run_stage(run.status)}}
        for run, period, group in sorted(runs, key=lambda r: (r[1].pay_date, r[2].code))
    ]
    deadlines.extend(_periods_without_runs(db, principal, f, {period.id for _, period, _ in runs}))
    if not population:
        empty_state = "NO_EMPLOYEES"
    elif counts["with_salary"] == 0:
        empty_state = "NO_SALARIES"
    else:
        empty_state = None
    return {
        "filters": {k: (v.isoformat() if isinstance(v, date) else v) for k, v in asdict(f).items() if v is not None},
        "kpis": counts,
        "consistency": {"eligible_equals_with_plus_without":
                        counts["eligible"] == counts["with_salary"] + counts["without_salary"]},
        "money": money,
        "converted": converted,
        "statutory": _statutory(db, principal, f, data),
        "exceptions": exceptions,
        "deadlines": sorted(deadlines, key=lambda d: (d["pay_date"], d["pay_group"])),
        "empty_state": empty_state,
        "processing_disabled_reason": ("Assign salaries before processing payroll."
                                       if empty_state == "NO_SALARIES" else
                                       "Add employees first." if empty_state == "NO_EMPLOYEES" else None),
        "definitions": {
            "eligible": "Employed at least one day of the period, counted per assignment in the selected pay groups.",
            "processed": "Included in a regular run that is finalized or later.",
            "paid": "Every payment instruction settled or reconciled with verified evidence.",
            "on_hold": "A payment or full hold, or a held payment instruction.",
            "money": "Calculated employees of the selected population, per currency and run stage.",
        },
    }


def _periods_without_runs(db: Session, principal: Principal, f: OverviewFilters, seen: Set[int]) -> List[dict]:
    query = (db.query(PayPeriod, PayGroup)
             .join(PayGroup, and_(PayGroup.company_id == PayPeriod.company_id, PayGroup.id == PayPeriod.pay_group_id))
             .filter(PayPeriod.company_id == principal.company_id, PayPeriod.period_end >= f.period_from,
                     PayPeriod.period_start <= f.period_to, PayGroup.status == "ACTIVE",
                     _overview_scope(principal, PayGroup.legal_entity_id, PayGroup.id)))
    if f.legal_entity_id is not None:
        query = query.filter(PayGroup.legal_entity_id == f.legal_entity_id)
    if f.pay_group_id is not None:
        query = query.filter(PayGroup.id == f.pay_group_id)
    if f.currency is not None:
        query = query.filter(PayGroup.currency == f.currency)
    return [{"pay_group_id": group.id, "pay_group": group.code, "period_start": period.period_start,
             "period_end": period.period_end, "input_cutoff": period.input_cutoff_on,
             "approval_deadline": period.approval_deadline_on, "pay_date": period.pay_date, "run": None}
            for period, group in query if period.id not in seen]


__all__ = [
    "BY_CODE", "CATALOGUE", "Filters", "OverviewData", "OverviewFilters", "catalogue", "count_exports", "export_file",
    "export_link", "get_export", "list_exports", "overview_data", "overview_filters", "overview_kpis", "overview_list",
    "overview_population",
    "preview", "request_export", "run_report", "run_stage", "suppress",
]
