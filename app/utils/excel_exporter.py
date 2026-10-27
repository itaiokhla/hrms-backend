from io import BytesIO
from typing import Any, Dict, List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.services.attendance_hours import decimal_hours


class ExcelExporter:
    HEADER_FILL = PatternFill(start_color="1e3a8a", end_color="1e3a8a", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="ffffff", size=11)
    BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # A text cell starting with one of these is read as a formula (or a command) by spreadsheet programs.
    FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

    @staticmethod
    def _neutralize_formulas(wb: Workbook) -> None:
        """Make every text cell inert before saving (LH BUG-26, LH-RPT-29): prefix a quote, store as text."""
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    # A lone "-" or "+" (the "no value" placeholder) cannot run anything and stays as it is.
                    if (isinstance(cell.value, str) and len(cell.value) > 1
                            and cell.value.startswith(ExcelExporter.FORMULA_PREFIXES)):
                        cell.value = "'" + cell.value
                        cell.data_type = "s"

    @staticmethod
    def _save(wb: Workbook) -> BytesIO:
        ExcelExporter._neutralize_formulas(wb)
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def _apply_header(cell):
        cell.font = ExcelExporter.HEADER_FONT
        cell.fill = ExcelExporter.HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = ExcelExporter.BORDER

    @staticmethod
    def _finalize_sheet(ws, column_widths: Dict[int, int]):
        for column, width in column_widths.items():
            ws.column_dimensions[get_column_letter(column)].width = width

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                if cell.row == 1:
                    continue
                cell.border = ExcelExporter.BORDER

    @staticmethod
    def _create_summary_sheet(wb: Workbook, title: str, summary: Dict[str, Any]) -> None:
        ws = wb.create_sheet(title=title)
        ws.append(["Metric", "Value"])
        ExcelExporter._apply_header(ws.cell(row=1, column=1))
        ExcelExporter._apply_header(ws.cell(row=1, column=2))

        for key, value in summary.items():
            ws.append([str(key).replace("_", " ").title(), value])

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=2):
            row[0].alignment = Alignment(horizontal="left")
            row[1].alignment = Alignment(horizontal="right")
            row[0].border = ExcelExporter.BORDER
            row[1].border = ExcelExporter.BORDER

        ExcelExporter._finalize_sheet(ws, {1: 24, 2: 16})

    @staticmethod
    def export_payroll_report(payrolls_data: List[Dict[str, Any]]) -> BytesIO:
        """Export payroll data to Excel with summary and details sheets.

        Totals are per currency: amounts in different currencies are never added together.
        """
        wb = Workbook()
        details = wb.active
        details.title = "Payroll Details"

        headers = [
            "Employee ID",
            "Employee Name",
            "Salary Month",
            "Currency",
            "Basic Salary",
            "HRA",
            "DA",
            "Conveyance",
            "Gross Salary",
            "Total Deductions",
            "Net Salary",
            "Status",
        ]

        for col, header in enumerate(headers, 1):
            cell = details.cell(row=1, column=col)
            cell.value = header
            ExcelExporter._apply_header(cell)

        totals: Dict[str, Dict[str, float]] = {}

        for row_idx, payroll in enumerate(payrolls_data, 2):
            currency = payroll.get("currency") or "-"
            details.cell(row=row_idx, column=1).value = payroll.get("employee_id")
            details.cell(row=row_idx, column=2).value = payroll.get("employee_name")
            details.cell(row=row_idx, column=3).value = payroll.get("salary_month")
            details.cell(row=row_idx, column=4).value = currency
            details.cell(row=row_idx, column=5).value = float(payroll.get("basic_salary", 0))
            details.cell(row=row_idx, column=6).value = float(payroll.get("hra", 0))
            details.cell(row=row_idx, column=7).value = float(payroll.get("da", 0))
            details.cell(row=row_idx, column=8).value = float(payroll.get("conveyance", 0))
            details.cell(row=row_idx, column=9).value = float(payroll.get("gross_salary", 0))
            details.cell(row=row_idx, column=10).value = float(payroll.get("total_deductions", 0))
            details.cell(row=row_idx, column=11).value = float(payroll.get("net_salary", 0))
            details.cell(row=row_idx, column=12).value = payroll.get("status")

            bucket = totals.setdefault(currency, {"count": 0, "gross": 0.0, "deductions": 0.0, "net": 0.0})
            bucket["count"] += 1
            bucket["gross"] += float(payroll.get("gross_salary", 0))
            bucket["deductions"] += float(payroll.get("total_deductions", 0))
            bucket["net"] += float(payroll.get("net_salary", 0))

            for col in range(1, 13):
                cell = details.cell(row=row_idx, column=col)
                if 5 <= col <= 11:
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left")
                cell.border = ExcelExporter.BORDER

        summary: Dict[str, Any] = {"Total Records": len(payrolls_data)}
        for currency, bucket in sorted(totals.items()):
            summary[f"Total Gross ({currency})"] = bucket["gross"]
            summary[f"Total Deductions ({currency})"] = bucket["deductions"]
            summary[f"Total Net ({currency})"] = bucket["net"]
            summary[f"Average Salary ({currency})"] = bucket["net"] / bucket["count"] if bucket["count"] else 0

        ExcelExporter._create_summary_sheet(wb, "Payroll Summary", summary)
        ExcelExporter._finalize_sheet(details, {
            1: 12,
            2: 20,
            3: 16,
            4: 10,
            5: 14,
            6: 10,
            7: 10,
            8: 12,
            9: 14,
            10: 16,
            11: 14,
            12: 12,
        })

        return ExcelExporter._save(wb)

    @staticmethod
    def export_attendance_report(attendance_data: List[Dict[str, Any]]) -> BytesIO:
        """Export attendance data to Excel with summary and details sheets."""
        wb = Workbook()
        details = wb.active
        details.title = "Attendance Details"

        headers = [
            "Employee ID",
            "Employee Name",
            "Date",
            "Check-in Time",
            "Check-out Time",
            "Status",
            "Face Match Score",
            "Gross Hours",
            "Lunch Hours",
            "Net Hours",
        ]

        for col, header in enumerate(headers, 1):
            cell = details.cell(row=1, column=col)
            cell.value = header
            ExcelExporter._apply_header(cell)

        summary = {
            "Total Records": len(attendance_data),
            "Total Present": 0,
            "Total Absent": 0,
            "Total Late": 0,
            "Total Half Day": 0,
        }
        totals = {"gross_seconds": 0, "lunch_seconds": 0, "net_seconds": 0}
        incomplete = 0

        for row_idx, record in enumerate(attendance_data, 2):
            details.cell(row=row_idx, column=1).value = record.get("employee_id")
            details.cell(row=row_idx, column=2).value = record.get("employee_name")
            details.cell(row=row_idx, column=3).value = record.get("date")
            details.cell(row=row_idx, column=4).value = record.get("check_in_time")
            details.cell(row=row_idx, column=5).value = record.get("check_out_time")
            status_value = record.get("status")
            details.cell(row=row_idx, column=6).value = status_value
            # Manual and button check-ins have no face match, which is not the same as a 0% match.
            score = record.get("face_match_score")
            details.cell(row=row_idx, column=7).value = f"{float(score) * 100:.1f}%" if score is not None else "-"
            # Hours from the canonical calculator; a missing punch-out is left blank, never estimated.
            for column, key in ((8, "gross_seconds"), (9, "lunch_seconds"), (10, "net_seconds")):
                details.cell(row=row_idx, column=column).value = decimal_hours(record.get(key))
                details.cell(row=row_idx, column=column).number_format = "0.00"
            if record.get("hours_complete", True):
                for key in totals:
                    totals[key] += record.get(key) or 0
            else:
                incomplete += 1

            if status_value == "present":
                summary["Total Present"] += 1
            elif status_value == "absent":
                summary["Total Absent"] += 1
            elif status_value == "late":
                summary["Total Late"] += 1
            elif status_value == "half-day":
                summary["Total Half Day"] += 1

            for col in range(1, 11):
                cell = details.cell(row=row_idx, column=col)
                if col >= 7:
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left")
                cell.border = ExcelExporter.BORDER

        # Totals are summed in seconds and converted once, so they match the API to the second.
        summary["Total Gross Hours"] = decimal_hours(totals["gross_seconds"])
        summary["Total Lunch Hours"] = decimal_hours(totals["lunch_seconds"])
        summary["Total Net Hours"] = decimal_hours(totals["net_seconds"])
        summary["Total Net Seconds"] = totals["net_seconds"]
        summary["Sessions Missing Punch-out"] = incomplete
        ExcelExporter._create_summary_sheet(wb, "Attendance Summary", summary)
        ExcelExporter._finalize_sheet(details, {
            1: 12,
            2: 20,
            3: 14,
            4: 18,
            5: 18,
            6: 12,
            7: 16,
            8: 12,
            9: 12,
            10: 12,
        })

        return ExcelExporter._save(wb)

    @staticmethod
    def export_leave_report(leaves_data: List[Dict[str, Any]]) -> BytesIO:
        """Export leave data to Excel with summary and details sheets."""
        wb = Workbook()
        details = wb.active
        details.title = "Leave Details"

        headers = [
            "Employee ID",
            "Employee Name",
            "Leave Type",
            "Start Date",
            "End Date",
            "Number of Days",
            "Reason",
            "Status",
        ]

        for col, header in enumerate(headers, 1):
            cell = details.cell(row=1, column=col)
            cell.value = header
            ExcelExporter._apply_header(cell)

        # Leave types are configured per company, so count whichever types occur.
        summary = {"Total Records": len(leaves_data)}

        for row_idx, leave in enumerate(leaves_data, 2):
            leave_type = leave.get("leave_type_name") or leave.get("leave_type") or "Unknown"
            summary[leave_type] = summary.get(leave_type, 0) + 1
            details.cell(row=row_idx, column=1).value = leave.get("employee_id")
            details.cell(row=row_idx, column=2).value = leave.get("employee_name")
            details.cell(row=row_idx, column=3).value = leave_type
            details.cell(row=row_idx, column=4).value = leave.get("start_date")
            details.cell(row=row_idx, column=5).value = leave.get("end_date")
            details.cell(row=row_idx, column=6).value = leave.get("days")
            details.cell(row=row_idx, column=7).value = leave.get("reason")
            details.cell(row=row_idx, column=8).value = leave.get("status")

            for col in range(1, 9):
                cell = details.cell(row=row_idx, column=col)
                if col == 6:
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")
                cell.border = ExcelExporter.BORDER

        ExcelExporter._create_summary_sheet(wb, "Leave Summary", summary)
        ExcelExporter._finalize_sheet(details, {
            1: 12,
            2: 20,
            3: 16,
            4: 14,
            5: 14,
            6: 14,
            7: 30,
            8: 14,
        })

        return ExcelExporter._save(wb)
