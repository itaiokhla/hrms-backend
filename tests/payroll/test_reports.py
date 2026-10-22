"""Payroll reports: catalogue, previews, scoped filters, per-currency and converted totals, small-group
suppression and exports (3.8 §5.13, 3.9 §5.4, contract §9).

EC-SEC-05 (scoped counts and totals; small groups hidden), EC-DOC-17 (CSV / XLSX cells escaped),
EC-MON-04 (never summed across currencies), EC-MON-05 (never converted 1:1), EC-SEC-01/02 (three tenants,
foreign ids are 404).

The world built here (``build_world`` / ``rc_world``) is shared with test_overview_kpis.py,
test_statutory_reports.py, test_compliance.py and test_accounting.py: runs are seeded directly (finalized
runs with lines, liabilities and run employees) so these tests do not depend on the calculation engine.
"""

import csv
import io
import itertools
from datetime import date
from decimal import Decimal

import pytest

from tests.helpers import create_employee, invite_and_activate, signup_company
from tests.payroll.test_wage_inputs import pay_group, place, session, set_setting, sql

PAY = "/api/v1/payroll"
REPORTS = f"{PAY}/reports"
EXPORTS = f"{PAY}/report-exports"
AUG = date(2026, 8, 1)
SEP = date(2026, 9, 1)
OCT = date(2026, 10, 1)
SEP_FILTER = {"period_from": "2026-09-01", "period_to": "2026-09-30"}
_numbers = itertools.count(900001)


def ok(response, status=200):
    assert response.status_code == status, response.text
    return response.json()


def money(amount, currency):
    return {"amount": amount, "currency": currency}


# ---------------------------------------------------------------------------
# Shared world builders
# ---------------------------------------------------------------------------
def make_entity(client, code, name, country, currency, timezone):
    """A legal entity with one branch. ``ZZ`` (the fictional test pack) is set directly in the database."""
    api_country = "DE" if country == "ZZ" else country
    entity = ok(client.post("/api/v1/legal-entities", json={
        "code": code, "legal_name": name, "country_code": api_country, "base_currency": currency,
        "timezone": timezone}))["id"]
    branch = ok(client.post("/api/v1/branches", json={"legal_entity_id": entity, "code": f"B-{code}",
                                                      "name": f"{name} office"}))["id"]
    if country == "ZZ":
        sql("UPDATE legal_entities SET country_code = 'ZZ' WHERE id = :e", e=entity)
        sql("UPDATE branches SET country_code = 'ZZ' WHERE id = :b", b=branch)
    return entity, branch


def move(client, employee, branch_id, entity_id, group_id):
    r = client.post(f"/api/v1/employees/{employee['id']}/assignments",
                    json={"branch_id": branch_id, "effective_from": "2024-01-16", "reason": "TRANSFER"})
    assert r.status_code == 200, r.text
    place(employee["id"], entity_id, group_id)


def assignment_of(employee_id, entity_id):
    return sql("SELECT id FROM employee_assignments WHERE employee_id = :e AND legal_entity_id = :le "
               "ORDER BY id DESC LIMIT 1", e=employee_id, le=entity_id)[0][0]


def add_periods(company_id, group_id, months=(AUG, SEP, OCT)):
    """A monthly calendar and its periods (month start -> period id), written directly."""
    import calendar

    from app.models.pay_runs import PayCalendar, PayPeriod

    result = {}
    with session(company_id) as (db, _):
        cal = PayCalendar(company_id=company_id, pay_group_id=group_id, frequency="MONTHLY",
                          rules={"monthly_start_day": 1}, effective_from=date(2026, 1, 1), status="ACTIVE",
                          row_version=1)
        db.add(cal)
        db.flush()
        for start in months:
            end = start.replace(day=calendar.monthrange(start.year, start.month)[1])
            period = PayPeriod(
                company_id=company_id, pay_calendar_id=cal.id, pay_group_id=group_id, frequency="MONTHLY",
                period_start=start, period_end=end, attendance_start=start, attendance_end=end, pay_date=end,
                nominal_pay_date=end, input_cutoff_on=end.replace(day=25), approval_deadline_on=end.replace(day=28),
                payment_submit_by=end.replace(day=29), tax_year=str(start.year), number_in_tax_year=start.month,
                is_last_in_tax_year=start.month == 12, is_transition=False, status="OPEN", pay_date_warnings=[],
                row_version=1)
            db.add(period)
            db.flush()
            result[start] = period.id
        db.commit()
    return result


def line(code, kind="EARNING", amount="0", *, category=None, statutory=False, fixture=False, units=None,
         period_ref=None, sources=None, rate=None, version_id=None):
    return {"code": code, "kind": kind, "amount": amount, "category": category, "statutory": statutory,
            "fixture": fixture, "units": units, "period_ref": period_ref, "sources": sources or [], "rate": rate,
            "version_id": version_id}


def member(employee, gross="0", deductions="0", net="0", employer_cost="0", *, lines=(), status="CALCULATED",
           hold_type=None, exclusion_reason=None, reimbursements="0"):
    return {"employee_id": employee["id"], "gross": gross, "deductions": deductions, "net": net,
            "employer_cost": employer_cost, "reimbursements": reimbursements, "lines": list(lines),
            "status": status, "hold_type": hold_type, "exclusion_reason": exclusion_reason}


FINAL = ("FINALIZED", "PAYMENT_READY", "SUBMITTED", "PARTIALLY_PAID", "PAID", "RECONCILED", "CLOSED")


def seed_run(company_id, *, entity, group, period, currency, members, run_type="REGULAR", status="FINALIZED",
             reverses=None, liabilities=()):
    """A run with run employees, lines and liabilities written directly, then moved to ``status``."""
    from app.models.pay_runs import PayrollLine, PayrollRun, PayrollRunEmployee, StatutoryLiability

    with session(company_id) as (db, _):
        run = PayrollRun(
            company_id=company_id, number=f"PR-2026-{next(_numbers):06d}", label=f"{run_type} run",
            legal_entity_id=entity, pay_group_id=group, pay_period_id=period, run_type=run_type, status="DRAFT",
            regular_sequence=1 if run_type == "REGULAR" else None, currency=currency, reverses_run_id=reverses,
            selection={}, pack_version_ids=[], totals={}, population={}, stale=False, stale_reasons=[],
            environment="test", row_version=1)
        db.add(run)
        db.flush()
        ids = {}
        for m in members:
            rem = PayrollRunEmployee(
                company_id=company_id, run_id=run.id, employee_id=m["employee_id"],
                employee_assignment_id=assignment_of(m["employee_id"], entity), status=m["status"],
                exclusion_reason=m["exclusion_reason"], hold_type=m["hold_type"],
                hold={"reason_code": "OTHER", "note": "test hold"} if m["hold_type"] else None,
                gross=Decimal(m["gross"]), deductions=Decimal(m["deductions"]),
                reimbursements=Decimal(m["reimbursements"]), net=Decimal(m["net"]),
                employer_cost=Decimal(m["employer_cost"]), currency=currency, flags=[], row_version=1)
            db.add(rem)
            db.flush()
            ids[m["employee_id"]] = rem.id
            for order, item in enumerate(m["lines"], start=1):
                db.add(PayrollLine(
                    company_id=company_id, run_id=run.id, run_employee_id=rem.id, component_code=item["code"],
                    component_version_id=item["version_id"], kind=item["kind"], category=item["category"],
                    stage=4, amount=Decimal(item["amount"]), currency=currency,
                    units=Decimal(item["units"]) if item["units"] is not None else None,
                    rate=Decimal(item["rate"]) if item["rate"] is not None else None, period_ref=item["period_ref"],
                    tags=[], visibility="EMPLOYEE_VISIBLE", sources=item["sources"], trace={},
                    explanation=f"{item['code']} test line", sort_order=order * 10, statutory=item["statutory"],
                    fixture=item["fixture"]))
        for item in liabilities:
            db.add(StatutoryLiability(
                company_id=company_id, run_id=run.id, legal_entity_id=entity, scheme=item["scheme"],
                payee_authority=item["payee"], employee_amount=Decimal(item.get("employee", "0")),
                employer_amount=Decimal(item.get("employer", "0")), currency=currency, due_date=item.get("due"),
                status="OPEN", fixture=item.get("fixture", False), details={}, row_version=1))
        db.commit()
        run_id = run.id
    if status != "DRAFT":
        sql("UPDATE payroll_runs SET status = :s, finalized_at = CASE WHEN :f THEN now() ELSE NULL END "
            "WHERE id = :r", s=status, f=status in FINAL, r=run_id)
    return {"id": run_id, "members": ids}


def grant_role(app, admin, first, grants, entity_ids=None):
    """A signed-in user whose only role has ``grants`` [(permission, scope)], targeted at ``entity_ids``."""
    from app.core.rate_limit import limiter

    limiter.reset()
    r = admin.post("/api/v1/roles", json={"name": f"Role {first}",
                                          "grants": [{"permission": p, "scope": s} for p, s in grants]})
    assert r.status_code == 200, r.text
    role_id = r.json()["id"]
    client = invite_and_activate(app, admin, create_employee(admin, first), [])
    unplace(client.employee["id"])
    ok(admin.put(f"/api/v1/users/{client.user_id}/roles", json={"role_ids": [role_id]}))
    if entity_ids:
        ok(admin.put(f"/api/v1/users/{client.user_id}/scope-targets",
                     json={"assignments": [{"role_id": role_id, "legal_entity_ids": entity_ids}]}))
    client.company_id = admin.company_id
    return client


def with_roles(app, admin, first, roles):
    from app.core.rate_limit import limiter

    limiter.reset()
    client = invite_and_activate(app, admin, create_employee(admin, first), roles)
    unplace(client.employee["id"])
    client.company_id = admin.company_id
    return client


def unplace(employee_id):
    """Staff users of the tests are not paid through payroll: no pay group (new hires join the entity's
    default pay group otherwise and would count as eligible)."""
    sql("UPDATE employee_assignments SET pay_group_id = NULL WHERE employee_id = :e", e=employee_id)


def add_salaries(app, w):
    """Approved CTC-based salaries for Emma and Ivan from 1 Aug 2026 (needs the catalogue)."""
    from tests.payroll.test_compensation import assign_and_approve
    from tests.payroll.test_components import active_version_id, make_structure, setup_catalogue

    acme = w["acme"]
    setup_catalogue(acme)
    version = active_version_id(make_structure(acme))
    approver = with_roles(app, acme, "Ava", ["payroll_approver"])
    for name in ("Emma", "Ivan"):
        assign_and_approve(acme, approver, w["emp"][name]["id"], version, "1200000", effective_from=AUG)
    w["approver"] = approver
    return approver


def build_world(app, *, salaries=False):
    """Acme with three legal entities:

    * LE01 (India, the IN FIXTURE pack): PG-IN (INR) with Emma, Ivan and Nora; finalized regular runs for
      August and September (Nora excluded, NO_SALARY), a fixture IN.epf liability.
    * ZZ01 (Zedland, the verified ZZ test pack): PG-ZZ (EUR) with Zoe and Zack; a finalized September run
      with pension and withholding lines and liabilities.
    * AE01 (UAE, a DRAFT pack: no statutory output): PG-AE (AED) with Gus; an APPROVED (not final)
      September run carrying a line coded like an Indian PF component (EC-UX-07 trap).

    Users: the founder (payroll manager, COMPANY) and Cleo (reports read/export, run read and statutory
    review at ZZ01 only).
    """
    acme = signup_company(app, "Acme Corp")
    cid = acme.company_id
    le01 = sql("SELECT id FROM legal_entities WHERE company_id = :c AND code = 'LE01'", c=cid)[0][0]
    zz01, zz_branch = make_entity(acme, "ZZ01", "Zedland Ltd", "ZZ", "EUR", "Europe/Berlin")
    ae01, ae_branch = make_entity(acme, "AE01", "Acme Gulf FZ-LLC", "AE", "AED", "Asia/Dubai")
    groups = {"IN": pay_group(cid, le01, "PG-IN"), "ZZ": pay_group(cid, zz01, "PG-ZZ", currency="EUR"),
              "AE": pay_group(cid, ae01, "PG-AE", currency="AED")}
    emp = {}
    for name in ("Emma", "Ivan", "Nora"):
        emp[name] = create_employee(acme, name, hire_date="2024-01-15")
        place(emp[name]["id"], le01, groups["IN"])
    for name in ("Zoe", "Zack"):
        emp[name] = create_employee(acme, name, hire_date="2024-01-15")
        move(acme, emp[name], zz_branch, zz01, groups["ZZ"])
    emp["Gus"] = create_employee(acme, "Gus", hire_date="2024-01-15")
    move(acme, emp["Gus"], ae_branch, ae01, groups["AE"])
    periods = {key: add_periods(cid, group) for key, group in groups.items()}
    w = {"acme": acme, "cid": cid, "le01": le01, "zz01": zz01, "ae01": ae01, "zz_branch": zz_branch,
         "ae_branch": ae_branch, "groups": groups, "emp": emp, "periods": periods}
    if salaries:
        add_salaries(app, w)

    runs = {}
    runs["IN_AUG"] = seed_run(cid, entity=le01, group=groups["IN"], period=periods["IN"][AUG], currency="INR", members=[
        member(emp["Emma"], "50000.00", "5000.00", "45000.00", "52000.00", lines=[
            line("BASIC", amount="50000.00", category="BASE"),
            line("IN_EPF_EE", "DEDUCTION", "1800.00", statutory=True, fixture=True),
            line("OTHER_DEDUCTION", "DEDUCTION", "3200.00"),
            line("IN_EPF_ER", "EMPLOYER_CONTRIBUTION", "1800.00", statutory=True, fixture=True)]),
        member(emp["Ivan"], "40000.00", "4000.00", "36000.00", "41500.00", lines=[
            line("BASIC", amount="40000.00", category="BASE"),
            line("IN_EPF_EE", "DEDUCTION", "1800.00", statutory=True, fixture=True),
            line("OTHER_DEDUCTION", "DEDUCTION", "2200.00"),
            line("IN_EPF_ER", "EMPLOYER_CONTRIBUTION", "1500.00", statutory=True, fixture=True)]),
    ])
    runs["IN_SEP"] = seed_run(cid, entity=le01, group=groups["IN"], period=periods["IN"][SEP], currency="INR", members=[
        member(emp["Emma"], "52000.00", "5000.00", "47000.00", "54000.00", lines=[
            line("BASIC", amount="50000.00", category="BASE"),
            line("OT_WORKDAY", amount="2000.00", category="OVERTIME", units="4"),
            line("IN_EPF_EE", "DEDUCTION", "1800.00", statutory=True, fixture=True),
            line("OTHER_DEDUCTION", "DEDUCTION", "3200.00"),
            line("IN_EPF_ER", "EMPLOYER_CONTRIBUTION", "2000.00", statutory=True, fixture=True)]),
        member(emp["Ivan"], "40000.00", "4000.00", "36000.00", "41500.00", lines=[
            line("BASIC", amount="40000.00", category="BASE"),
            line("IN_EPF_EE", "DEDUCTION", "1800.00", statutory=True, fixture=True),
            line("OTHER_DEDUCTION", "DEDUCTION", "2200.00"),
            line("IN_EPF_ER", "EMPLOYER_CONTRIBUTION", "1500.00", statutory=True, fixture=True)]),
        member(emp["Nora"], status="EXCLUDED", exclusion_reason="NO_SALARY"),
    ], liabilities=[{"scheme": "IN.epf", "payee": "EPFO", "employee": "3600.00", "employer": "3500.00",
                     "due": date(2026, 10, 15), "fixture": True}])
    runs["ZZ_SEP"] = seed_run(cid, entity=zz01, group=groups["ZZ"], period=periods["ZZ"][SEP], currency="EUR", members=[
        member(emp["Zoe"], "4000.00", "500.00", "3500.00", "4280.00", lines=[
            line("BASIC", amount="4000.00", category="BASE"),
            line("ZZ_PENSION_EE", "DEDUCTION", "200.00", statutory=True),
            line("ZZ_WHT", "DEDUCTION", "300.00", statutory=True),
            line("ZZ_PENSION_ER", "EMPLOYER_CONTRIBUTION", "280.00", statutory=True)]),
        member(emp["Zack"], "3000.00", "350.00", "2650.00", "3210.00", lines=[
            line("BASIC", amount="3000.00", category="BASE"),
            line("ZZ_PENSION_EE", "DEDUCTION", "150.00", statutory=True),
            line("ZZ_WHT", "DEDUCTION", "200.00", statutory=True),
            line("ZZ_PENSION_ER", "EMPLOYER_CONTRIBUTION", "210.00", statutory=True)]),
    ], liabilities=[
        {"scheme": "ZZ.pension", "payee": "Zedland Pension Fund (fictional)", "employee": "350.00",
         "employer": "490.00", "due": date(2026, 10, 15)},
        {"scheme": "ZZ.wht", "payee": "Zedland Revenue (fictional)", "employee": "500.00",
         "due": date(2026, 10, 10)},
    ])
    runs["AE_SEP"] = seed_run(cid, entity=ae01, group=groups["AE"], period=periods["AE"][SEP], currency="AED",
                              status="APPROVED", members=[
        member(emp["Gus"], "10000.00", "100.00", "9900.00", "10000.00", lines=[
            line("BASIC", amount="10000.00", category="BASE"),
            line("IN_EPF_EE", "DEDUCTION", "100.00")]),
    ])
    w["runs"] = runs
    w["cleo"] = grant_role(app, acme, "Cleo", [(p, "LEGAL_ENTITY") for p in (
        "payroll:reports:read", "payroll:reports:export", "payroll:run:read", "payroll:statutory:review")], [zz01])
    return w


@pytest.fixture
def rc_world(app):
    return build_world(app)


def other_tenant(app, name, *, currency="INR"):
    """A second tenant with one finalized run (isolation tests)."""
    client = signup_company(app, name)
    cid = client.company_id
    entity = sql("SELECT id FROM legal_entities WHERE company_id = :c AND code = 'LE01'", c=cid)[0][0]
    group = pay_group(cid, entity, "PG-X", currency=currency)
    employee = create_employee(client, f"{name.split()[0]}Emp", hire_date="2024-01-15")
    place(employee["id"], entity, group)
    periods = add_periods(cid, group, months=(SEP,))
    run = seed_run(cid, entity=entity, group=group, period=periods[SEP], currency=currency, members=[
        member(employee, "7777.00", "777.00", "7000.00", "8000.00", lines=[line("BASIC", amount="7777.00")])],
        liabilities=[{"scheme": "IN.epf", "payee": "EPFO", "employee": "77.00", "employer": "77.00",
                      "due": date(2026, 10, 15)}])
    return {"client": client, "cid": cid, "entity": entity, "group": group, "employee": employee, "run": run,
            "period": periods[SEP]}


def preview(client, code, expect=200, **params):
    return ok(client.get(f"{REPORTS}/{code}/preview", params=params), expect)


def export_ready(client, code, fmt="CSV", filters=None):
    created = ok(client.post(f"{REPORTS}/{code}/exports", json={"format": fmt, "filters": filters or {}}), 202)
    status = ok(client.get(f"{EXPORTS}/{created['id']}"))
    assert status["status"] == "READY", status
    r = client.get(f"{EXPORTS}/{created['id']}/download")
    assert r.status_code == 200, r.text
    return created["id"], r


def csv_rows(content: bytes):
    return list(csv.reader(io.StringIO(content.decode("utf-8"))))


# ---------------------------------------------------------------------------
# Catalogue and permissions
# ---------------------------------------------------------------------------
def test_catalogue_groups_every_report_and_hides_reports_needing_more_permissions(app, rc_world):
    from app.payroll.services import reports

    w = rc_world
    catalogue = ok(w["acme"].get(REPORTS))
    codes = {r["code"] for g in catalogue["groups"] for r in g["reports"]}
    expected = {
        "payroll_register", "gross_to_net", "variance", "ytd_by_employee", "corrected_view", "employer_cost",
        "ctc_vs_actual", "overtime_by_bucket", "lop_and_absence", "holiday_premiums", "late_penalties",
        "bonus_incentives", "task_pay", "arrears", "payment_register", "failures_returns",
        "reconciliation_statement", "outstanding_obligations", "exclusions", "loans_outstanding", "reimbursements",
        "leave_encashments", "liabilities_remittances", "statutory_filings", "compensation_changes"}
    audit_reports = {"settings_policy_changes", "bank_detail_events"}
    # The founder (company administrator, which implies the non-confidential payroll:audit:read, plus the
    # payroll manager role) sees the whole catalogue.
    assert codes == expected | audit_reports == {r.code for r in reports.CATALOGUE}
    # A payroll processor holds no payroll:audit:read: the audit-trail reports are not offered.
    processor = with_roles(app, w["acme"], "Pia", ["payroll_processor"])
    processor_codes = {r["code"] for g in ok(processor.get(REPORTS))["groups"] for r in g["reports"]}
    assert processor_codes == expected
    assert processor.get(f"{REPORTS}/settings_policy_changes/preview").status_code == 403
    assert catalogue["preview_rows"] == 50
    assert [g["name"] for g in catalogue["groups"]][:2] == ["Payroll", "Cost"]
    register = next(r for g in catalogue["groups"] for r in g["reports"] if r["code"] == "payroll_register")
    assert register["exportable"] is True and register["formats"] == ["CSV", "XLSX", "PDF"]

    auditor = with_roles(app, w["acme"], "Audrey", ["auditor"])
    audit_codes = {r["code"] for g in ok(auditor.get(REPORTS))["groups"] for r in g["reports"]}
    assert {"settings_policy_changes", "bank_detail_events"} <= audit_codes
    assert all(not r["exportable"] for g in ok(auditor.get(REPORTS))["groups"] for r in g["reports"])
    assert "compensation_changes" not in audit_codes  # needs payroll:compensation:read
    # Function-level denials are 403, unknown reports 404.
    assert auditor.post(f"{REPORTS}/payroll_register/exports", json={"format": "CSV"}).status_code == 403
    assert w["acme"].get(f"{REPORTS}/no_such_report/preview").status_code == 404
    employee = with_roles(app, w["acme"], "Eve", ["employee"])
    assert employee.get(REPORTS).status_code == 403
    assert employee.get(f"{REPORTS}/payroll_register/preview").status_code == 403


def test_payroll_register_preview_totals_per_currency_and_run_stage_filters(rc_world):
    w = rc_world
    data = preview(w["acme"], "payroll_register", **SEP_FILTER)
    assert data["row_count"] == 4 and data["truncated"] is False
    by_code = {r["employee_code"]: r for r in data["rows"]}
    assert {by_code[w["emp"][n]["employee_code"]]["currency"] for n in ("Emma", "Ivan")} == {"INR"}
    emma = by_code[w["emp"]["Emma"]["employee_code"]]
    assert emma["gross"] == money("52000.00", "INR") and emma["net"] == money("47000.00", "INR")
    assert emma["run_stage"] == "Final" and emma["legal_entity"] == "LE01" and emma["pay_group"] == "PG-IN"
    # EC-MON-04: totals per currency, never added across currencies.
    assert data["totals"] == {
        "EUR": {"deductions": "850.00", "employer_cost": "7490.00", "gross": "7000.00", "net": "6150.00",
                "reimbursements": "0.00"},
        "INR": {"deductions": "9000.00", "employer_cost": "95500.00", "gross": "92000.00", "net": "83000.00",
                "reimbursements": "0.00"},
    }
    assert data["converted"] is None
    # Filters narrow the same scoped query.
    assert preview(w["acme"], "payroll_register", legal_entity_id=w["zz01"], **SEP_FILTER)["row_count"] == 2
    assert preview(w["acme"], "payroll_register", currency="INR", **SEP_FILTER)["row_count"] == 2
    assert preview(w["acme"], "payroll_register", pay_group_id=w["groups"]["IN"], period_from="2026-08-01",
                   period_to="2026-09-30")["row_count"] == 4
    # Final runs by default; APPROVED / ALL include the approved UAE run labelled by its stage.
    approved = preview(w["acme"], "payroll_register", status="APPROVED", **SEP_FILTER)
    assert [r["run_stage"] for r in approved["rows"]] == ["Approved"] and approved["totals"].keys() == {"AED"}
    assert preview(w["acme"], "payroll_register", status="ALL", **SEP_FILTER)["row_count"] == 5
    # Bad filters fail clearly.
    assert w["acme"].get(f"{REPORTS}/payroll_register/preview",
                         params={"period_from": "2026-09-30", "period_to": "2026-09-01"}).json()["code"] == \
        "INVALID_PERIOD_RANGE"
    assert w["acme"].get(f"{REPORTS}/payroll_register/preview", params={"currency": "inr"}).status_code == 422
    assert w["acme"].get(f"{REPORTS}/payroll_register/preview", params={"currency": "ZZZ"}).json()["code"] == \
        "UNKNOWN_CURRENCY"
    assert w["acme"].get(f"{REPORTS}/payroll_register/preview", params={"status": "BOGUS"}).json()["code"] == \
        "INVALID_STATUS"


def test_gross_to_net_variance_overtime_and_exclusions(rc_world):
    w = rc_world
    g2n = preview(w["acme"], "gross_to_net", **SEP_FILTER)["rows"]
    inr = {(r["section"], r["component_code"]): r for r in g2n if r["currency"] == "INR"}
    assert inr[("EARNING", "BASIC")]["amount"] == money("90000.00", "INR")
    assert inr[("EARNING", "BASIC")]["employees"] == 2
    assert inr[("DEDUCTION", "IN_EPF_EE")]["label"] == "FIXTURE — not law"
    assert inr[("TOTAL", "NET")]["amount"] == money("83000.00", "INR")
    eur = {(r["section"], r["component_code"]): r for r in g2n if r["currency"] == "EUR"}
    assert eur[("DEDUCTION", "ZZ_PENSION_EE")]["label"] == "Statutory"
    assert eur[("EMPLOYER_CONTRIBUTION", "ZZ_PENSION_ER")]["amount"] == money("490.00", "EUR")

    variance = {r["employee_code"]: r for r in preview(w["acme"], "variance", **SEP_FILTER)["rows"]}
    emma = variance[w["emp"]["Emma"]["employee_code"]]
    assert emma["previous_net"] == money("45000.00", "INR") and emma["net_delta"] == money("2000.00", "INR")
    assert emma["net_delta_percent"] == "4.44" and emma["flag"] is None
    assert variance[w["emp"]["Ivan"]["employee_code"]]["net_delta"] == money("0.00", "INR")
    assert variance[w["emp"]["Zoe"]["employee_code"]]["flag"] == "NEW"

    overtime = preview(w["acme"], "overtime_by_bucket", **SEP_FILTER)
    assert [(r["component_code"], r["units"], r["amount"]) for r in overtime["rows"]] == [
        ("OT_WORKDAY", "4.000000", money("2000.00", "INR"))]
    exclusions = preview(w["acme"], "exclusions", **SEP_FILTER)["rows"]
    assert [(r["employee_code"], r["status"], r["reason"]) for r in exclusions] == [
        (w["emp"]["Nora"]["employee_code"], "EXCLUDED", "NO_SALARY")]


def test_corrected_view_shows_original_delta_and_corrected_amounts(rc_world):
    w = rc_world
    aug_run = w["runs"]["IN_AUG"]["id"]
    seed_run(w["cid"], entity=w["le01"], group=w["groups"]["IN"], period=w["periods"]["IN"][SEP], currency="INR",
             run_type="RETRO_ADJUSTMENT", members=[
                 member(w["emp"]["Emma"], "1000.00", "0.00", "1000.00", "1000.00", lines=[
                     line("BASIC", amount="1000.00", category="RETRO_ADJUSTMENT",
                          period_ref={"run_id": aug_run, "start": "2026-08-01", "end": "2026-08-31"})])])
    rows = preview(w["acme"], "corrected_view", **SEP_FILTER)["rows"]
    assert len(rows) == 1
    row = rows[0]
    assert row["correction_type"] == "RETRO_ADJUSTMENT" and row["component_code"] == "BASIC"
    assert row["original_amount"] == money("50000.00", "INR") and row["delta"] == money("1000.00", "INR")
    assert row["corrected_amount"] == money("51000.00", "INR")
    assert row["original_period"] == "2026-08-01 – 2026-08-31"
    arrears = preview(w["acme"], "arrears", **SEP_FILTER)["rows"]
    assert [(r["direction"], r["original_period_start"]) for r in arrears] == [("ARREARS", "2026-08-01")]
    # Cleo (ZZ01 only) never sees the Indian correction nor its original.
    assert preview(w["cleo"], "corrected_view", **SEP_FILTER)["rows"] == []


def test_ec_sec_05_scoped_users_get_scoped_rows_counts_and_totals(rc_world):
    w = rc_world
    cleo = w["cleo"]
    register = preview(cleo, "payroll_register", **SEP_FILTER)
    assert register["row_count"] == 2
    assert {r["legal_entity"] for r in register["rows"]} == {"ZZ01"}
    assert register["totals"].keys() == {"EUR"}
    # Asking for an entity outside the scope reveals nothing (no rows, no totals).
    outside = preview(cleo, "payroll_register", legal_entity_id=w["le01"], **SEP_FILTER)
    assert outside["row_count"] == 0 and outside["totals"] == {}
    assert {r["currency"] for r in preview(cleo, "gross_to_net", **SEP_FILTER)["rows"]} == {"EUR"}
    liabilities = preview(cleo, "liabilities_remittances")
    assert {r["scheme"] for r in liabilities["rows"]} == {"ZZ.pension", "ZZ.wht"}
    assert liabilities["totals"] == {"EUR": {"employee_amount": "850.00", "employer_amount": "490.00",
                                             "total": "1340.00"}}
    full = preview(w["acme"], "liabilities_remittances")
    assert full["totals"]["INR"] == {"employee_amount": "3600.00", "employer_amount": "3500.00", "total": "7100.00"}
    assert full["rows"][0]["label"] in (None, "FIXTURE — not law")


def test_compensation_changes_need_compensation_access_beyond_self(app):
    w = build_world(app, salaries=True)  # salaries before the runs: no arrears for finalized periods
    founder = preview(w["acme"], "compensation_changes")
    assert {r["employee_code"] for r in founder["rows"]} == {w["emp"][n]["employee_code"] for n in ("Emma", "Ivan")}
    assert founder["totals"] == {"INR": {"pay_basis_amount": "2400000.00"}}
    # Reports access plus only the SELF compensation grant every employee holds: no one else's salary.
    reader = grant_role(app, w["acme"], "Rita", [("payroll:reports:read", "COMPANY"),
                                                  ("payroll:compensation:read", "SELF")])
    assert preview(reader, "compensation_changes")["rows"] == []


def test_ec_sec_05_small_groups_are_suppressed_in_salary_analytics(rc_world):
    from app.payroll.services.reports import suppress

    w = rc_world
    data = preview(w["acme"], "employer_cost", group_by="legal_entity", **SEP_FILTER)
    assert len(data["rows"]) == 2 and all(r.get("suppressed") for r in data["rows"])
    assert all(r["employer_cost"] is None and r["employees"] is None for r in data["rows"])
    assert data["totals"] == {} and data["suppressed_groups"] == 2
    assert any("suppression" in note for note in data["notes"])

    set_setting(w["cid"], "privacy.small_group_suppression_size", 2)
    data = preview(w["acme"], "employer_cost", group_by="legal_entity", **SEP_FILTER)
    rows = {r["group"]: r for r in data["rows"]}
    assert rows["LE01"]["employer_cost"] == money("95500.00", "INR") and rows["LE01"]["employees"] == 2
    assert rows["ZZ01"]["employer_contributions"] == money("490.00", "EUR")
    assert data["suppressed_groups"] == 0 and data["totals"]["EUR"]["gross"] == "7000.00"

    # Complementary suppression: a single small group of a currency hides the next smallest too, so it
    # cannot be derived from the total.
    rows = [{"currency": "INR", "employees": 1, "cost": Decimal("10")},
            {"currency": "INR", "employees": 9, "cost": Decimal("90")},
            {"currency": "INR", "employees": 7, "cost": Decimal("70")},
            {"currency": "EUR", "employees": 8, "cost": Decimal("80")}]
    assert suppress(rows, ["cost"], 5) == 2
    assert [r.get("suppressed", False) for r in rows] == [True, False, True, False]
    assert rows[2]["cost"] is None and rows[3]["cost"] == Decimal("80")


def test_ec_mon_05_converted_total_is_labelled_and_never_one_to_one(rc_world):
    from app.database.db import system_session
    from app.models.catalogue import FxRate
    from app.payroll.services import fx

    w = rc_world
    data = preview(w["acme"], "payroll_register", convert_to="INR", fx_date="2026-09-30", **SEP_FILTER)
    assert data["converted"]["available"] is False
    assert [i["currency"] for i in data["converted"]["issues"]] == ["EUR"]
    assert data["totals"].keys() == {"EUR", "INR"}  # amounts stay per currency

    with system_session() as db:
        db.add(FxRate(company_id=w["cid"], base_currency="EUR", quote_currency="INR", rate=Decimal("90.5"),
                      rate_type="MID", source="MANUAL", as_of=date(2026, 9, 30), fetched_at=fx.clock.utcnow(),
                      status="ACTIVE"))
        db.commit()
    data = preview(w["acme"], "payroll_register", convert_to="INR", fx_date="2026-09-30", **SEP_FILTER)
    converted = data["converted"]
    assert converted["available"] is True and converted["currency"] == "INR"
    assert converted["totals"]["gross"] == money("725500.00", "INR")  # 92,000 + 7,000 x 90.5
    [rate] = converted["rates"]
    assert (rate["base"], rate["quote"], Decimal(rate["rate"]), rate["source"], rate["as_of"]) == (
        "EUR", "INR", Decimal("90.5"), "MANUAL", "2026-09-30")
    assert converted["label"]


def test_ec_doc_17_exports_escape_formulas_and_are_private_to_the_requester(app, rc_world):
    from openpyxl import load_workbook

    from tests.payroll.test_wage_inputs import audit_actions

    w = rc_world
    sql("UPDATE statutory_liabilities SET status = 'REMITTED', remittance_ref = '=HYPERLINK(\"http://x\")' "
        "WHERE company_id = :c AND scheme = 'ZZ.wht'", c=w["cid"])
    export_id, response = export_ready(w["acme"], "liabilities_remittances")
    assert response.headers["content-type"].startswith("text/csv")
    assert "no-store" in response.headers["cache-control"]
    rows = csv_rows(response.content)
    assert rows[0][0].startswith("Confidential — exported by ")
    cells = [cell for row in rows for cell in row]
    assert "'=HYPERLINK(\"http://x\")" in cells and "=HYPERLINK(\"http://x\")" not in cells
    assert any(row and row[0] == "Total EUR" for row in rows)
    assert audit_actions(w["cid"], "report_export", export_id) == [
        "REPORT_EXPORT_REQUESTED", "REPORT_EXPORTED", "REPORT_EXPORT_DOWNLOADED"]

    _, xlsx = export_ready(w["acme"], "liabilities_remittances", "XLSX")
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    values = [c.value for row in sheet.iter_rows() for c in row if isinstance(c.value, str)]
    assert "'=HYPERLINK(\"http://x\")" in values and "=HYPERLINK(\"http://x\")" not in values
    _, pdf = export_ready(w["acme"], "payroll_register", "PDF", SEP_FILTER)
    assert pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == "application/pdf"

    # Only the requester sees or downloads an export; another tenant gets 404 too.
    processor = with_roles(app, w["acme"], "Pia", ["payroll_processor"])
    assert processor.get(f"{EXPORTS}/{export_id}").status_code == 404
    assert processor.get(f"{EXPORTS}/{export_id}/download").status_code == 404
    assert all(e["id"] != export_id for e in ok(processor.get(EXPORTS))["items"])
    beta = other_tenant(app, "Beta Industries")["client"]
    assert beta.get(f"{EXPORTS}/{export_id}/download").status_code == 404
    # Exports are strict requests.
    assert w["acme"].post(f"{REPORTS}/payroll_register/exports", json={"format": "DOCX"}).status_code == 422
    assert w["acme"].post(f"{REPORTS}/payroll_register/exports",
                          json={"format": "CSV", "filters": {"company_id": 1}}).status_code == 422


def test_exports_mask_sensitive_columns_unless_the_exporter_has_full_access(app, rc_world):
    from app.payroll import audit as payroll_audit
    from app.services.audit_log_service import AuditLogService

    w = rc_world
    founder_user = sql("SELECT id FROM users WHERE email = :m", m=w["acme"].email)[0][0]
    with session(w["cid"], founder_user) as (db, principal):
        AuditLogService.log(db, principal, payroll_audit.PAYROLL_SETTINGS, "REPORT_TEST_EVENT",
                            "=SUM(A1:A9) changed", "setting_value", 1, company_id=w["cid"], ip_address="203.0.113.77")
        db.commit()
    audit_grants = [("payroll:audit:read", "COMPANY")]
    full = grant_role(app, w["acme"], "Fiona", audit_grants + [("payroll:reports:read", "COMPANY"),
                                                               ("payroll:reports:export", "COMPANY")])
    scoped = grant_role(app, w["acme"], "Sam", audit_grants + [("payroll:reports:read", "LEGAL_ENTITY"),
                                                               ("payroll:reports:export", "LEGAL_ENTITY")],
                        [w["zz01"]])

    def exported(client):
        _, response = export_ready(client, "settings_policy_changes")
        rows = csv_rows(response.content)
        header = next(i for i, row in enumerate(rows) if row and row[0] == "At")
        return rows[header], [dict(zip(rows[header], row)) for row in rows[header + 1:] if row]

    _, full_rows = exported(full)
    change = next(r for r in full_rows if r["Action"] == "REPORT_TEST_EVENT")
    assert change["Actor"] == w["acme"].email and change["IP address"] == "203.0.113.77"
    assert change["Description"].startswith("'=SUM(")  # EC-DOC-17
    _, masked_rows = exported(scoped)
    change = next(r for r in masked_rows if r["Action"] == "REPORT_TEST_EVENT")
    assert change["Actor"] != w["acme"].email and change["Actor"].endswith(w["acme"].email[-4:])
    assert "203.0.113.77" not in change["IP address"]


def test_ec_sec_01_three_tenants_never_see_each_others_report_rows(app, rc_world):
    w = rc_world
    beta = other_tenant(app, "Beta Industries")
    cobalt = other_tenant(app, "Cobalt Labs", currency="USD")
    for tenant, count, currency in ((beta, 1, "INR"), (cobalt, 1, "USD")):
        data = preview(tenant["client"], "payroll_register", **SEP_FILTER)
        assert data["row_count"] == count and data["totals"].keys() == {currency}
        assert data["rows"][0]["employee_code"] == tenant["employee"]["employee_code"]
        # Another tenant's entity id as a filter reveals nothing.
        assert preview(tenant["client"], "payroll_register", legal_entity_id=w["le01"], **SEP_FILTER)["row_count"] == 0
        assert preview(tenant["client"], "liabilities_remittances")["totals"].keys() == {currency}
    acme = preview(w["acme"], "payroll_register", **SEP_FILTER)
    assert acme["row_count"] == 4 and acme["totals"].keys() == {"EUR", "INR"}
    assert {beta["employee"]["employee_code"], cobalt["employee"]["employee_code"]}.isdisjoint(
        r["employee_code"] for r in acme["rows"])
    assert preview(w["acme"], "payroll_register", legal_entity_id=beta["entity"], **SEP_FILTER)["row_count"] == 0
