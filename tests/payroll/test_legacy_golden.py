"""M0 golden tests: freeze today's legacy payroll behaviour (README §8 M0, 3.3 §10 safety net).

Scenarios L01-L17 of audit 02 §11 (L18, the inventory, is in test_legacy_inventory.py).
Every assertion records what the legacy code returns *today*, known defects included;
each defect is named (audit 02 §9, D1-D24) in the test docstring. The v2 legacy-compatible
policy must reproduce the amounts; the defects are what v2 deliberately fixes.

All periods are fixed months of 2025 (or earlier), far from the real clock.
"""

import hashlib
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest

from tests.attendance_helpers import owner_rows
from tests.helpers import create_employee, invite_and_activate, signup_company
from tests.payroll.legacy_fixtures import (
    CALC_MONEY_KEYS,
    VECTOR_BY_ID,
    VECTORS,
    calc_url,
    dec,
    documents,
    holiday,
    insert_rows,
    legacy_reference,
    mark,
    payroll_row,
    payroll_tenant,
    pdf_text,
    structure,
    weekdays,
)

SOD_DETAIL = "Separation of duties: payroll must be approved by someone other than its preparer."


@pytest.fixture
def co(app):
    return payroll_tenant(app, "Golden Ledger Ltd")


def _gulf(co):
    """A second legal entity (AE01, AED) with a Dubai branch."""
    r = co.admin.post("/api/v1/legal-entities", json={"code": "AE01", "legal_name": "Golden Gulf FZ-LLC",
                                                      "country_code": "AE", "base_currency": "AED",
                                                      "timezone": "Asia/Dubai"})
    assert r.status_code == 200, r.text
    entity = r.json()
    r = co.admin.post("/api/v1/branches", json={"legal_entity_id": entity["id"], "code": "DXB", "name": "Dubai"})
    assert r.status_code == 200, r.text
    return entity, r.json()


def _generate(client, employee_id, year, month, ot="0", bonus="0") -> dict:
    r = client.post(calc_url(employee_id, year, month, ot, bonus, action="generate"))
    assert r.status_code == 200, r.text
    return r.json()


def _approve(client, payroll_id) -> dict:
    r = client.post(f"/api/v1/payroll/{payroll_id}/approve")
    assert r.status_code == 200, r.text
    return r.json()


def _run_vector(co, vector):
    employee = create_employee(co.admin, f"Gold{vector.id}")
    structure(co.admin, employee["id"], **vector.structure)
    if vector.weekday_holidays:
        for day in weekdays(vector.year, vector.month):
            holiday(co.admin, day.isoformat())
    for day, status in vector.attendance:
        mark(co.admin, employee["id"], day, status)
    r = co.preparer.post(calc_url(employee["id"], vector.year, vector.month, vector.ot, vector.bonus))
    assert r.status_code == 200, r.text
    generated = _generate(co.preparer, employee["id"], vector.year, vector.month, vector.ot, vector.bonus)
    return employee, r.json(), generated


# ---------------------------------------------------------------------------
# L01-L03, L05, L17: golden vectors (audit 02 §2)
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("vector", VECTORS, ids=[v.id for v in VECTORS])
def test_L01_reference_oracle_reproduces_golden_vectors(vector):
    """The pure oracle (legacy_fixtures.legacy_reference) gives the frozen amounts of every vector."""
    result = legacy_reference(vector.parts, vector.wdays, vector.count("absent"), vector.count("half-day"),
                              vector.count("late"), vector.ot, vector.bonus)
    for key, expected in vector.expected.items():
        assert result[key] == Decimal(expected), key


@pytest.mark.parametrize("vector_id", [v.id for v in VECTORS])
def test_L01_golden_vector_calculate_and_stored_row(co, vector_id):
    """calculate returns every amount as a float (D3) equal to the vector; generate stores 2-dp Decimals.

    Stored row: gross excludes OT and bonus (D9), other_deductions lumps absent+half+late (D10),
    leave_days holds the half-day count (D10), net is clipped at zero (D2, G4).
    """
    vector = VECTOR_BY_ID[vector_id]
    employee, calc, generated = _run_vector(co, vector)
    for key in CALC_MONEY_KEYS:
        assert isinstance(calc[key], float), key
    for key, expected in vector.expected.items():
        if key != "raw_net":
            assert dec(calc[key]) == Decimal(expected), key
    present, late = vector.count("present"), vector.count("late")
    assert (calc["working_days"], calc["present_days"], calc["absent_days"], calc["half_days"], calc["late_days"]) == (
        vector.wdays, present + late, vector.count("absent"), vector.count("half-day"), late)
    assert calc["employee_id"] == employee["id"] and calc["employee_name"] == f"Gold{vector_id} Test"
    assert dec(calc["basic_salary"]) == Decimal(vector.structure["basic_salary"])
    assert dec(calc["hra"]) == Decimal(vector.structure.get("house_rent_allowance", "0"))

    row = payroll_row(generated["id"])
    e = vector.expected
    assert str(row["gross_salary"]) == e["gross_salary"]
    assert str(row["other_deductions"]) == e["total_deductions"]
    assert str(row["total_deductions"]) == e["total_deductions"]
    assert str(row["net_salary"]) == e["net_salary"]
    assert str(row["overtime_hours"]) == e["overtime_hours"]
    assert str(row["overtime_amount"]) == e["overtime_amount"]
    assert str(row["bonus"]) == e["bonus"]
    assert (row["working_days"], row["present_days"], row["absent_days"], row["leave_days"]) == (
        vector.wdays, present + late, vector.count("absent"), vector.count("half-day"))
    assert row["salary_month"] == date(vector.year, vector.month, 1) and row["status"] == "draft"


def test_L02_stored_row_versus_api_for_G2(co):
    """G2 as stored: the three attendance deductions and the late count are not kept separately (D10),
    statutory columns are always 0, gross excludes OT and bonus (D9), hours are an informational snapshot."""
    vector = VECTOR_BY_ID["G2"]
    employee, calc, generated = _run_vector(co, vector)
    assert set(generated) == {"id", "employee_id", "employee_name", "employee_code", "salary_month", "basic_salary",
                              "gross_salary", "total_deductions", "net_salary", "status", "currency",
                              "legal_entity_id", "net_work_seconds", "created_by_user_id", "approved_by_user_id",
                              "created_at"}
    assert (generated["gross_salary"], generated["total_deductions"], generated["net_salary"]) == (
        35000.0, 4466.66, 32471.34)
    assert (generated["status"], generated["currency"], generated["legal_entity_id"]) == ("draft", "INR", co.le01["id"])
    assert generated["created_by_user_id"] == co.preparer.user_id and generated["approved_by_user_id"] is None

    row = payroll_row(generated["id"])
    assert str(row["gross_salary"]) == "35000.00"
    assert str(row["other_deductions"]) == str(row["total_deductions"]) == "4466.66"
    assert row["leave_days"] == 1  # half-day count stored in the leave column
    assert "late_days" not in row and "late_penalty" not in row and "absent_deduction" not in row
    for statutory in ("provident_fund", "insurance", "professional_tax", "income_tax"):
        assert str(row[statutory]) == "0.00"
    assert (str(row["basic_salary"]), str(row["house_rent_allowance"]), str(row["dearness_allowance"]),
            str(row["conveyance_allowance"]), str(row["other_allowances"])) == (
        "30000.00", "5000.00", "0.00", "0.00", "0.00")
    assert row["created_by_user_id"] == co.preparer.user_id
    assert row["approved_by_user_id"] is None and row["processed_by_user_id"] is None
    assert row["remarks"] is None
    # G2's net equation holds (no clipping): gross + OT + bonus - deductions.
    assert row["gross_salary"] + row["overtime_amount"] + row["bonus"] - row["total_deductions"] == row["net_salary"]
    # Hours snapshot: six manual 8-hour rows, whatever their status.
    hours = calc["attendance_hours"]
    assert (row["gross_work_seconds"], row["net_work_seconds"], row["break_seconds"], row["incomplete_sessions"]) == (
        hours["gross_seconds"], hours["net_seconds"], hours["lunch_seconds"], hours["incomplete_sessions"])
    assert hours["gross_seconds"] == 6 * 8 * 3600 and hours["incomplete_sessions"] == 0


def test_L03_net_clipped_at_zero_EC_RUN_10_defect_D2(co):
    """DEFECT D2 / EC-RUN-10 (frozen): G4's raw net is -400.00 but the API, the stored row, the slip and the
    payslip all say 0.00, while total deductions stay 1400.00, so the stored row breaks its own net equation."""
    vector = VECTOR_BY_ID["G4"]
    employee, calc, generated = _run_vector(co, vector)
    assert calc["net_salary"] == 0.0 and calc["total_deductions"] == 1400.0 and calc["absent_days"] == 28
    assert calc["working_days"] == 20  # 8 of the 28 absent rows fall on weekends and still deduct (D6)
    row = payroll_row(generated["id"])
    assert str(row["net_salary"]) == "0.00" and str(row["total_deductions"]) == "1400.00"
    assert row["gross_salary"] + row["overtime_amount"] + row["bonus"] - row["total_deductions"] == Decimal("-400.00")
    _approve(co.approver, generated["id"])
    slip = co.admin.get(f"/api/v1/payroll/{generated['id']}/slip").json()
    assert slip["net_salary"] == 0.0 and slip["deductions"]["attendance_deductions"] == 1400.0
    assert documents(generated["id"])[0]["snapshot"]["net"] == 0.0


def test_L05_zero_working_days_pays_in_full_without_a_flag_defect_D6(co):
    """DEFECT D6 (frozen): when every weekday is a holiday the per-day rate is 0, nothing is deducted or
    flagged, and the row is generated, approved and printed without error (G5 has the amounts)."""
    employee = create_employee(co.admin, "Zero")
    structure(co.admin, employee["id"], basic_salary="1000")
    for day in weekdays(2025, 6):
        holiday(co.admin, day.isoformat())
    mark(co.admin, employee["id"], "2025-06-02", "absent")
    calc = co.preparer.post(calc_url(employee["id"], 2025, 6)).json()
    assert (calc["working_days"], calc["absent_deduction"], calc["net_salary"]) == (0, 0.0, 1000.0)
    generated = _generate(co.preparer, employee["id"], 2025, 6)
    assert payroll_row(generated["id"])["working_days"] == 0
    _approve(co.approver, generated["id"])
    pdf = co.admin.get(f"/api/v1/payroll/{generated['id']}/slip.pdf")
    assert pdf.status_code == 200 and "₹1,000.00" in pdf_text(pdf.content)


def test_L17_rounding_edges_half_up_on_hours_and_bonus(co):
    """Hours and bonus are rounded to 2 dp HALF_UP after a float round trip (Decimal(str(float))):
    OT 1.005 h -> 1.01 (G7); bonus 0.005 -> 0.01, 0.004 -> 0.00; OT 0.005 h -> 0.01 h -> 2.08.
    Query bounds: negative bonus and more than 500 OT hours are 422."""
    employee = create_employee(co.admin, "Round")
    structure(co.admin, employee["id"], basic_salary="30000", house_rent_allowance="5000")

    def calc(**params):
        r = co.preparer.post(calc_url(employee["id"], 2025, 3, **params))
        assert r.status_code == 200, r.text
        return r.json()

    up = calc(bonus="0.005")
    assert dec(up["bonus"]) == Decimal("0.01") and dec(up["net_salary"]) == Decimal("35000.01")
    down = calc(bonus="0.004")
    assert dec(down["bonus"]) == Decimal("0") and dec(down["net_salary"]) == Decimal("35000")
    tiny = calc(ot="0.005")
    assert dec(tiny["overtime_hours"]) == Decimal("0.01") and dec(tiny["overtime_amount"]) == Decimal("2.08")
    g7 = legacy_reference(["30000", "5000"], 21, ot="1.005")
    assert g7["overtime_hours"] == Decimal("1.01")
    assert co.preparer.post(calc_url(employee["id"], 2025, 3, bonus="-1")).status_code == 422
    assert co.preparer.post(calc_url(employee["id"], 2025, 3, ot="500.01")).status_code == 422
    assert co.preparer.post(calc_url(employee["id"], 2025, 13)).status_code == 422


# ---------------------------------------------------------------------------
# L04: working days = Monday-Friday minus non-optional holidays for the employee's location
# ---------------------------------------------------------------------------
def test_L04_holidays_reduce_weekdays_by_exact_location_match(co):
    """National weekday -1, optional 0, weekend 0, regional for the employee's location -1, another location 0.
    DEFECT D6 (frozen): Monday-Friday is hard-coded and the location match is an exact, case-sensitive string."""
    employee = create_employee(co.admin, "Pune", location="Pune")
    structure(co.admin, employee["id"], basic_salary="21000")

    def wdays():
        r = co.preparer.post(calc_url(employee["id"], 2025, 8))
        assert r.status_code == 200, r.text
        return r.json()["working_days"]

    assert wdays() == 21
    steps = [
        ({"date": "2025-08-15", "name": "Independence Day"}, 20),
        ({"date": "2025-08-18", "name": "Optional Day", "is_optional": True}, 20),
        ({"date": "2025-08-16", "name": "Saturday Fair"}, 20),
        ({"date": "2025-08-19", "name": "Pune Day", "location": "Pune"}, 19),
        ({"date": "2025-08-20", "name": "Delhi Day", "location": "Delhi"}, 19),
        ({"date": "2025-08-21", "name": "Lowercase Pune", "location": "pune"}, 19),
    ]
    for body, expected in steps:
        holiday(co.admin, body.pop("date"), **body)
        assert wdays() == expected, body["name"]
    calc = co.preparer.post(calc_url(employee["id"], 2025, 8)).json()
    assert calc["net_salary"] == 21000.0  # holidays only change the per-day rate
    generated = _generate(co.preparer, employee["id"], 2025, 8)
    assert payroll_row(generated["id"])["working_days"] == 19


# ---------------------------------------------------------------------------
# L06-L07: structures and dead settings
# ---------------------------------------------------------------------------
def test_L06_structure_selection_mid_month_future_and_missing_defect_D7(co):
    """The latest structure starting on or before the month end applies to the WHOLE month (no proration,
    DEFECT D7); a structure starting after the month is ignored; no structure -> 400 NO_SALARY_STRUCTURE."""
    employee = create_employee(co.admin, "Revised")
    structure(co.admin, employee["id"], basic_salary="30000")
    structure(co.admin, employee["id"], effective_from="2025-03-15", basic_salary="40000")
    structure(co.admin, employee["id"], effective_from="2025-06-01", basic_salary="50000")
    windows = {s["effective_from"]: s["effective_to"] for s in
               co.admin.get(f"/api/v1/salary-structures?employee_id={employee['id']}").json()}
    assert windows == {"2024-01-01": "2025-03-14", "2025-03-15": "2025-05-31", "2025-06-01": None}

    def gross(year, month):
        r = co.preparer.post(calc_url(employee["id"], year, month))
        assert r.status_code == 200, r.text
        return r.json()["gross_salary"]

    assert gross(2025, 2) == 30000.0
    assert gross(2025, 3) == 40000.0  # starts on the 15th, applies from the 1st
    assert gross(2025, 5) == 40000.0  # the June revision is ignored
    assert gross(2025, 6) == 50000.0

    r = co.preparer.post(calc_url(employee["id"], 2023, 12))
    assert r.status_code == 400 and r.json()["code"] == "NO_SALARY_STRUCTURE"
    assert r.json()["detail"] == "No salary structure covers 12/2023 for Revised Test."
    bare = create_employee(co.admin, "Bare")
    r = co.preparer.post(calc_url(bare["id"], 2025, 3, action="generate"))
    assert r.status_code == 400 and r.json()["code"] == "NO_SALARY_STRUCTURE"
    assert owner_rows("SELECT count(*) AS n FROM payroll")[0]["n"] == 0


def test_L07_pf_tax_overtime_settings_are_dead_EC_MIG_04_EC_MIG_05_defect_D8(co):
    """DEFECT D8 / EC-MIG-04 / EC-MIG-05 (frozen): the company settings overtime, pf and tax are stored but
    never read; the late penalty (EC-MIG-03, D4) and hours per day come from global env config."""
    from app.config import settings

    vector = VECTOR_BY_ID["G2"]
    employee = create_employee(co.admin, "Dead")
    structure(co.admin, employee["id"], **vector.structure)
    for day, status in vector.attendance:
        mark(co.admin, employee["id"], day, status)
    url = calc_url(employee["id"], 2025, 3, vector.ot, vector.bonus)
    before = co.preparer.post(url).json()
    r = co.admin.put("/api/v1/settings", json={"overtime": 3, "pf": 50, "tax": 50})
    assert r.status_code == 200, r.text
    stored = co.admin.get("/api/v1/settings").json()
    assert (stored["overtime"], stored["pf"], stored["tax"]) == (3.0, 50.0, 50.0)
    after = co.preparer.post(url).json()
    assert after == before
    assert (after["late_penalty"], after["overtime_amount"]) == (300.0, 937.5)
    assert (settings.late_penalty_amount, settings.default_work_hours_per_day) == (100.0, 8.0)


# ---------------------------------------------------------------------------
# L08-L09: currency and legal entity stamping
# ---------------------------------------------------------------------------
def test_L08_aed_entity_gets_the_global_late_penalty_in_AED_EC_MIG_03_defect_D4(co):
    """DEFECT D4 / EC-MIG-03 (frozen): 100 per late mark from env config, applied in whatever currency the
    row is in. An AE01 employee with 2 late marks: currency AED, legal entity AE01, late 200.00 AED."""
    entity, dubai = _gulf(co)
    employee = create_employee(co.admin, "Fatima", branch_id=dubai["id"])
    structure(co.admin, employee["id"], basic_salary="12000")
    for day in ("2025-06-09", "2025-06-10"):
        mark(co.admin, employee["id"], day, "late")
    calc = co.preparer.post(calc_url(employee["id"], 2025, 6)).json()
    assert (calc["late_penalty"], calc["total_deductions"], calc["net_salary"], calc["present_days"]) == (
        200.0, 200.0, 11800.0, 2)
    generated = _generate(co.preparer, employee["id"], 2025, 6)
    assert (generated["currency"], generated["legal_entity_id"]) == ("AED", entity["id"])
    row = payroll_row(generated["id"])
    assert (row["currency"], row["legal_entity_id"], str(row["other_deductions"]), str(row["net_salary"])) == (
        "AED", entity["id"], "200.00", "11800.00")
    slip = co.admin.get(f"/api/v1/payroll/{generated['id']}/slip").json()
    assert slip["currency"] == "AED" and slip["employer"]["legal_name"] == "Golden Gulf FZ-LLC"
    assert slip["employer"]["code"] == "AE01" and slip["employer"]["country_code"] == "AE"


def test_L09_currency_and_entity_stamping_transfer_default_and_trigger(co):
    """Entity at the END of the month (transfer on the 20th -> new entity for the whole month); no
    assignment at month end and none earlier -> the tenant default entity (LE01) even for an employee who
    only ever worked for AE01; rows inserted without currency get them from the payroll_fill_defaults
    trigger, which has no default-entity fallback (legal_entity_id stays NULL: frozen discrepancy)."""
    entity, dubai = _gulf(co)
    mover = create_employee(co.admin, "Mover")
    structure(co.admin, mover["id"], basic_salary="10000")
    r = co.admin.post(f"/api/v1/employees/{mover['id']}/assignments",
                      json={"branch_id": dubai["id"], "effective_from": "2025-07-20", "reason": "TRANSFER"})
    assert r.status_code == 200, r.text
    june = _generate(co.preparer, mover["id"], 2025, 6)
    july = _generate(co.preparer, mover["id"], 2025, 7)
    assert (june["currency"], june["legal_entity_id"]) == ("INR", co.le01["id"])
    assert (july["currency"], july["legal_entity_id"]) == ("AED", entity["id"])
    assert july["net_salary"] == 10000.0  # same amount, relabelled AED; nothing is converted or prorated

    gulf = create_employee(co.admin, "Gulf", branch_id=dubai["id"])  # hired 2024-01-15 into AE01
    structure(co.admin, gulf["id"], effective_from="2023-01-01", basic_salary="9000")
    before_hire = _generate(co.preparer, gulf["id"], 2023, 12)
    assert (before_hire["currency"], before_hire["legal_entity_id"]) == ("INR", co.le01["id"])
    first_month = _generate(co.preparer, gulf["id"], 2024, 1)
    assert (first_month["currency"], first_month["legal_entity_id"]) == ("AED", entity["id"])

    covered, uncovered = insert_rows(co.company_id, [
        {"employee_id": gulf["id"], "salary_month": date(2024, 2, 1), "basic_salary": Decimal("9000"),
         "gross_salary": Decimal("9000"), "net_salary": Decimal("9000")},
        {"employee_id": gulf["id"], "salary_month": date(2023, 11, 1), "basic_salary": Decimal("9000"),
         "gross_salary": Decimal("9000"), "net_salary": Decimal("9000")},
    ])
    assert (payroll_row(covered)["currency"], payroll_row(covered)["legal_entity_id"]) == ("AED", entity["id"])
    assert (payroll_row(uncovered)["currency"], payroll_row(uncovered)["legal_entity_id"]) == ("INR", None)
    assert payroll_row(covered)["status"] == "draft"


# ---------------------------------------------------------------------------
# L10: generate-month
# ---------------------------------------------------------------------------
def test_L10_generate_month_buckets_and_inactive_leavers_defect_D7_D18(co):
    """Active employees only (DEFECT D7: a leaver in the month is skipped); per-employee buckets success /
    already_exists / failed (no structure), each with the AppError detail (D18: synchronous loop)."""
    paid = create_employee(co.admin, "Paid")
    bare = create_employee(co.admin, "Bare")
    existing = create_employee(co.admin, "Existing")
    leaver = create_employee(co.admin, "Leaver")
    for emp in (paid, existing, leaver):
        structure(co.admin, emp["id"], basic_salary="20000")
    prior = _generate(co.preparer, existing["id"], 2025, 3)
    assert co.admin.post(f"/api/v1/employees/{leaver['id']}/deactivate").status_code == 200

    r = co.preparer.post("/api/v1/payroll/generate-month?year=2025&month=3")
    assert r.status_code == 200, r.text
    body = r.json()
    # Active: founder, preparer, approver, paid, bare, existing (the leaver is excluded).
    assert body["summary"] == {"total_employees": 6, "generated": 1, "failed": 4, "already_exists": 1}
    assert [x["employee_id"] for x in body["results"]["success"]] == [paid["id"]]
    assert [x["employee_id"] for x in body["results"]["already_exists"]] == [existing["id"]]
    assert body["results"]["already_exists"][0]["error"] == "Payroll already exists for this month."
    failed = {x["employee_id"]: x["error"] for x in body["results"]["failed"]}
    assert failed[bare["id"]] == "No salary structure covers 03/2025 for Bare Test."
    assert {co.preparer.employee["id"], co.approver.employee["id"], bare["id"]} <= set(failed)
    assert leaver["id"] not in failed
    months = owner_rows("SELECT employee_id, id FROM payroll WHERE salary_month = '2025-03-01' ORDER BY id")
    assert [m["employee_id"] for m in months] == [existing["id"], paid["id"]] and months[0]["id"] == prior["id"]
    again = co.preparer.post("/api/v1/payroll/generate-month?year=2025&month=3").json()["summary"]
    assert again == {"total_employees": 6, "generated": 0, "failed": 4, "already_exists": 2}


# ---------------------------------------------------------------------------
# L11: state machine and separation of duties
# ---------------------------------------------------------------------------
def test_L11_state_machine_and_separation_of_duties_EC_RUN_05_EC_RUN_06_defect_D12(co, app):
    """draft -> approved -> processed only. EC-RUN-05: the preparer cannot approve while SoD is ON (default).
    EC-RUN-06: nobody approves or processes their own pay. DEFECT D12 (frozen): the preparer and the
    approver may both process. DEFECT D11: no reject, reopen or regenerate (409 PAYROLL_EXISTS)."""
    worker = create_employee(co.admin, "Worker")
    structure(co.admin, worker["id"], basic_salary="20000")
    staff = invite_and_activate(app, co.admin, create_employee(co.admin, "Staff"), ["employee"])
    pid = _generate(co.preparer, worker["id"], 2025, 3)["id"]

    dup = co.preparer.post(calc_url(worker["id"], 2025, 3, action="generate"))
    assert dup.status_code == 409 and dup.json()["code"] == "PAYROLL_EXISTS"
    r = co.preparer.post(f"/api/v1/payroll/{pid}/approve")
    assert r.status_code == 403 and r.json()["detail"] == SOD_DETAIL and r.json()["code"] == "FORBIDDEN"
    r = co.approver.post(f"/api/v1/payroll/{pid}/process")
    assert r.status_code == 400 and r.json()["detail"] == "Payroll must be approved before it is processed."
    for path in ("approve", "process"):
        assert staff.post(f"/api/v1/payroll/{pid}/{path}").status_code == 403
    for url in (calc_url(worker["id"], 2025, 3), calc_url(worker["id"], 2025, 4, action="generate"),
                "/api/v1/payroll/generate-month?year=2025&month=3", "/api/v1/payroll/process-bulk?year=2025&month=3"):
        assert staff.post(url).status_code == 403, url

    approved = _approve(co.approver, pid)
    assert approved["status"] == "approved" and approved["approved_by_user_id"] == co.approver.user_id
    r = co.approver.post(f"/api/v1/payroll/{pid}/approve")
    assert r.status_code == 400 and r.json()["detail"] == "Only draft payroll can be approved (current: approved)."
    processed = co.preparer.post(f"/api/v1/payroll/{pid}/process")  # D12: the preparer pays
    assert processed.status_code == 200 and processed.json()["status"] == "processed"
    row = payroll_row(pid)
    assert (row["created_by_user_id"], row["approved_by_user_id"], row["processed_by_user_id"]) == (
        co.preparer.user_id, co.approver.user_id, co.preparer.user_id)
    assert row["approved_at"] is not None and row["processed_at"] is not None
    r = co.approver.post(f"/api/v1/payroll/{pid}/process")
    assert r.status_code == 400 and r.json()["detail"] == "Payroll must be approved before it is processed."
    r = co.approver.post(f"/api/v1/payroll/{pid}/approve")
    assert r.status_code == 400 and r.json()["detail"] == "Only draft payroll can be approved (current: processed)."

    # EC-RUN-06: own pay. The approver's pay, prepared by the preparer, needs another approver and payer.
    structure(co.admin, co.approver.employee["id"], basic_salary="30000")
    own = _generate(co.preparer, co.approver.employee["id"], 2025, 3)["id"]
    r = co.approver.post(f"/api/v1/payroll/{own}/approve")
    assert r.status_code == 403 and r.json()["detail"] == "You cannot approve your own pay."
    _approve(co.admin, own)
    r = co.approver.post(f"/api/v1/payroll/{own}/process")
    assert r.status_code == 403 and r.json()["detail"] == "You cannot process your own pay."
    # The preparer may prepare their own pay (legacy allows it) but not approve or process it.
    structure(co.admin, co.preparer.employee["id"], basic_salary="25000")
    mine = _generate(co.preparer, co.preparer.employee["id"], 2025, 3)["id"]
    r = co.preparer.post(f"/api/v1/payroll/{mine}/approve")
    assert r.status_code == 403 and r.json()["detail"] == "You cannot approve your own pay."
    _approve(co.approver, mine)

    # process-bulk: only approved rows of the month; own pay fails per row and stays approved.
    draft = _generate(co.preparer, worker["id"], 2025, 4)["id"]
    bulk = co.preparer.post("/api/v1/payroll/process-bulk?year=2025&month=3")
    assert bulk.status_code == 200, bulk.text
    body = bulk.json()
    assert body["summary"] == {"processed": 1, "failed": 1}
    assert [x["payroll_id"] for x in body["results"]["success"]] == [own]
    assert body["results"]["failed"] == [{"payroll_id": mine, "employee_id": co.preparer.employee["id"],
                                          "error": "You cannot process your own pay."}]
    assert payroll_row(mine)["status"] == "approved" and payroll_row(draft)["status"] == "draft"
    empty = co.approver.post("/api/v1/payroll/process-bulk?year=2025&month=4").json()
    assert empty["summary"] == {"processed": 0, "failed": 0}

    # SoD OFF (administrators only): the preparer may approve what they prepared; own pay is still refused.
    r = co.preparer.put("/api/v1/settings", json={"payrollSeparationOfDuties": False})
    assert r.status_code == 403
    r = co.admin.put("/api/v1/settings", json={"payrollSeparationOfDuties": False})
    assert r.status_code == 200, r.text
    assert _approve(co.preparer, draft)["approved_by_user_id"] == co.preparer.user_id
    mine_april = _generate(co.preparer, co.preparer.employee["id"], 2025, 4)["id"]
    r = co.preparer.post(f"/api/v1/payroll/{mine_april}/approve")
    assert r.status_code == 403 and r.json()["detail"] == "You cannot approve your own pay."

    logs = owner_rows("SELECT action, description, entity_id FROM audit_logs WHERE module = 'Payroll' "
                      "AND entity_id = :p ORDER BY id", p=str(pid))
    assert [entry["action"] for entry in logs] == ["GENERATE", "APPROVE", "PROCESS"]
    assert logs[1]["description"] == f"Approved payroll #{pid}; payslip PS-{co.company_id}-{pid:06d} issued"
    assert logs[2]["description"] == f"Processed payroll #{pid}"


def test_L11_processed_means_marked_paid_without_payment_evidence_defect_D1(co):
    """DEFECT D1 (frozen): processing only flips the status and stamps who/when; there is no payment
    reference, no settlement evidence and nothing to undo it."""
    worker = create_employee(co.admin, "Payee")
    structure(co.admin, worker["id"], basic_salary="20000")
    pid = _generate(co.preparer, worker["id"], 2025, 3)["id"]
    _approve(co.approver, pid)
    before = payroll_row(pid)
    co.approver.post(f"/api/v1/payroll/{pid}/process")
    after = payroll_row(pid)
    changed = {k for k in after if after[k] != before[k]}
    assert changed == {"status", "processed_by_user_id", "processed_at", "updated_at"}
    assert after["status"] == "processed"
    assert not any("payment" in column or "reference" in column for column in after)


# ---------------------------------------------------------------------------
# L12: payslip issuance and the immutable snapshot
# ---------------------------------------------------------------------------
SNAPSHOT_KEYS = {"template_version", "issued_at", "document_number", "status", "brand", "employer", "employee",
                 "period", "currency", "earnings", "deductions", "gross", "total_deductions", "net", "attendance"}


def test_L12_payslip_issued_on_approve_with_frozen_snapshot_EC_WL_09_defects_D9_D19(co):
    """One IssuedDocument per approved row (PS-{company}-{id:06d}, PAYSLIP_V2, sha256 of the stored bytes,
    current brand version). Drafts render live with a DRAFT watermark and are never stored. The snapshot
    keeps float amounts (D3), drops zero lines, and its gross excludes the OT and Bonus lines it lists (D9).
    DEFECT D19: the number leaks internal ids."""
    from app.services import brand_service

    vector = VECTOR_BY_ID["G2"]
    employee, _, generated = _run_vector(co, vector)
    pid = generated["id"]

    draft = co.admin.get(f"/api/v1/payroll/{pid}/slip.pdf")
    assert draft.status_code == 200 and draft.headers["content-type"] == "application/pdf"
    assert "x-document-number" not in draft.headers and draft.headers["cache-control"] == "private, no-store"
    draft_text = pdf_text(draft.content)
    assert "DRAFT" in draft_text and "(DRAFT - not issued)" in draft_text
    assert documents(pid) == []

    _approve(co.approver, pid)
    docs = documents(pid)
    assert len(docs) == 1
    doc = docs[0]
    number = f"PS-{co.company_id}-{pid:06d}"
    assert (doc["template_version"], doc["employee_id"], doc["legal_entity_id"], doc["issued_by_user_id"]) == (
        "PAYSLIP_V2", employee["id"], co.le01["id"], co.approver.user_id)
    assert doc["brand_version"] == brand_service.document_brand(co.company_id, co.le01["id"])["brand_version"]
    snapshot = doc["snapshot"]
    assert set(snapshot) == SNAPSHOT_KEYS
    assert (snapshot["document_number"], snapshot["status"], snapshot["period"], snapshot["currency"]) == (
        number, "issued", "March 2025", "INR")
    assert snapshot["template_version"] == "PAYSLIP_V2" and snapshot["brand"]["brand_version"] == doc["brand_version"]
    assert snapshot["earnings"] == [["Basic salary", 30000.0], ["House rent allowance", 5000.0],
                                    ["Overtime", 937.5], ["Bonus", 1000.5]]
    assert snapshot["deductions"] == [["Attendance deductions", 4466.66]]
    assert (snapshot["gross"], snapshot["total_deductions"], snapshot["net"]) == (35000.0, 4466.66, 32471.34)
    assert snapshot["attendance"] == {"working_days": 21, "present_days": 3, "absent_days": 2, "half_days": 1}
    assert snapshot["employee"]["name"] == "GoldG2 Test" and snapshot["employee"]["employee_code"] == employee["employee_code"]
    assert snapshot["employer"]["code"] == "LE01"

    issued = co.admin.get(f"/api/v1/payroll/{pid}/slip.pdf")
    sha = hashlib.sha256(issued.content).hexdigest()
    assert issued.headers["x-document-sha256"] == sha == doc["sha256"]
    assert issued.headers["x-document-number"] == number
    assert issued.headers["x-brand-version"] == doc["brand_version"]
    assert issued.headers["content-disposition"] == (
        f'attachment; filename="payslip-{employee["employee_code"]}-March-2025.pdf"')
    text = pdf_text(issued.content)
    for label in ("Gross earnings", "Total deductions", "Net pay", f"Document {number}", "₹35,000.00",
                  "₹32,471.34", "₹4,466.66", "Overtime", "Bonus"):
        assert label in text, label
    assert "DRAFT" not in text

    # The issued snapshot and bytes never change; the JSON slip reads live employee data.
    r = co.admin.put(f"/api/v1/employees/{employee['id']}", json={"first_name": "Renamed"})
    assert r.status_code == 200, r.text
    assert co.admin.get(f"/api/v1/payroll/{pid}/slip.pdf").content == issued.content
    assert documents(pid)[0]["snapshot"]["employee"]["name"] == "GoldG2 Test"
    slip = co.admin.get(f"/api/v1/payroll/{pid}/slip").json()
    assert slip["employee"]["name"] == "Renamed Test"
    assert co.approver.post(f"/api/v1/payroll/{pid}/process").status_code == 200
    assert len(documents(pid)) == 1

    from sqlalchemy import text as sql

    from app.database.db import engine

    with engine.begin() as conn:
        conn.execute(sql("SELECT set_config('app.company_id', :c, true)"), {"c": str(co.company_id)})
        with pytest.raises(Exception):
            conn.execute(sql("UPDATE issued_documents SET snapshot = '{}'::jsonb"))


def test_L12_legacy_row_without_document_is_issued_on_first_download_defect_D14(co):
    """DEFECT D14 (frozen): a row approved before documents existed is issued inside the first GET, by
    whoever downloads it, with the brand in force at download time; later downloads return the same bytes."""
    worker = create_employee(co.admin, "Legacy")
    structure(co.admin, worker["id"], basic_salary="20000")
    pid = _generate(co.preparer, worker["id"], 2025, 2)["id"]
    from sqlalchemy import text as sql

    from app.database.db import system_engine

    with system_engine.begin() as conn:
        conn.execute(sql("UPDATE payroll SET status = 'approved' WHERE id = :p"), {"p": pid})
    assert documents(pid) == []
    first = co.admin.get(f"/api/v1/payroll/{pid}/slip.pdf")
    assert first.status_code == 200
    docs = documents(pid)
    assert len(docs) == 1 and docs[0]["issued_by_user_id"] == co.admin.user_id
    assert first.headers["x-document-sha256"] == docs[0]["sha256"] == hashlib.sha256(first.content).hexdigest()
    assert co.admin.get(f"/api/v1/payroll/{pid}/slip.pdf").content == first.content
    assert len(documents(pid)) == 1


# ---------------------------------------------------------------------------
# L13: visibility, listing and tenant isolation
# ---------------------------------------------------------------------------
def test_L13_self_visibility_listing_and_cross_tenant_404_EC_DOC_01_EC_DOC_02(co, app):
    """SELF sees own approved/processed rows only (EC-DOC-02: draft -> 404); another employee's slip -> 404
    (EC-DOC-01); /payroll/my returns at most 36 rows, newest first; company lists need COMPANY scope (403);
    another tenant gets 404 for every object route and its own empty lists."""
    worker_emp = create_employee(co.admin, "Selfie")
    other_emp = create_employee(co.admin, "Other")
    worker = invite_and_activate(app, co.admin, worker_emp, ["employee"])
    other = invite_and_activate(app, co.admin, other_emp, ["employee"])
    structure(co.admin, worker_emp["id"], basic_salary="20000")
    draft = _generate(co.preparer, worker_emp["id"], 2025, 3)["id"]
    approved = _generate(co.preparer, worker_emp["id"], 2025, 4)["id"]
    processed = _generate(co.preparer, worker_emp["id"], 2025, 5)["id"]
    for pid in (approved, processed):
        _approve(co.approver, pid)
    co.approver.post(f"/api/v1/payroll/{processed}/process")

    assert [p["id"] for p in worker.get("/api/v1/payroll/my").json()] == [processed, approved]
    for suffix in ("", "/slip", "/slip.pdf"):
        assert worker.get(f"/api/v1/payroll/{draft}{suffix}").status_code == 404
        assert worker.get(f"/api/v1/payroll/{approved}{suffix}").status_code == 200
        assert other.get(f"/api/v1/payroll/{approved}{suffix}").status_code == 404
    assert other.get("/api/v1/payroll/my").json() == []
    assert worker.get("/api/v1/payroll").status_code == 403
    assert worker.get("/api/v1/payroll/summary/2025/4").status_code == 403
    assert worker.get("/api/v1/reports/payroll?year=2025").status_code == 403

    insert_rows(co.company_id, [
        {"employee_id": worker_emp["id"], "salary_month": date(2021 + (i // 12), i % 12 + 1, 1), "status": "approved",
         "basic_salary": Decimal("100"), "gross_salary": Decimal("100"), "net_salary": Decimal("100"),
         "currency": "INR"}
        for i in range(37)
    ])
    mine = worker.get("/api/v1/payroll/my").json()
    assert len(mine) == 36
    months = [p["salary_month"] for p in mine]
    assert months == sorted(months, reverse=True) and months[0] == "2025-05-01"
    assert months[-1] == "2021-04-01"  # the three oldest of 39 visible rows are cut off

    rival = signup_company(app, "Rival Payroll Co")
    for suffix in ("", "/slip", "/slip.pdf"):
        assert rival.get(f"/api/v1/payroll/{approved}{suffix}").status_code == 404
    for pid, action in ((draft, "approve"), (approved, "process")):
        r = rival.post(f"/api/v1/payroll/{pid}/{action}")
        assert r.status_code == 404 and r.json()["code"] == "NOT_FOUND"
    assert rival.post(calc_url(worker_emp["id"], 2025, 6)).status_code == 404
    assert rival.post(calc_url(worker_emp["id"], 2025, 6, action="generate")).status_code == 404
    assert rival.get("/api/v1/payroll").json()["total"] == 0
    assert rival.get("/api/v1/payroll/summary/2025/4").json()["total_employees"] == 0
    assert payroll_row(draft)["status"] == "draft" and payroll_row(approved)["status"] == "approved"


# ---------------------------------------------------------------------------
# L14: summary, report and exports
# ---------------------------------------------------------------------------
EXCEL_HEADERS = ["Employee ID", "Employee Name", "Salary Month", "Currency", "Basic Salary", "HRA", "DA",
                 "Conveyance", "Gross Salary", "Total Deductions", "Net Salary", "Status"]


def _excel(client, year, month):
    from openpyxl import load_workbook

    r = client.get(f"/api/v1/reports/payroll/export/excel?year={year}&month={month}")
    assert r.status_code == 200, r.text
    book = load_workbook(BytesIO(r.content))
    details, summary = book.worksheets[0], book.worksheets[1]
    headers = [details.cell(row=1, column=c).value for c in range(1, 13)]
    labels = {row[0].value: row[1].value for row in summary.iter_rows(min_row=2)}
    return book, headers, labels


def test_L14_summary_report_and_exports_single_and_mixed_currency_defects_D3_D9(co):
    """Float totals per currency, never summed across currencies; mixed -> None plus by_currency.
    DEFECT D3: floats everywhere, Excel float sums. DEFECT D9: no OT/bonus/other columns in Excel.
    Frozen quirk: Excel summary labels are title-cased, so "(INR)" prints as "(Inr)"."""
    empty = co.admin.get("/api/v1/payroll/summary/2025/3").json()
    assert empty == {"month": "2025-03", "total_employees": 0, "currency": "INR", "mixed_currency": False,
                     "total_gross_salary": 0, "total_deductions": 0, "total_net_salary": 0, "average_salary": 0,
                     "by_currency": {}, "status_breakdown": {}}

    first = create_employee(co.admin, "Alpha")
    second = create_employee(co.admin, "Beta")
    structure(co.admin, first["id"], basic_salary="30000", house_rent_allowance="5000")
    structure(co.admin, second["id"], basic_salary="10000")
    mark(co.admin, second["id"], "2025-03-12", "half-day")
    p1 = _generate(co.preparer, first["id"], 2025, 3)["id"]
    _generate(co.preparer, second["id"], 2025, 3)
    _approve(co.approver, p1)

    summary = co.admin.get("/api/v1/payroll/summary/2025/3").json()
    assert summary == {
        "month": "2025-03", "total_employees": 2, "currency": "INR", "mixed_currency": False,
        "total_gross_salary": 45000.0, "total_deductions": 238.1, "total_net_salary": 44761.9,
        "average_salary": 22380.95,
        "by_currency": {"INR": {"employees": 2, "gross": 45000.0, "deductions": 238.1, "net": 44761.9}},
        "status_breakdown": {"approved": 1, "draft": 1},
    }
    report = co.admin.get("/api/v1/reports/payroll?year=2025&month=3").json()
    assert report["currency"] == "INR" and report["currencies"] == ["INR"] and report["total_records"] == 2
    assert report["summary"] == {"total_gross": 45000.0, "total_deductions": 238.1, "total_net": 44761.9,
                                 "average_salary": 22380.95}
    assert set(report["records"][0]) == {"employee_id", "employee_code", "employee_name", "salary_month",
                                         "basic_salary", "hra", "da", "conveyance", "gross_salary",
                                         "total_deductions", "net_salary", "currency", "status"}
    _, headers, labels = _excel(co.admin, 2025, 3)
    assert headers == EXCEL_HEADERS
    assert labels == {"Total Records": 2, "Total Gross (Inr)": 45000.0, "Total Deductions (Inr)": 238.1,
                      "Total Net (Inr)": 44761.9, "Average Salary (Inr)": 22380.95}
    pdf = co.admin.get("/api/v1/reports/payroll/export/pdf?year=2025&month=3")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    text = pdf_text(pdf.content)
    assert "Payroll Report" in text and "₹45,000.00" in text and "₹44,761.90" in text

    entity, dubai = _gulf(co)
    gulf = create_employee(co.admin, "Gamma", branch_id=dubai["id"])
    structure(co.admin, gulf["id"], basic_salary="12000")
    _generate(co.preparer, gulf["id"], 2025, 3)
    mixed = co.admin.get("/api/v1/payroll/summary/2025/3").json()
    assert mixed == {
        "month": "2025-03", "total_employees": 3, "currency": None, "mixed_currency": True,
        "total_gross_salary": None, "total_deductions": None, "total_net_salary": None, "average_salary": None,
        "by_currency": {"AED": {"employees": 1, "gross": 12000.0, "deductions": 0.0, "net": 12000.0},
                        "INR": {"employees": 2, "gross": 45000.0, "deductions": 238.1, "net": 44761.9}},
        "status_breakdown": {"approved": 1, "draft": 2},
    }
    report = co.admin.get("/api/v1/reports/payroll?year=2025&month=3").json()
    assert report["currency"] is None and report["currencies"] == ["AED", "INR"]
    assert report["summary"] == {"total_gross_AED": 12000.0, "total_net_AED": 12000.0,
                                 "total_gross_INR": 45000.0, "total_net_INR": 44761.9}
    _, headers, labels = _excel(co.admin, 2025, 3)
    assert headers == EXCEL_HEADERS
    assert labels == {"Total Records": 3, "Total Gross (Aed)": 12000.0, "Total Deductions (Aed)": 0.0,
                      "Total Net (Aed)": 12000.0, "Average Salary (Aed)": 12000.0,
                      "Total Gross (Inr)": 45000.0, "Total Deductions (Inr)": 238.1,
                      "Total Net (Inr)": 44761.9, "Average Salary (Inr)": 22380.95}
    text = pdf_text(co.admin.get("/api/v1/reports/payroll/export/pdf?year=2025&month=3").content)
    assert "AED 12,000.00" in text and "₹45,000.00" in text
    assert "different currencies are never added together" in text


# ---------------------------------------------------------------------------
# L15-L16: the PAYSLIP_V2 renderer and money formatting (no database needed)
# ---------------------------------------------------------------------------
FIXED_SNAPSHOT = {
    "template_version": "PAYSLIP_V2", "issued_at": "2025-04-01T10:00:00+00:00", "document_number": "PS-7-000042",
    "status": "issued",
    "brand": {"display_name": "Golden Ledger Ltd", "product_name": "Intertech HRMS", "accent": "#3730A3",
              "accent_strip": "#4F46E5", "show_powered_by": True, "powered_by_label": "Powered by Intertech HRMS",
              "legal_footer": "Registered office: Bengaluru"},
    "employer": {"legal_name": "Golden Ledger Private Limited", "country_code": "IN", "code": "LE01",
                 "address": {"line1": "1 MG Road", "city": "Bengaluru"},
                 "registration_lines": [{"label": "CIN", "value": "U72200KA2019PTC000042"}]},
    "employee": {"name": "Priya Sharma", "employee_code": "E0042", "department": "Engineering",
                 "designation": "Engineer", "branch": "Main"},
    "period": "March 2025", "currency": "INR",
    "earnings": [["Basic salary", 30000.0], ["House rent allowance", 5000.0], ["Overtime", 937.5], ["Bonus", 1000.5]],
    "deductions": [["Attendance deductions", 4466.66]],
    "gross": 35000.0, "total_deductions": 4466.66, "net": 32471.34,
    "attendance": {"working_days": 21, "present_days": 3, "absent_days": 2, "half_days": 1},
}


def test_L15_payslip_v2_renderer_is_deterministic_for_a_fixed_snapshot():
    """PAYSLIP_V2 renders the same bytes for the same snapshot (invariant=1), which is what makes a
    re-render from the stored snapshot reproduce the original (EC-WL-09)."""
    from app.utils.pdf_generator import PAYSLIP_TEMPLATE_VERSION, SalarySlipGenerator

    assert PAYSLIP_TEMPLATE_VERSION == "PAYSLIP_V2"
    first = SalarySlipGenerator.generate_pdf(dict(FIXED_SNAPSHOT)).getvalue()
    second = SalarySlipGenerator.generate_pdf(dict(FIXED_SNAPSHOT)).getvalue()
    assert hashlib.sha256(first).hexdigest() == hashlib.sha256(second).hexdigest()
    text = pdf_text(first)
    for label in ("Payslip · March 2025", "Earnings", "Deductions", "Gross earnings", "Total deductions", "Net pay",
                  "Document PS-7-000042 · issued 2025-04-01", "₹35,000.00", "₹32,471.34", "₹1,000.50",
                  "Working days", "Half days", "Golden Ledger Private Limited", "CIN: U72200KA2019PTC000042",
                  "Registered office: Bengaluru", "Powered by Intertech HRMS", "Page 1"):
        assert label in text, label
    assert "DRAFT" not in text
    draft = SalarySlipGenerator.generate_pdf({**FIXED_SNAPSHOT, "status": "draft"}).getvalue()
    assert draft != first and "(DRAFT - not issued)" in pdf_text(draft)


@pytest.mark.parametrize("amount, currency, expected", [
    (Decimal("1234567.89"), "INR", "₹12,34,567.89"),
    (Decimal("-1234"), "INR", "₹-1,234.00"),
    (12000, "AED", "AED 12,000.00"),
    (6200, "USD", "$6,200.00"),
    (1234, "JPY", "¥1,234.00"),  # DEFECT D16 / EC-MON-01: zero-decimal currency printed with 2 dp
    (Decimal("1234.5"), "JPY", "¥1,234.50"),
    (Decimal("12.345"), "KWD", "KWD 12.35"),  # DEFECT D16 / EC-MON-02: three-decimal currency cut to 2 dp
    (Decimal("1.005"), "BHD", "BHD 1.00"),  # ... through float, so 1.005 rounds down
    (None, None, "0.00"),
    (None, "INR", "₹0.00"),
])
def test_L16_pdf_money_formatting_is_frozen(amount, currency, expected):
    """pdf_fonts.money as PAYSLIP_V2 prints it: Indian grouping for INR, symbol or ISO code, always 2 dp."""
    from app.utils.pdf_fonts import money

    assert money(amount, currency) == expected


def _glyph_map():
    from reportlab.pdfbase import pdfmetrics

    from app.utils import pdf_fonts

    pdf_fonts.register()
    return pdfmetrics.getFont(pdf_fonts.SANS).face.charToGlyph


def test_L16_rupee_glyph_is_embedded_EC_DOC_23():
    """EC-DOC-23: ₹ and the other mapped symbols except ฿ and ৳ exist in the embedded NotoSans."""
    from app.utils.pdf_fonts import CURRENCY_SYMBOLS

    glyphs = _glyph_map()
    assert 0x20B9 in glyphs
    missing = {code for code, symbol in CURRENCY_SYMBOLS.items() if any(ord(ch) not in glyphs for ch in symbol)}
    assert missing == {"THB", "BDT"}


@pytest.mark.xfail(strict=True, reason="DEFECT D16 / EC-DOC-23: NotoSans has no ฿ (U+0E3F) or ৳ (U+09F3) glyph")
def test_L16_thai_baht_and_taka_symbols_render_known_gap():
    glyphs = _glyph_map()
    assert 0x0E3F in glyphs and 0x09F3 in glyphs
