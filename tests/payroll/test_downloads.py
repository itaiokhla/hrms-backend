"""Signed downloads, the access trail and the download center (3.6 §7.3, §8).

EC-DOC-06 (signed links: five minutes, one user, one session; expired or reused by another user is denied),
EC-DOC-17 (CSV/XLSX cells escaped), EC-DOC-20 (bulk ZIP built in chunks and streamed from the download center),
EC-DOC-21 (exports enforce permission, scope and masking by role).
"""

import csv
import hashlib
import io
import zipfile
from datetime import date, timedelta

from tests.helpers import create_employee, login, new_client
from tests.payroll.test_compensation import login_as
from tests.payroll.test_payslips_v2 import (  # noqa: F401  (docs_world is a fixture)
    ME,
    PAY,
    docs_world,
    emma_run,
    finalized_run,
    generate,
    payslip_of,
    pdf_text,
)
from tests.payroll.test_wage_inputs import audit_actions, place, sql


def link(client, url):
    r = client.get(url)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["expires_in_seconds"] == 300 and body["url"].startswith(f"{PAY}/downloads/")
    return body["url"]


# ---------------------------------------------------------------------------
# Signed links (EC-DOC-06)
# ---------------------------------------------------------------------------
def test_ec_doc_06_signed_link_is_bound_to_user_session_and_expiry(app, docs_world, monkeypatch):
    from app.payroll.services import downloads

    w = docs_world
    run_id, rows = emma_run(w)
    generate(w["cid"], run_id)
    doc_id, _, number, _, sha, *_ = payslip_of(rows[w["emp"]["Emma"]["id"]])
    emma = w["logins"]["Emma"]
    url = link(emma, f"{ME}/payslips/{doc_id}/download")
    r = emma.get(url)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert hashlib.sha256(r.content).hexdigest() == sha and r.headers["x-document-sha256"] == sha
    assert r.headers["cache-control"] == "private, no-store" and number in pdf_text(r.content)
    assert emma.get(url).status_code == 200  # the same user and session may retry within five minutes
    # Another user, even with full access to the document, cannot use Emma's link.
    assert w["acme"].get(url).status_code == 404
    assert w["logins"]["Ivan"].get(url).status_code == 404
    assert w["beta"]["client"].get(url).status_code == 404
    # The same user in another session cannot use it either.
    other_session = login(app, emma.email)
    assert other_session.get(url).status_code == 404
    # Tampered tokens are 404, never 500.
    token = url.rsplit("/", 1)[1]
    assert emma.get(f"{PAY}/downloads/{token[:-2]}AA").status_code == 404
    assert emma.get(f"{PAY}/downloads/not-a-token").status_code == 404
    # Expired after five minutes.
    real = downloads.clock.utcnow
    monkeypatch.setattr(downloads.clock, "utcnow", lambda: real() + timedelta(minutes=5, seconds=2))
    r = emma.get(url)
    assert r.status_code == 410 and r.json()["code"] == "DOWNLOAD_LINK_EXPIRED"
    monkeypatch.setattr(downloads.clock, "utcnow", real)
    # Unauthenticated: never.
    assert new_client(app).get(url).status_code == 401


def test_access_is_rechecked_when_the_link_is_used(docs_world):
    w = docs_world
    run_id, rows = emma_run(w)
    generate(w["cid"], run_id)
    doc_id = payslip_of(rows[w["emp"]["Emma"]["id"]])[0]
    url = link(w["acme"], f"{PAY}/payslips/{doc_id}/download")
    # The HR user loses the payroll role after asking for the link: the link no longer works.
    sql("DELETE FROM user_role_assignments WHERE user_id = (SELECT id FROM users WHERE email = :m) AND role_id IN "
        "(SELECT id FROM roles WHERE system_key = 'payroll_manager')", m=w["acme"].email)
    sql("UPDATE users SET permission_version = permission_version + 1 WHERE email = :m", m=w["acme"].email)
    assert w["acme"].get(url).status_code in (403, 404)


def test_views_and_downloads_are_logged_and_the_first_own_view_is_recorded(docs_world):
    w = docs_world
    run_id, rows = emma_run(w)
    generate(w["cid"], run_id)
    doc_id = payslip_of(rows[w["emp"]["Emma"]["id"]])[0]
    emma = w["logins"]["Emma"]
    assert w["acme"].get(f"{PAY}/payslips/{doc_id}").status_code == 200  # HR view: not the employee's first view
    assert sql("SELECT first_viewed_at FROM documents WHERE id = :d", d=doc_id)[0][0] is None
    body = emma.get(f"{ME}/payslips/{doc_id}").json()
    assert body["status"] == "VIEWED" and body["content"]["totals"]["net"] == "68700.00"
    assert "publish_reason" not in body and "template_version" not in body  # staff-only fields
    emma.get(link(emma, f"{ME}/payslips/{doc_id}/download"))
    trail = sql("SELECT action, user_id, ip IS NOT NULL FROM document_access_log WHERE document_id = :d ORDER BY id",
                d=doc_id)
    actions = [a for a, *_ in trail]
    # payslip.email_notify is OFF in docs_world: no EMAIL entry (test_publication.py covers the e-mail trail).
    assert actions.count("VIEW") == 2 and actions.count("DOWNLOAD") == 1 and "EMAIL" not in actions
    assert all(has_ip for action, _, has_ip in trail if action in ("VIEW", "DOWNLOAD"))
    logged = audit_actions(w["cid"], "document", doc_id)
    assert "DOCUMENT_VIEWED" in logged and "DOCUMENT_DOWNLOADED" in logged


# ---------------------------------------------------------------------------
# Download center (EC-DOC-20, EC-DOC-17, EC-DOC-21)
# ---------------------------------------------------------------------------
def test_ec_doc_20_bulk_zip_is_chunked_watermarked_and_owned_by_the_requester(app, docs_world, monkeypatch):
    from app.payroll.services import downloads

    w = docs_world
    extra = [create_employee(w["acme"], f"Bulk{i}", hire_date="2024-01-15") for i in range(3)]
    for e in extra:
        place(e["id"], w["le01"], w["group"])
    run_id, rows = finalized_run(w["cid"], w["le01"], w["group"], date(2026, 9, 1),
                                 [w["emp"]["Emma"]["id"], w["emp"]["Ivan"]["id"], *[e["id"] for e in extra]])
    generate(w["cid"], run_id)
    progress = []
    real_progress = downloads.build_payslip_zip

    def spy(ctx_progress, *args, **kwargs):
        return real_progress(lambda done, total=None: (progress.append((done, total)), ctx_progress(done, total)),
                             *args, **kwargs)

    monkeypatch.setattr(downloads, "CHUNK", 2)
    monkeypatch.setattr(downloads, "build_payslip_zip", spy)
    r = w["acme"].post(f"{PAY}/payslips/zip", json={"run_id": run_id})
    assert r.status_code == 202, r.text
    export_id = r.json()["id"]
    export = w["acme"].get(f"{PAY}/exports/{export_id}").json()
    assert export["status"] == "READY" and export["item_count"] == 5, export
    assert progress == [(2, 5), (4, 5), (5, 5)]  # three chunks of at most two documents
    assert export["watermark"].startswith("Confidential — exported by ")
    url = w["acme"].post(f"{PAY}/exports/{export_id}/download").json()["url"]
    r = w["acme"].get(url)
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    archive = zipfile.ZipFile(io.BytesIO(r.content))
    pdfs = [n for n in archive.namelist() if n.endswith(".pdf")]
    assert len(pdfs) == 5 and "README.txt" in archive.namelist()
    assert export["watermark"] in archive.read("README.txt").decode("utf-8")
    assert all("Confidential — exported by" in pdf_text(archive.read(n)) for n in pdfs)
    # The stored payslips are untouched by the watermark.
    sha = payslip_of(rows[w["emp"]["Emma"]["id"]])[4]
    assert all(hashlib.sha256(archive.read(n)).hexdigest() != sha for n in pdfs)
    # Only the requester sees or downloads the export.
    processor = login_as(app, w["acme"], create_employee(w["acme"], "Pia", hire_date="2024-01-15"),
                         ["payroll_processor"])
    assert processor.get(f"{PAY}/exports/{export_id}").status_code == 404
    assert processor.post(f"{PAY}/exports/{export_id}/download").status_code == 404
    assert processor.get(url).status_code == 404
    assert processor.get(f"{PAY}/exports").json()["total"] == 0
    assert w["acme"].get(f"{PAY}/exports").json()["total"] == 1
    trail = sql("SELECT action FROM document_access_log WHERE export_job_id = :e ORDER BY id", e=export_id)
    assert [a for (a,) in trail] == ["EXPORT", "DOWNLOAD"]


def test_ec_doc_17_register_csv_and_xlsx_escape_formula_cells(docs_world):
    from openpyxl import load_workbook

    w = docs_world
    sql("UPDATE employees SET first_name = '=HYPERLINK(\"http://evil\",\"x\")' WHERE id = :e", e=w["emp"]["Emma"]["id"])
    run_id, _ = emma_run(w)
    generate(w["cid"], run_id)
    for kind in ("REGISTER_CSV", "REGISTER_XLSX"):
        r = w["acme"].post(f"{PAY}/exports", json={"kind": kind, "run_id": run_id})
        assert r.status_code == 202, r.text
        export_id = r.json()["id"]
        assert w["acme"].get(f"{PAY}/exports/{export_id}").json()["status"] == "READY"
        data = w["acme"].get(w["acme"].post(f"{PAY}/exports/{export_id}/download").json()["url"]).content
        if kind == "REGISTER_CSV":
            cells = [c for row in csv.reader(io.StringIO(data.decode("utf-8-sig"))) for c in row]
        else:
            book = load_workbook(io.BytesIO(data), read_only=True)
            cells = [str(c) for row in book.active.iter_rows(values_only=True) for c in row if c is not None]
        assert any(c.startswith("'=HYPERLINK") for c in cells)
        assert not any(c.startswith("=") for c in cells)
        assert cells[0].startswith("Confidential — exported by ")
        assert "68700.00" in cells  # exact decimal strings, never floats


def test_ec_doc_21_exports_by_role_and_scope(docs_world):
    w = docs_world
    run_id, _ = emma_run(w)
    generate(w["cid"], run_id)
    emma = w["logins"]["Emma"]
    assert emma.post(f"{PAY}/exports", json={"kind": "REGISTER_CSV", "run_id": run_id}).status_code == 403
    assert emma.post(f"{PAY}/payslips/zip", json={"run_id": run_id}).status_code == 403
    assert w["clerk"].post(f"{PAY}/exports", json={"kind": "REGISTER_CSV", "run_id": run_id}).status_code == 404
    assert w["beta"]["client"].post(f"{PAY}/exports", json={"kind": "REGISTER_CSV", "run_id": run_id}
                                    ).status_code in (403, 404)
    bad = w["acme"].post(f"{PAY}/exports", json={"kind": "REGISTER_CSV", "run_id": run_id, "company_id": 2})
    assert bad.status_code == 422
    assert w["acme"].post(f"{PAY}/exports", json={"kind": "EVERYTHING", "run_id": run_id}).status_code == 422


def test_exports_expire(docs_world):
    w = docs_world
    run_id, _ = emma_run(w)
    generate(w["cid"], run_id)
    export_id = w["acme"].post(f"{PAY}/exports", json={"kind": "REGISTER_CSV", "run_id": run_id}).json()["id"]
    url = w["acme"].post(f"{PAY}/exports/{export_id}/download").json()["url"]
    sql("UPDATE export_jobs SET expires_at = now() - interval '1 minute' WHERE id = :e", e=export_id)
    r = w["acme"].get(url)
    assert r.status_code == 410 and r.json()["code"] == "EXPORT_EXPIRED"
    assert w["acme"].get(f"{PAY}/exports").json()["items"][0]["status"] == "EXPIRED"
    assert w["acme"].post(f"{PAY}/exports/{export_id}/download").status_code == 409
