"""LH-M0 safe fixes (WI-0.6; LH 13 §5): one regression test per bug, failing on the old code.

BUG-01 cancel reason, BUG-02 type flags, BUG-03 idempotency and per-person lock (LH-REQ-02, LH-SEC-07, LH-SEC-09),
BUG-07 local dates, BUG-26 spreadsheet formulas (LH-RPT-29), BUG-30 notification text (LH-NTF-15 pre-state,
LH-PRV-06). BUG-10 is tests/leave_unit/test_fix_links.py; BUG-24 is a Jest test of the leaves page.
"""

import threading
import uuid
from datetime import date, datetime, timedelta, timezone
from io import BytesIO

import pytest
import pytz

from tests.attendance_helpers import IST, fake_clock, owner_rows  # noqa: F401  (fake_clock is a fixture)
from tests.helpers import create_employee, invite_and_activate, login, signup_company

NY = pytz.timezone("America/New_York")


@pytest.fixture
def team(app):
    admin = signup_company(app, "Leave Fix Co")  # company zone Asia/Kolkata
    boss = create_employee(admin, "Arjun", "Mehta")
    worker = create_employee(admin, "Priya", "Shah", manager_id=boss["id"])
    return {"admin": admin, "boss": invite_and_activate(app, admin, boss, ["manager"]),
            "worker": invite_and_activate(app, admin, worker, ["employee"]), "worker_id": worker["id"]}


def _apply(client, start, end=None, leave_type="casual", key=None, **extra):
    headers = {"Idempotency-Key": key} if key else None
    return client.post("/api/v1/leaves/apply", headers=headers,
                       json={"leave_type": leave_type, "start_date": start, "end_date": end or start, **extra})


_SOON = (datetime.now(timezone.utc) + timedelta(days=40)).date()
MONDAY = _SOON + timedelta(days=(7 - _SOON.weekday()) % 7)  # a Monday at least 40 days ahead


def _day(offset: int) -> str:
    """The offset-th working day (Mon-Fri) from MONDAY: the default work week has no leave-free dates there."""
    weeks, day = divmod(offset, 5)
    return (MONDAY + timedelta(days=7 * weeks + day)).isoformat()


# ----------------------------------------------------------------------------------------------- BUG-01
def test_bug01_cancel_reason_saved(team):
    leave = _apply(team["worker"], _day(0)).json()
    r = team["worker"].post(f"/api/v1/leaves/{leave['id']}/cancel", json={"reason": "Plans changed"})
    assert r.status_code == 200, r.text
    assert owner_rows("SELECT status, cancellation_reason FROM leaves WHERE id = :i", i=leave["id"]) == [
        {"status": "cancelled", "cancellation_reason": "Plans changed"}]
    audit = owner_rows("SELECT description, details::text AS d FROM audit_logs WHERE action = 'CANCEL'")
    assert all("Plans changed" not in (a["description"] + (a["d"] or "")) for a in audit)


def test_bug01_cancel_without_body_still_works(team):
    leave = _apply(team["worker"], _day(1)).json()
    assert team["worker"].post(f"/api/v1/leaves/{leave['id']}/cancel").status_code == 200


# ----------------------------------------------------------------------------------------------- BUG-02
def test_bug02_type_flags_round_trip(team):
    r = team["admin"].post("/api/v1/leaves/types", json={"code": "comp_off", "name": "Comp off", "kind": "COMP_OFF",
                                                         "allow_half_day": False, "counts_non_working_days": True})
    assert r.status_code == 200, r.text
    row = next(t for t in team["admin"].get("/api/v1/leaves/types").json() if t["code"] == "comp_off")
    assert (row["kind"], row["allow_half_day"], row["counts_non_working_days"]) == ("COMP_OFF", False, True)


def test_bug02_legacy_put_keeps_flags(team):
    created = team["admin"].post("/api/v1/leaves/types", json={"code": "comp_off", "name": "Comp off",
                                                               "kind": "COMP_OFF", "allow_half_day": False}).json()
    # Today's policy tab sends only these fields.
    r = team["admin"].put(f"/api/v1/leaves/types/{created['id']}", json={
        "code": "comp_off", "name": "Compensatory off", "annual_quota": 0, "is_paid": True, "is_unlimited": False,
        "is_active": True})
    assert r.status_code == 200, r.text
    assert (r.json()["kind"], r.json()["allow_half_day"], r.json()["name"]) == ("COMP_OFF", False, "Compensatory off")


# ----------------------------------------------------------------------------------------------- BUG-03
def test_lh_req_02_same_key_replays(team):
    key = uuid.uuid4().hex
    first = _apply(team["worker"], _day(2), key=key)
    second = _apply(team["worker"], _day(2), key=key)
    assert first.status_code == second.status_code == 200, second.text
    assert second.headers.get("Idempotent-Replayed") == "true" and second.json()["id"] == first.json()["id"]
    assert len(owner_rows("SELECT id FROM leaves WHERE employee_id = :e", e=team["worker_id"])) == 1


def test_lh_sec_07_key_reused_other_body_409(team):
    key = uuid.uuid4().hex
    assert _apply(team["worker"], _day(3), key=key).status_code == 200
    r = _apply(team["worker"], _day(4), key=key)
    assert r.status_code == 409 and r.json()["code"] == "IDEMPOTENCY_KEY_REUSED"
    bad = _apply(team["worker"], _day(4), key="not a valid key!")
    assert bad.status_code == 400 and bad.json()["code"] == "IDEMPOTENCY_KEY_INVALID"


def test_lh_rel_10_key_after_the_24_hour_window_runs_again(team, fake_clock):
    """The server keeps a key 24 h (idempotency.TTL); a retry after that is a new request, never a replay."""
    fake_clock.now = IST.localize(datetime(2026, 11, 2, 9, 0)).astimezone(timezone.utc)
    key = uuid.uuid4().hex
    first = _apply(team["worker"], "2026-11-23", key=key)
    assert first.status_code == 200, first.text
    replay = _apply(team["worker"], "2026-11-23", key=key)
    assert replay.headers.get("Idempotent-Replayed") == "true"
    fake_clock.now += timedelta(hours=25)
    later = _apply(team["worker"], "2026-11-24", key=key)
    assert later.status_code == 200 and later.headers.get("Idempotent-Replayed") is None
    assert later.json()["id"] != first.json()["id"]


def test_lh_sec_01_foreign_tenant_leave_id_is_404_on_every_route(app, team):
    leave = _apply(team["worker"], _day(7)).json()
    other = signup_company(app, "Foreign Co")
    foreign = invite_and_activate(app, other, create_employee(other, "Omar", "Haddad"), ["hr_admin"])
    missing = 999999
    for action in ("approve", "reject", "cancel"):
        r = foreign.post(f"/api/v1/leaves/{leave['id']}/{action}", json={})
        ghost = foreign.post(f"/api/v1/leaves/{missing}/{action}", json={})
        assert r.status_code == ghost.status_code == 404, (action, r.text)
        assert r.json() == ghost.json()  # same body as an id that does not exist
    assert foreign.get(f"/api/v1/leaves/balance/{team['worker_id']}").status_code == 404
    listed = foreign.get("/api/v1/leaves", params={"employee_id": team["worker_id"]}).json()
    assert listed["total"] == 0
    assert owner_rows("SELECT status FROM leaves WHERE id = :i", i=leave["id"]) == [{"status": "pending"}]


def test_lh_sec_18_every_tenant_table_enforces_rls(app):
    """The contract test of tests/test_platform_contracts.py is the gate (release gate 1)."""
    from tests.test_platform_contracts import test_every_tenant_table_enforces_row_level_security

    test_every_tenant_table_enforces_row_level_security(app)


def test_lh_sec_19_rls_status_reports_enforcement_and_admin_group_membership():
    from sqlalchemy import text

    from app.database.db import engine, rls_status

    assert rls_status() == {"database": "postgresql", "rls_enforced": True}
    with engine.connect() as conn:  # the CASE guard: no error on a server without neon_superuser
        assert conn.execute(text("SELECT count(*) FROM pg_roles WHERE rolname = 'neon_superuser'")).scalar() == 0


def test_lh_sec_20_reason_text_in_a_job_payload_is_refused():
    from app.payroll.services import jobs

    for key in ("reason", "note", "message", "comment", "diagnosis", "file_name", "cancellationReason"):
        with pytest.raises(ValueError):
            jobs.check_payload({key: "anything"})
    jobs.check_payload({"reason_code": "SICK", "employee_ids": [1, 2], "local_date": "2026-11-02",
                        "message_id": 7, "note_status": "OPEN"})


def test_lh_sec_09_parallel_applies_against_three_available_days(app, team):
    """20 parallel 2-day applies on different dates against 3.00 available days: exactly one is reserved, the
    others are refused for balance, and available never drops below zero (LCK-02 lock before the checks)."""
    from app.core.rate_limit import limiter

    r = team["admin"].post("/api/v1/leaves/types", json={"code": "study", "name": "Study leave", "annual_quota": 3})
    assert r.status_code == 200, r.text
    email = team["worker"].email
    clients = []
    for _ in range(20):
        limiter.reset()  # twenty sign-ins from one test address
        clients.append(login(app, email))
    # Twenty disjoint two-working-day windows inside one leave year (Feb-Mar of next year).
    feb = date(MONDAY.year + 1, 2, 1)
    base = feb + timedelta(days=(7 - feb.weekday()) % 7)
    year = base.year
    results = []

    def window(n):
        weeks, day = divmod(2 * n, 5)
        first = base + timedelta(days=7 * weeks + day)
        second = first + timedelta(days=3 if first.weekday() == 4 else 1)
        return first.isoformat(), second.isoformat()

    def go(n, client):
        start, end = window(n)
        results.append(_apply(client, start, end, leave_type="study", key=uuid.uuid4().hex).json().get("code", "OK"))

    threads = [threading.Thread(target=go, args=(n, c)) for n, c in enumerate(clients)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert results.count("OK") == 1, results
    assert set(results) <= {"OK", "INSUFFICIENT_BALANCE", "LEAVE_OVERLAP"}, results
    balance = team["worker"].get("/api/v1/leaves/my/balance", params={"year": year}).json()["study"]
    assert balance["pending"] == 2 and balance["available"] == 1


# ----------------------------------------------------------------------------------------------- BUG-07
def test_bug07_apply_past_date_uses_applicant_zone(app, monkeypatch, fake_clock):
    """A New York employee of an India-zoned company may still apply for 'today' in New York (LH-TZ-06)."""
    from app.services import company_service

    admin = signup_company(app, "Zones Co")
    entity = admin.post("/api/v1/legal-entities", json={"code": "US1", "legal_name": "Zones Inc", "country_code": "US",
                                                         "base_currency": "USD", "timezone": "America/New_York"}).json()
    branch = admin.post("/api/v1/branches", json={"legal_entity_id": entity["id"], "code": "NYC", "name": "NYC",
                                                  "timezone": "America/New_York"}).json()
    emp = create_employee(admin, "Carlos", "Diaz", legal_entity_id=entity["id"], branch_id=branch["id"],
                          hire_date="2026-01-05")
    worker = invite_and_activate(app, admin, emp, ["employee"])
    # Mon 2026-11-02 22:00 in New York = Tue 2026-11-03 08:30 in India.
    fake_clock.now = NY.localize(datetime(2026, 11, 2, 22, 0)).astimezone(timezone.utc)
    monkeypatch.setattr(company_service, "utcnow", lambda: fake_clock.now)  # the old check read the company zone
    r = _apply(worker, "2026-11-02")
    assert r.status_code == 200, r.text


def test_bug07_dashboard_balance_year_is_the_employees(app, team, fake_clock):
    fake_clock.now = IST.localize(datetime(2026, 11, 2, 9, 0)).astimezone(timezone.utc)
    leave = _apply(team["worker"], "2026-11-16").json()
    assert team["boss"].post(f"/api/v1/leaves/{leave['id']}/approve").status_code == 200
    # 31 Dec 2026 20:00 UTC is already 1 Jan 2027 in India: the new year's balance is shown.
    fake_clock.now = datetime(2026, 12, 31, 20, 0, tzinfo=timezone.utc)
    balance = team["worker"].get("/api/v1/dashboard").json()["me"]["leave_balance"]["casual"]
    assert balance["used"] == 0


def test_bug07_delegation_uses_delegator_zone(app, team, fake_clock):
    fake_clock.now = IST.localize(datetime(2026, 6, 15, 10, 0)).astimezone(timezone.utc)  # far from the real date
    deputy = create_employee(team["admin"], "Deepa", "Nair")
    deputy_client = invite_and_activate(app, team["admin"], deputy, ["manager"])
    r = team["boss"].post("/api/v1/org/delegations", json={"delegate_employee_id": deputy["id"],
                                                           "starts_on": "2026-06-15", "ends_on": "2026-06-15"})
    assert r.status_code == 200, r.text
    from app.database.db import SessionLocal, bind_tenant
    from app.services import org_service

    db = SessionLocal()
    try:
        bind_tenant(db, team["admin"].company_id)
        boss_id = team["boss"].me()["employee"]["id"]
        assert org_service.active_delegators_for(db, deputy["id"]) == [boss_id]
    finally:
        db.close()
    assert deputy_client is not None


# ----------------------------------------------------------------------------------------------- BUG-26
def test_lh_rpt_29_formula_cells_escaped():
    from openpyxl import load_workbook

    from app.utils.excel_exporter import ExcelExporter

    rows = [{"employee_id": 1, "employee_name": "=HYPERLINK(\"http://x\")", "leave_type": "casual",
             "start_date": "2026-11-02", "end_date": "2026-11-02", "days": 1, "reason": "+1 call me @home",
             "status": "approved"},
            {"employee_id": 2, "employee_name": "\tTab", "leave_type": "-sick", "start_date": "2026-11-03",
             "end_date": "2026-11-03", "days": 1, "reason": "\rCR", "status": "-"}]
    wb = load_workbook(BytesIO(ExcelExporter.export_leave_report(rows).read()))
    values = [c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert values, "the export has text cells"
    assert not any(len(v) > 1 and v.startswith(("=", "+", "-", "@", "\t", "\r")) for v in values), values
    assert "-" in values  # the lone "no value" placeholder is kept as it is
    assert all(c.data_type != "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row)


# ----------------------------------------------------------------------------------------------- BUG-30
def test_bug30_notification_has_no_type_name(team):
    r = team["admin"].post("/api/v1/leaves/types", json={"code": "fertility_care", "name": "Fertility care",
                                                          "annual_quota": 5})
    assert r.status_code == 200, r.text
    leave = _apply(team["worker"], _day(6), leave_type="fertility_care").json()
    assert team["boss"].post(f"/api/v1/leaves/{leave['id']}/approve").status_code == 200
    rows = owner_rows("SELECT title, message FROM notifications")
    assert rows, "the request and the decision notify someone"
    assert all("fertility" not in (r["title"] + r["message"]).lower() for r in rows), rows
    assert any(r["title"] == "Leave request approved" for r in rows)  # asserted by e2e/05-leave.spec.ts
