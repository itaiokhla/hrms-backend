"""Punch In / Lunch In / Lunch Out / Punch Out state machine, hours and their consumers."""

import threading
from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from sqlalchemy import text

from tests.attendance_helpers import (  # noqa: F401  (fake_clock is a fixture)
    IST, act, employee_with_login, fake_clock, key, meta, owner_rows, set_policy, today,
)
from tests.helpers import create_employee, invite_and_activate, signup_company

DAY = "2026-10-19"


@pytest.fixture
def team(app):
    admin = signup_company(app, "Lifecycle Co")
    set_policy(admin, requireFaceForSelfCheckIn=False)
    emp = employee_with_login(app, admin, "Asha")
    return {"admin": admin, "emp": emp}


def _types(view):
    return [e["event_type"] for e in view["events"]]


def test_full_day_matches_goal_example(team, fake_clock):
    emp = team["emp"]
    fake_clock.at_ist(DAY, "09:00")
    r = act(emp, "PUNCH_IN", meta(fake_clock))
    assert r.status_code == 200, r.text
    view = r.json()
    assert view["attendance_state"] == "WORKING" and view["state"] == "working"
    assert view["actions"] == ["LUNCH_IN", "PUNCH_OUT"]

    fake_clock.at_ist(DAY, "13:00")
    view = act(emp, "LUNCH_IN", meta(fake_clock)).json()
    assert view["attendance_state"] == "ON_LUNCH" and view["actions"] == ["LUNCH_OUT", "PUNCH_OUT"]

    fake_clock.at_ist(DAY, "13:30")
    during = today(emp)["record"]["hours"]
    assert during["provisional"] and during["lunch_seconds"] == 30 * 60 and during["net_seconds"] == 4 * 3600

    fake_clock.at_ist(DAY, "13:45")
    assert act(emp, "LUNCH_OUT", meta(fake_clock)).json()["attendance_state"] == "WORKING"
    fake_clock.at_ist(DAY, "18:00")
    view = act(emp, "PUNCH_OUT", meta(fake_clock)).json()
    hours = view["record"]["hours"]
    assert view["attendance_state"] == "PUNCHED_OUT" and view["actions"] == []
    assert (hours["gross_seconds"], hours["lunch_seconds"], hours["net_seconds"]) == (9 * 3600, 45 * 60, 8 * 3600 + 15 * 60)
    assert not hours["provisional"] and hours["complete"]
    assert _types(view) == ["PUNCH_IN", "LUNCH_IN", "LUNCH_OUT", "PUNCH_OUT"]
    assert view["record"]["working_seconds"] == 29700
    # The database holds exactly those server times.
    rows = owner_rows("SELECT event_type, occurred_at FROM attendance_events ORDER BY id")
    assert [r["occurred_at"].astimezone(IST).strftime("%H:%M") for r in rows] == ["09:00", "13:00", "13:45", "18:00"]


def test_two_lunches_and_single_lunch_policy(app, team, fake_clock):
    emp, admin = team["emp"], team["admin"]
    for hhmm, action in (("09:00", "PUNCH_IN"), ("12:30", "LUNCH_IN"), ("13:00", "LUNCH_OUT"),
                         ("15:30", "LUNCH_IN"), ("15:45", "LUNCH_OUT"), ("18:00", "PUNCH_OUT")):
        fake_clock.at_ist(DAY, hhmm)
        assert act(emp, action, meta(fake_clock)).status_code == 200
    hours = today(emp)["record"]["hours"]
    assert hours["net_seconds"] == 8 * 3600 + 15 * 60 and hours["lunch_seconds"] == 45 * 60

    set_policy(admin, allowMultipleLunches=False)
    second = employee_with_login(app, admin, "Bina")
    fake_clock.at_ist("2026-10-20", "09:00")
    act(second, "PUNCH_IN", meta(fake_clock))
    fake_clock.at_ist("2026-10-20", "12:00")
    act(second, "LUNCH_IN", meta(fake_clock))
    fake_clock.at_ist("2026-10-20", "12:30")
    view = act(second, "LUNCH_OUT", meta(fake_clock)).json()
    assert view["actions"] == ["PUNCH_OUT"]  # no second lunch offered...
    fake_clock.at_ist("2026-10-20", "15:00")
    r = act(second, "LUNCH_IN", meta(fake_clock))
    assert r.status_code == 409 and r.json()["code"] == "LUNCH_LIMIT"  # ...or accepted


def test_invalid_transitions_are_refused(team, fake_clock):
    emp = team["emp"]
    fake_clock.at_ist(DAY, "08:00")
    assert act(emp, "LUNCH_IN", meta(fake_clock)).json()["code"] == "NOT_CHECKED_IN"
    assert act(emp, "LUNCH_OUT", meta(fake_clock)).json()["code"] == "NOT_CHECKED_IN"
    r = act(emp, "PUNCH_OUT", meta(fake_clock))
    assert r.status_code == 400 and r.json()["code"] == "NOT_CHECKED_IN"

    fake_clock.at_ist(DAY, "09:00")
    assert act(emp, "PUNCH_IN", meta(fake_clock)).status_code == 200
    r = act(emp, "PUNCH_IN", meta(fake_clock))
    assert r.status_code == 409 and r.json()["code"] == "ALREADY_CHECKED_IN"
    assert act(emp, "LUNCH_OUT", meta(fake_clock)).json()["code"] == "NOT_ON_LUNCH"
    fake_clock.at_ist(DAY, "09:00", 30)
    assert act(emp, "PUNCH_OUT", meta(fake_clock)).json()["code"] == "TOO_SOON"
    fake_clock.at_ist(DAY, "12:00")
    assert act(emp, "LUNCH_IN", meta(fake_clock)).status_code == 200
    assert act(emp, "LUNCH_IN", meta(fake_clock)).json()["code"] == "ALREADY_ON_LUNCH"

    # Punch-out during lunch needs an explicit confirmation, then closes the break at the punch-out.
    fake_clock.at_ist(DAY, "12:40")
    r = act(emp, "PUNCH_OUT", meta(fake_clock))
    assert r.status_code == 409 and r.json()["code"] == "LUNCH_OPEN_CONFIRM"
    view = act(emp, "PUNCH_OUT", meta(fake_clock, confirm_close_lunch=True)).json()
    assert view["attendance_state"] == "PUNCHED_OUT"
    brk = view["record"]["breaks"]
    assert len(brk) == 1 and brk[0]["end_reason"] == "AUTO_CLOSED_AT_PUNCH_OUT" and brk[0]["ended_at"] == view["record"]["check_out_time"]
    assert view["record"]["hours"]["net_seconds"] == 3 * 3600
    assert owner_rows("SELECT count(*) AS n FROM attendance_breaks WHERE ended_at IS NULL")[0]["n"] == 0
    # A finished day cannot be punched again or extended with lunch.
    assert act(emp, "PUNCH_IN", meta(fake_clock)).json()["code"] == "ALREADY_CHECKED_IN"
    assert act(emp, "LUNCH_IN", meta(fake_clock)).json()["code"] == "NOT_CHECKED_IN"


def test_block_policy_requires_lunch_out_first(team, fake_clock):
    emp, admin = team["emp"], team["admin"]
    set_policy(admin, lunchOpenAtPunchOut="BLOCK")
    fake_clock.at_ist(DAY, "09:00")
    act(emp, "PUNCH_IN", meta(fake_clock))
    fake_clock.at_ist(DAY, "13:00")
    view = act(emp, "LUNCH_IN", meta(fake_clock)).json()
    assert view["actions"] == ["LUNCH_OUT"]
    r = act(emp, "PUNCH_OUT", meta(fake_clock, confirm_close_lunch=True))
    assert r.status_code == 409 and r.json()["code"] == "LUNCH_OPEN"


def test_idempotent_retries_and_conflicting_keys(team, fake_clock):
    emp = team["emp"]
    fake_clock.at_ist(DAY, "09:00")
    body = meta(fake_clock)
    first = act(emp, "PUNCH_IN", body)
    retry = act(emp, "PUNCH_IN", body)  # network retry / double click with the same intent
    assert first.status_code == retry.status_code == 200 and retry.json()["replayed"] is True
    assert owner_rows("SELECT count(*) AS n FROM attendance_events")[0]["n"] == 1
    fake_clock.at_ist(DAY, "12:00")
    r = act(emp, "LUNCH_IN", {**meta(fake_clock), "idempotency_key": body["idempotency_key"]})
    assert r.status_code == 409 and r.json()["code"] == "IDEMPOTENCY_CONFLICT"


def test_concurrent_requests_change_state_once(app, team, fake_clock):
    emp = team["emp"]
    fake_clock.at_ist(DAY, "09:00")
    act(emp, "PUNCH_IN", meta(fake_clock))
    fake_clock.at_ist(DAY, "13:00")
    # Five tabs press "Start lunch" at the same moment with different request ids.
    from tests.helpers import login

    tabs = [login(app, emp.email) for _ in range(5)]
    results = []
    barrier = threading.Barrier(len(tabs))

    def press(tab):
        barrier.wait()
        results.append(act(tab, "LUNCH_IN", meta(fake_clock)).status_code)

    threads = [threading.Thread(target=press, args=(t,)) for t in tabs]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert sorted(results) == [200, 409, 409, 409, 409]
    assert owner_rows("SELECT count(*) AS n FROM attendance_breaks")[0]["n"] == 1


def test_stale_tab_and_inactive_employee(team, fake_clock):
    emp, admin = team["emp"], team["admin"]
    fake_clock.at_ist(DAY, "09:00")
    session_id = act(emp, "PUNCH_IN", meta(fake_clock)).json()["record"]["id"]
    fake_clock.at_ist(DAY, "12:00")
    r = act(emp, "LUNCH_IN", meta(fake_clock, attendance_id=session_id + 999))
    assert r.status_code == 409 and r.json()["code"] == "STALE_SESSION"
    assert act(emp, "LUNCH_IN", meta(fake_clock, attendance_id=session_id)).status_code == 200


def test_overnight_shift_punches_out_next_day(team, fake_clock):
    emp = team["emp"]
    fake_clock.at_ist("2026-10-31", "22:00")
    record = act(emp, "PUNCH_IN", meta(fake_clock)).json()["record"]
    assert record["date"] == "2026-10-31"
    fake_clock.at_ist("2026-11-01", "02:00")
    view = today(emp)
    assert view["attendance_state"] == "WORKING" and view["overnight"] is True
    act(emp, "LUNCH_IN", meta(fake_clock))
    fake_clock.at_ist("2026-11-01", "02:30")
    act(emp, "LUNCH_OUT", meta(fake_clock))
    fake_clock.at_ist("2026-11-01", "06:30")
    r = act(emp, "PUNCH_OUT", meta(fake_clock))
    assert r.status_code == 200, r.text
    body = r.json()
    hours = body["hours"]  # the shift acted on (it belongs to 31 Oct)
    assert hours["gross_seconds"] == 8 * 3600 + 30 * 60 and hours["net_seconds"] == 8 * 3600
    assert body["attendance_state"] == "NOT_PUNCHED_IN" and body["record"] is None
    assert body["last_session"]["id"] == record["id"] and body["last_session"]["hours"] == hours
    # The new calendar day starts fresh.
    fake_clock.at_ist("2026-11-01", "21:55")
    assert act(emp, "PUNCH_IN", meta(fake_clock)).status_code == 200


def test_dst_shift_counts_real_hours(app, fake_clock):
    admin = signup_company(app, "Eastern Co")
    admin.put("/api/v1/settings", json={"timezone": "America/New_York"})
    set_policy(admin, requireFaceForSelfCheckIn=False)
    emp = employee_with_login(app, admin, "Dana")
    import pytz

    ny = pytz.timezone("America/New_York")
    fake_clock.now = ny.localize(datetime(2026, 3, 7, 22, 0)).astimezone(timezone.utc)
    act(emp, "PUNCH_IN", meta(fake_clock))
    fake_clock.now = ny.localize(datetime(2026, 3, 8, 6, 0)).astimezone(timezone.utc)
    hours = act(emp, "PUNCH_OUT", meta(fake_clock)).json()["hours"]
    assert hours["gross_seconds"] == 7 * 3600  # clocks sprang forward overnight


def test_shift_left_open_is_closed_without_inventing_a_punch_out(team, fake_clock):
    emp = team["emp"]
    fake_clock.at_ist(DAY, "09:00")
    act(emp, "PUNCH_IN", meta(fake_clock))
    fake_clock.at_ist("2026-10-20", "01:30")  # 16.5 h later (company limit 16 h)
    view = today(emp)
    assert view["attendance_state"] == "NOT_PUNCHED_IN"  # yesterday's shift was closed; a new day can start
    rows = owner_rows("SELECT check_out_time, closed_reason FROM attendance")
    assert rows == [{"check_out_time": None, "closed_reason": "MISSING_PUNCH_OUT"}]
    assert [r["event_type"] for r in owner_rows("SELECT event_type FROM attendance_events ORDER BY id")] == ["PUNCH_IN", "SESSION_CLOSED"]
    history = emp.get("/api/v1/attendance/me/history", params={"start_date": "2026-10-18", "end_date": "2026-10-20"}).json()
    hours = history["items"][0]["hours"]
    assert hours["complete"] is False and hours["net_seconds"] is None and "MISSING_PUNCH_OUT" in hours["issues"]
    fake_clock.at_ist("2026-10-20", "09:00")
    assert act(emp, "PUNCH_IN", meta(fake_clock)).status_code == 200


def test_face_rules_and_inactive_employment(app, fake_clock):
    admin = signup_company(app, "Face Co")
    emp_record = create_employee(admin, "Fara")
    emp = invite_and_activate(app, admin, emp_record, ["employee"])
    fake_clock.at_ist(DAY, "09:00")
    r = act(emp, "PUNCH_IN", meta(fake_clock))
    assert r.status_code == 400 and r.json()["code"] == "FACE_REQUIRED"  # default: face required at punches
    set_policy(admin, requireFaceForSelfCheckIn=False)
    assert act(emp, "PUNCH_IN", meta(fake_clock)).status_code == 200
    fake_clock.at_ist(DAY, "12:00")
    assert act(emp, "LUNCH_IN", meta(fake_clock)).status_code == 200  # lunch needs no face by default
    set_policy(admin, requireFaceForLunch=True)
    fake_clock.at_ist(DAY, "12:30")
    r = act(emp, "LUNCH_OUT", meta(fake_clock))
    assert r.status_code == 400 and r.json()["code"] == "FACE_REQUIRED"
    assert admin.post(f"/api/v1/employees/{emp_record['id']}/deactivate").status_code == 200
    r = act(emp, "LUNCH_OUT", meta(fake_clock))
    assert r.status_code in (401, 403)  # the account is deactivated with the employment


def test_manual_correction_keeps_history(team, fake_clock):
    emp, admin = team["emp"], team["admin"]
    fake_clock.at_ist(DAY, "09:00")
    act(emp, "PUNCH_IN", meta(fake_clock))
    fake_clock.at_ist(DAY, "13:00")
    act(emp, "LUNCH_IN", meta(fake_clock))
    fake_clock.at_ist(DAY, "13:20")
    act(emp, "LUNCH_OUT", meta(fake_clock))
    fake_clock.at_ist(DAY, "17:00")
    act(emp, "PUNCH_OUT", meta(fake_clock))

    def ist(hhmm):
        return IST.localize(datetime.strptime(f"{DAY} {hhmm}", "%Y-%m-%d %H:%M")).isoformat()

    employee_id = emp.employee["id"]
    r = admin.post("/api/v1/attendance/manual", json={
        "employee_id": employee_id, "date": DAY, "check_in_time": ist("09:00"), "check_out_time": ist("18:00"),
        "status": "present", "notes": "Badge reader missed the evening punch",
        "breaks": [{"started_at": ist("13:00"), "ended_at": ist("13:45")}],
    })
    assert r.status_code == 200, r.text
    hours = r.json()["hours"]
    assert (hours["gross_seconds"], hours["lunch_seconds"], hours["net_seconds"]) == (32400, 2700, 29700)
    events = owner_rows("SELECT event_type, details FROM attendance_events ORDER BY id")
    assert [e["event_type"] for e in events] == ["PUNCH_IN", "LUNCH_IN", "LUNCH_OUT", "PUNCH_OUT", "CORRECTION"]
    # The replaced lunch end (13:20 IST = 07:50 UTC) is kept in the correction event.
    assert '"before"' in events[-1]["details"] and "07:50:00" in events[-1]["details"]

    # Corrections cannot erase the punch-out or push breaks outside the session.
    bad = admin.post("/api/v1/attendance/manual", json={
        "employee_id": employee_id, "date": DAY, "check_in_time": ist("09:00")})
    assert bad.status_code == 409 and bad.json()["code"] == "CANNOT_REOPEN"
    bad = admin.post("/api/v1/attendance/manual", json={
        "employee_id": employee_id, "date": DAY, "check_in_time": ist("14:00"), "check_out_time": ist("18:00")})
    assert bad.status_code == 409 and bad.json()["code"] == "BREAKS_OUTSIDE_SESSION"
    bad = admin.post("/api/v1/attendance/manual", json={
        "employee_id": employee_id, "date": DAY, "check_in_time": ist("09:00"), "check_out_time": ist("18:00"),
        "breaks": [{"started_at": ist("13:00"), "ended_at": ist("13:45")}, {"started_at": ist("13:30"), "ended_at": ist("14:00")}]})
    assert bad.status_code == 409 and bad.json()["code"] == "BREAKS_OVERLAP"
    # Employees cannot correct their own attendance.
    assert emp.post("/api/v1/attendance/manual", json={
        "employee_id": employee_id, "date": DAY, "check_in_time": ist("09:00")}).status_code == 403


def test_database_enforces_append_only_events_and_break_rules(team, fake_clock):
    emp = team["emp"]
    fake_clock.at_ist(DAY, "09:00")
    att_id = act(emp, "PUNCH_IN", meta(fake_clock)).json()["record"]["id"]
    from app.database.db import engine, system_engine

    with pytest.raises(Exception, match="append-only"):
        with system_engine.begin() as conn:  # even the owner cannot rewrite history
            conn.execute(text("UPDATE attendance_events SET occurred_at = now()"))
    with pytest.raises(Exception, match="permission denied"):
        with engine.begin() as conn:  # the application role has no UPDATE/DELETE grant at all
            conn.execute(text("SELECT set_config('app.company_id', :c, true)"), {"c": str(emp.me()["company"]["id"])})
            conn.execute(text("DELETE FROM attendance_events"))
    row = owner_rows("SELECT company_id, employee_id, check_in_time FROM attendance WHERE id = :i", i=att_id)[0]
    with pytest.raises(Exception, match="overlap"):
        with system_engine.begin() as conn:
            for start, end in ((0, 60), (30, 90)):
                conn.execute(text(
                    "INSERT INTO attendance_breaks (company_id, attendance_id, employee_id, kind, started_at, ended_at, "
                    "is_paid, created_at, updated_at) VALUES (:c, :a, :e, 'LUNCH', :s, :t, false, now(), now())"),
                    {"c": row["company_id"], "a": att_id, "e": row["employee_id"],
                     "s": row["check_in_time"] + timedelta(minutes=start), "t": row["check_in_time"] + timedelta(minutes=end)})
    with pytest.raises(Exception, match="before the punch-in"):
        with system_engine.begin() as conn:
            conn.execute(text(
                "INSERT INTO attendance_breaks (company_id, attendance_id, employee_id, kind, started_at, is_paid, "
                "created_at, updated_at) VALUES (:c, :a, :e, 'LUNCH', :s, false, now(), now())"),
                {"c": row["company_id"], "a": att_id, "e": row["employee_id"], "s": row["check_in_time"] - timedelta(hours=1)})


def test_reports_payroll_and_exports_reconcile(app, team, fake_clock):
    from openpyxl import load_workbook

    emp, admin = team["emp"], team["admin"]
    plan = {"2026-09-07": [("09:00", "PUNCH_IN"), ("13:00", "LUNCH_IN"), ("13:45", "LUNCH_OUT"), ("18:00", "PUNCH_OUT")],
            "2026-09-08": [("09:00", "PUNCH_IN"), ("12:30", "LUNCH_IN"), ("13:00", "LUNCH_OUT"), ("15:30", "LUNCH_IN"),
                           ("15:45", "LUNCH_OUT"), ("18:00", "PUNCH_OUT")],
            "2026-09-09": [("10:00", "PUNCH_IN"), ("16:00", "PUNCH_OUT")]}
    for day, steps in plan.items():
        for hhmm, action in steps:
            fake_clock.at_ist(day, hhmm)
            assert act(emp, action, meta(fake_clock)).status_code == 200
    expected_net = 29700 + 29700 + 6 * 3600
    params = {"start_date": "2026-09-01", "end_date": "2026-09-30"}
    report = admin.get("/api/v1/reports/attendance", params=params).json()
    assert report["summary"]["total_net_seconds"] == expected_net
    assert sum(r["net_seconds"] for r in report["records"]) == expected_net
    summary = admin.get("/api/v1/attendance/hours", params=params).json()
    assert summary["totals"]["net_seconds"] == expected_net and summary["totals"]["lunch_seconds"] == 5400
    listing = admin.get("/api/v1/attendance", params=params).json()
    assert sum(i["hours"]["net_seconds"] for i in listing["items"]) == expected_net

    admin.post("/api/v1/salary-structures", json={"employee_id": emp.employee["id"], "basic_salary": "30000",
                                                   "effective_from": "2026-01-01"})
    payroll_user = employee_with_login(app, admin, "Pay", roles=("payroll_manager",))
    calc = payroll_user.post(f"/api/v1/payroll/calculate?employee_id={emp.employee['id']}&year=2026&month=9").json()
    assert calc["attendance_hours"]["net_seconds"] == expected_net
    run = payroll_user.post(f"/api/v1/payroll/generate?employee_id={emp.employee['id']}&year=2026&month=9").json()
    slip = admin.get(f"/api/v1/payroll/{run['id']}/slip").json()
    assert slip["attendance"]["net_work_seconds"] == expected_net
    stored = owner_rows("SELECT net_work_seconds, gross_work_seconds FROM payroll WHERE id = :i", i=run["id"])[0]
    assert stored == {"net_work_seconds": expected_net, "gross_work_seconds": 32400 + 32400 + 21600}

    book = load_workbook(BytesIO(admin.get("/api/v1/reports/attendance/export/excel", params=params).content))
    details, summary_sheet = book.worksheets[0], book.worksheets[1]
    assert [details.cell(row=1, column=c).value for c in (8, 9, 10)] == ["Gross Hours", "Lunch Hours", "Net Hours"]
    assert sorted(details.cell(row=r, column=10).value for r in (2, 3, 4)) == [6.0, 8.25, 8.25]
    totals = {row[0].value: row[1].value for row in summary_sheet.iter_rows(min_row=2)}
    assert totals["Total Net Seconds"] == expected_net and totals["Total Net Hours"] == 22.5
    assert admin.get("/api/v1/reports/attendance/export/pdf", params=params).status_code == 200
