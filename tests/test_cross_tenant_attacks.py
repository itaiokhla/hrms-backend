"""Cross-tenant attack simulations (three companies), complementing test_tenant_isolation.

Each test plants data in one company and attacks it from another through the API:
forged tenant parameters, foreign ids in filters, nested references, exports, files,
report aggregates, concurrent interleaved requests and stale privileges.
"""

import io
from concurrent.futures import ThreadPoolExecutor

from tests.helpers import create_employee, invite_and_activate
from tests.white_label_helpers import png


def _xlsx_text(content: bytes) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content), read_only=True)
    return " ".join(str(cell) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True) for cell in row
                    if cell is not None)


def _ids(body):
    return {item["id"] for item in body["items"]}


def test_forged_tenant_query_parameters_are_ignored(three_companies):
    a, b = three_companies["A"], three_companies["B"]
    zed = create_employee(b, "Zed")
    own = _ids(a.get("/api/v1/employees").json())
    forged = a.get("/api/v1/employees", params={"company_id": b.company_id, "tenant_id": b.company_id}).json()
    assert _ids(forged) == own and zed["id"] not in own
    assert a.get("/api/v1/payroll", params={"company_id": b.company_id}).json()["total"] == 0
    assert a.get("/api/v1/auth/me", params={"company_id": b.company_id}).json()["company"]["id"] == a.company_id


def test_filters_and_pagination_with_foreign_ids_return_nothing(three_companies):
    a, b = three_companies["A"], three_companies["B"]
    b_dept = b.post("/api/v1/departments", json={"name": "Fleet"}).json()["id"]
    b_mgr = create_employee(b, "Manny")
    create_employee(b, "Reportee", manager_id=b_mgr["id"], department_id=b_dept)
    assert a.get("/api/v1/employees", params={"department_id": b_dept}).json()["total"] == 0
    assert a.get("/api/v1/employees", params={"manager_id": b_mgr["id"]}).json()["total"] == 0
    own_total = a.get("/api/v1/employees").json()["total"]
    page = a.get("/api/v1/employees", params={"skip": 0, "limit": 200, "include_inactive": True}).json()
    assert page["total"] == own_total and len(page["items"]) == own_total
    beyond = a.get("/api/v1/employees", params={"skip": 10_000, "limit": 200}).json()
    assert beyond["items"] == [] and beyond["total"] == own_total


def test_nested_references_to_another_tenant_are_rejected(three_companies):
    a, b = three_companies["A"], three_companies["B"]
    b_dept = b.post("/api/v1/departments", json={"name": "Fleet"}).json()["id"]
    b_emp = create_employee(b, "Bea")
    a_emp = create_employee(a, "Ann")
    before = a.get("/api/v1/employees").json()["total"]
    payload = {"first_name": "Nest", "last_name": "Ed", "email": "nested@acme.io", "employee_code": "N-1", "hire_date": "2024-01-02"}
    assert a.post("/api/v1/employees", json={**payload, "department_id": b_dept}).status_code in (400, 404)
    assert a.post("/api/v1/employees", json={**payload, "manager_id": b_emp["id"]}).status_code in (400, 404)
    assert a.put(f"/api/v1/employees/{a_emp['id']}", json={"department_id": b_dept}).status_code in (400, 404)
    assert a.get("/api/v1/employees").json()["total"] == before
    asset = a.post("/api/v1/assets", json={"asset_tag": "LAP-9", "name": "Laptop", "category": "IT"}).json()
    assert a.post(f"/api/v1/assets/{asset['id']}/assign", json={"employee_id": b_emp["id"]}).status_code == 404
    line = {"employee_id": a_emp["id"], "manager_id": b_emp["id"], "relation_type": "functional", "effective_from": "2024-01-01"}
    assert a.post("/api/v1/org/secondary-lines", json=line).status_code == 404
    b_holiday = b.post("/api/v1/holidays", json={"name": "National Day", "date": "2026-12-02"}).json()
    assert a.put(f"/api/v1/holidays/{b_holiday['id']}", json={"name": "Hijacked", "date": "2026-12-02"}).status_code == 404
    assert a.delete(f"/api/v1/holidays/{b_holiday['id']}").status_code == 404
    holidays = b.get("/api/v1/holidays", params={"year": 2026}).json()
    items = holidays if isinstance(holidays, list) else holidays.get("items", holidays.get("holidays", []))
    assert [h["name"] for h in items] == ["National Day"]


def test_leave_records_and_reports_never_cross_tenants(three_companies):
    a, b = three_companies["A"], three_companies["B"]
    b_me = b.me()["employee"]["id"]
    leave = b.post("/api/v1/leaves/apply", json={"leave_type": "casual", "start_date": "2099-05-04", "end_date": "2099-05-04"})
    assert leave.status_code == 200, leave.text
    leave_id = leave.json()["id"]
    assert a.post(f"/api/v1/leaves/{leave_id}/approve").status_code == 404
    assert a.post(f"/api/v1/leaves/{leave_id}/reject", json={"reason": "x"}).status_code in (404, 422)
    assert a.get(f"/api/v1/leaves/balance/{b_me}").status_code == 404
    report = a.get("/api/v1/reports/leaves", params={"year": 2099, "employee_id": b_me})
    assert report.status_code in (200, 404)
    if report.status_code == 200:
        assert str(b.email) not in report.text and "2099-05-04" not in report.text
    export = a.get("/api/v1/reports/leaves/export/excel", params={"year": 2099, "employee_id": b_me})
    assert export.status_code in (200, 404)
    if export.status_code == 200:
        text = _xlsx_text(export.content)
        assert "2099-05-04" not in text and b.email not in text


def test_payroll_aggregates_and_exports_count_only_the_own_tenant(three_companies):
    a, b = three_companies["A"], three_companies["B"]
    payee = create_employee(b, "Payee")
    assert b.post("/api/v1/salary-structures", json={"employee_id": payee["id"], "basic_salary": "91234",
                                                     "effective_from": "2024-01-01"}).status_code == 200
    assert b.post(f"/api/v1/payroll/generate?employee_id={payee['id']}&year=2024&month=2").status_code == 200
    summary = a.get("/api/v1/payroll/summary/2024/2")
    assert summary.status_code == 200 and "91234" not in summary.text and "91,234" not in summary.text
    report = a.get("/api/v1/reports/payroll", params={"year": 2024, "month": 2})
    assert report.status_code == 200 and "Payee" not in report.text and "91234" not in report.text
    export = a.get("/api/v1/reports/payroll/export/excel", params={"year": 2024, "month": 2})
    assert export.status_code == 200 and "Payee" not in _xlsx_text(export.content)
    assert "Payee" in _xlsx_text(b.get("/api/v1/reports/payroll/export/excel", params={"year": 2024, "month": 2}).content)
    workbook = a.get("/api/v1/employees/export")
    assert workbook.status_code == 200 and "Payee" not in _xlsx_text(workbook.content)
    dashboard = a.get("/api/v1/dashboard").json()
    assert "Payee" not in str(dashboard)
    analytics = a.get("/api/v1/reports/department-analytics", params={"year": 2024, "month": 2})
    assert analytics.status_code == 200 and "91234" not in analytics.text


def test_another_tenants_files_are_not_served(three_companies):
    a, b = three_companies["A"], three_companies["B"]
    b_emp = create_employee(b, "Pic")
    r = b.post(f"/api/v1/employees/{b_emp['id']}/photo", files={"file": ("p.png", io.BytesIO(png(200, 200)), "image/png")})
    assert r.status_code == 200, r.text
    assert b.get(f"/api/v1/employees/{b_emp['id']}/photo").status_code == 200
    assert a.get(f"/api/v1/employees/{b_emp['id']}/photo").status_code == 404
    # A's own logo endpoint never falls back to another tenant's file.
    assert a.get("/api/v1/settings/logo").status_code == 404


def test_interleaved_concurrent_requests_never_mix_tenants(three_companies):
    a, b, c = three_companies["A"], three_companies["B"], three_companies["C"]
    for client, name in ((a, "Alpha"), (b, "Beta"), (c, "Gamma")):
        for i in range(3):
            create_employee(client, f"{name}{i}")
    expected = {id(client): _ids(client.get("/api/v1/employees").json()) for client in (a, b, c)}

    def probe(client):
        return id(client), _ids(client.get("/api/v1/employees").json()), client.get("/api/v1/auth/me").json()["company"]["id"]

    with ThreadPoolExecutor(max_workers=6) as pool:
        results = list(pool.map(probe, [a, b, c] * 8))
    companies = {id(a): a.company_id, id(b): b.company_id, id(c): c.company_id}
    for key, ids, company_id in results:
        assert ids == expected[key] and company_id == companies[key]
    # The same employee code in two tenants at the same moment: both succeed, each in its own tenant.
    payload = {"first_name": "Same", "last_name": "Code", "employee_code": "DUP-1", "hire_date": "2024-01-02"}
    with ThreadPoolExecutor(max_workers=2) as pool:
        made = list(pool.map(lambda client: client.post("/api/v1/employees", json={**payload, "email": f"dup@{client.company_id}.io"}),
                             [a, b]))
    assert [r.status_code for r in made] == [200, 200]
    assert a.get("/api/v1/employees", params={"q": "Same"}).json()["total"] == 1


def test_revoked_privilege_takes_effect_on_the_live_session(three_companies, app):
    a = three_companies["A"]
    hr_emp = create_employee(a, "Hira")
    hr = invite_and_activate(app, a, hr_emp, ["hr_admin"])
    assert hr.get("/api/v1/employees").status_code == 200
    employee_role = [r["id"] for r in a.get("/api/v1/roles").json() if r["system_key"] == "employee"]
    assert a.put(f"/api/v1/users/{hr.user_id}/roles", json={"role_ids": employee_role}).status_code == 200
    # Same cookies, no new sign-in: the downgrade applies to the next request.
    listing = hr.get("/api/v1/employees")
    assert listing.status_code == 200 and listing.json()["total"] == 1  # only their own record now
    assert hr.get("/api/v1/users").status_code == 403
