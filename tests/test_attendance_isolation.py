"""Three companies, reporting hierarchies and adversarial access to attendance and location data.

Company Alpha                       Company Beta            Company Gamma
├── Admin A, HR A                   ├── Admin B, HR B       ├── Admin C, HR C
├── Manager A1                      └── Manager B1          └── Manager C1
│   ├── Employee A1                     ├── Employee B1         ├── Employee C1
│   ├── Employee A2                     └── Employee B2         └── Employee C2
│   └── Team Lead A1 (nested)
│       └── Employee A4
└── Manager A2
    └── Employee A3

Alpha requires a location at punches and tracks every 120 min; Beta makes it
optional and tracks every 60 min; Gamma collects no location at all.
"""

from io import BytesIO

import pytest
from sqlalchemy import text

from app.services import geo
from tests.attendance_helpers import (  # noqa: F401  (fake_clock is a fixture)
    BLR, act, employee_with_login, fake_clock, fix, meta, owner_rows, sample, set_policy, tracking,
)
from tests.helpers import create_employee, invite_and_activate, signup_company

DAY = "2026-10-19"
TEAM_LEAD = [("employee:read", "DIRECT_REPORTS"), ("attendance:read", "DIRECT_REPORTS"),
             ("location:read", "DIRECT_REPORTS"), ("location:review", "DIRECT_REPORTS")]


def _user(app, admin, employee, roles):
    from app.core.rate_limit import limiter

    limiter.reset()  # one test client IP creates every account here; the limiter is tested elsewhere
    return invite_and_activate(app, admin, employee, roles)


@pytest.fixture
def world(app, fake_clock):
    fake_clock.at_ist(DAY, "09:00")
    people, sessions = {}, {}

    # Alpha
    admin_a = signup_company(app, "Alpha Logistics")
    set_policy(admin_a, requireFaceForSelfCheckIn=False, locationPunchPolicy="REQUIRED", locationLunchPolicy="OPTIONAL",
               trackingEnabled=True, trackingIntervalMinutes=120)
    tl_role = admin_a.post("/api/v1/roles", json={"name": "Team Lead", "grants": [
        {"permission": p, "scope": s} for p, s in TEAM_LEAD]}).json()["id"]
    m1 = create_employee(admin_a, "ManagerA1")
    m2 = create_employee(admin_a, "ManagerA2")
    e1 = create_employee(admin_a, "EmployeeA1", manager_id=m1["id"])
    e2 = create_employee(admin_a, "EmployeeA2", manager_id=m1["id"])
    tl = create_employee(admin_a, "LeadA1", manager_id=m1["id"])
    e4 = create_employee(admin_a, "EmployeeA4", manager_id=tl["id"])
    e3 = create_employee(admin_a, "EmployeeA3", manager_id=m2["id"])
    people.update({
        "admin_a": admin_a,
        "hr_a": employee_with_login(app, admin_a, "HRA", roles=("hr_admin",)),
        "m1": _user(app, admin_a, m1, ["manager"]), "m2": _user(app, admin_a, m2, ["manager"]),
        "e1": _user(app, admin_a, e1, ["employee"]), "e2": _user(app, admin_a, e2, ["employee"]),
        "e3": _user(app, admin_a, e3, ["employee"]), "e4": _user(app, admin_a, e4, ["employee"]),
    })
    lead = _user(app, admin_a, tl, ["employee"])
    from tests.helpers import role_id

    admin_a.put(f"/api/v1/users/{lead.user_id}/roles", json={"role_ids": [role_id(admin_a, "employee"), tl_role]})
    people["tl"] = lead

    # Beta
    admin_b = signup_company(app, "Beta Retail")
    set_policy(admin_b, requireFaceForSelfCheckIn=False, locationPunchPolicy="OPTIONAL", trackingEnabled=True,
               trackingIntervalMinutes=60)
    mb = create_employee(admin_b, "ManagerB1")
    people.update({
        "admin_b": admin_b, "hr_b": employee_with_login(app, admin_b, "HRB", roles=("hr_admin",)),
        "mb1": _user(app, admin_b, mb, ["manager"]),
        "b1": _user(app, admin_b, create_employee(admin_b, "EmployeeB1", manager_id=mb["id"]), ["employee"]),
        "b2": _user(app, admin_b, create_employee(admin_b, "EmployeeB2", manager_id=mb["id"]), ["employee"]),
    })

    # Gamma (no location collection at all)
    admin_c = signup_company(app, "Gamma Health")
    set_policy(admin_c, requireFaceForSelfCheckIn=False)
    mc = create_employee(admin_c, "ManagerC1")
    people.update({
        "admin_c": admin_c, "mc1": _user(app, admin_c, mc, ["manager"]),
        "c1": _user(app, admin_c, create_employee(admin_c, "EmployeeC1", manager_id=mc["id"]), ["employee"]),
        "c2": _user(app, admin_c, create_employee(admin_c, "EmployeeC2", manager_id=mc["id"]), ["employee"]),
    })

    # Everyone punches in with a (synthetic) reading at a distinct place.
    for index, name in enumerate(["e1", "e2", "e3", "e4", "tl", "b1", "b2", "c1", "c2"]):
        lat, lng = geo.destination(*BLR, 10 * index, 100 + 50 * index)
        r = act(people[name], "PUNCH_IN", meta(fake_clock, fix(fake_clock, lat, lng)))
        assert r.status_code == 200, (name, r.text)
        sessions[name] = r.json()["record"]["id"]
    return {"p": people, "s": sessions}


def _ids(client, path="/api/v1/attendance", **params):
    r = client.get(path, params={"start_date": DAY, "end_date": DAY, **params})
    assert r.status_code == 200, r.text
    return {item["employee_name"].split()[0] for item in r.json()["items"]}


def test_hierarchy_scopes_attendance_and_location_separately(world, fake_clock):
    p, s = world["p"], world["s"]
    assert _ids(p["m1"]) == {"EmployeeA1", "EmployeeA2", "LeadA1", "EmployeeA4"}  # whole reporting tree
    assert _ids(p["m2"]) == {"EmployeeA3"}
    assert _ids(p["tl"]) == {"LeadA1", "EmployeeA4"}  # own record + direct report
    assert _ids(p["e1"]) == {"EmployeeA1"}
    assert len(_ids(p["hr_a"])) == 5

    # Manager A1 never reaches the sibling branch.
    assert p["m1"].get(f"/api/v1/attendance/sessions/{s['e3']}").status_code == 404
    assert p["m1"].get(f"/api/v1/attendance/statistics?employee_id={p['e3'].employee['id']}&month=10&year=2026").status_code == 404

    # Precise location is a separate, narrower grant: direct reports only for a manager.
    direct = p["m1"].get(f"/api/v1/attendance/sessions/{s['e1']}", params={"include_locations": True}).json()
    assert direct["location_access"] is True and direct["samples"][0]["latitude"] is not None
    nested = p["m1"].get(f"/api/v1/attendance/sessions/{s['e4']}", params={"include_locations": True}).json()
    assert nested["employee_name"].startswith("EmployeeA4") and nested["location_access"] is False
    assert nested["samples"] is None and nested["location"]["captured"] == 1  # counts only, no coordinates
    led = p["tl"].get(f"/api/v1/attendance/sessions/{s['e4']}", params={"include_locations": True}).json()
    assert led["location_access"] is True and led["samples"][0]["longitude"] is not None
    assert p["tl"].get(f"/api/v1/attendance/sessions/{s['e1']}").status_code == 404

    # Employees see their own coordinates and nobody else's.
    own = p["e1"].get(f"/api/v1/attendance/sessions/{s['e1']}", params={"include_locations": True}).json()
    assert own["samples"][0]["latitude"] is not None
    for other in ("e2", "e3", "e4"):
        assert p["e1"].get(f"/api/v1/attendance/sessions/{s[other]}", params={"include_locations": True}).status_code == 404
    assert p["e1"].get("/api/v1/attendance/location-review").status_code == 403

    # Coordinate access by others is audited; self-access is not noise in the trail.
    logs = p["admin_a"].get("/api/v1/audit-logs?module=Location").json()["logs"]
    viewers = {(log["actor_email"], log["description"].split(" of ")[1].split(" for ")[0]) for log in logs}
    assert (p["m1"].email, "EmployeeA1 Test") in viewers and (p["tl"].email, "EmployeeA4 Test") in viewers
    assert not any(log["actor_email"] == p["e1"].email for log in logs)

    # Movement flags reach the right reviewer only.
    fake_clock.advance(minutes=125)
    for name in ("e1", "e4"):
        lat, lng = geo.destination(*BLR, 0, 9000)
        assert sample(p[name], fake_clock, s[name], fix(fake_clock, lat, lng)).status_code == 200
    m1_queue = p["m1"].get("/api/v1/attendance/location-review").json()
    tl_queue = p["tl"].get("/api/v1/attendance/location-review").json()
    assert {i["employee_name"] for i in m1_queue} == {"EmployeeA1 Test"}
    assert {i["employee_name"] for i in tl_queue} == {"EmployeeA4 Test"}
    e4_item = tl_queue[0]["id"]
    assert p["m1"].post(f"/api/v1/attendance/location-review/{e4_item}", json={"decision": "ACKNOWLEDGED"}).status_code == 404
    assert p["e4"].post(f"/api/v1/attendance/location-review/{e4_item}", json={"decision": "ACKNOWLEDGED"}).status_code == 403
    ok = p["tl"].post(f"/api/v1/attendance/location-review/{e4_item}", json={"decision": "ACKNOWLEDGED", "note": "Client site visit"})
    assert ok.status_code == 200 and ok.json()["review_status"] == "ACKNOWLEDGED"
    again = p["tl"].post(f"/api/v1/attendance/location-review/{e4_item}", json={"decision": "ACKNOWLEDGED"})
    assert again.status_code == 409

    # Managers get no payroll through the attendance pages.
    for path in ("/api/v1/payroll", "/api/v1/reports/payroll?year=2026"):
        assert p["m1"].get(path).status_code == 403
    summary = p["m1"].get("/api/v1/attendance/summary/today").json()
    # In scope: the manager themself plus the four people in their tree.
    assert summary["present_today"] == 4 and summary["total_employees"] == 5


def test_cross_company_probes_are_denied_without_leaking(world, fake_clock):
    p, s = world["p"], world["s"]
    alpha_sample = owner_rows("SELECT id FROM location_samples WHERE attendance_id = :a", a=s["e1"])[0]["id"]
    alpha_employee = p["e1"].employee["id"]
    for name in ("admin_b", "hr_b", "mb1", "admin_c"):
        client = p[name]
        assert client.get(f"/api/v1/attendance/sessions/{s['e1']}", params={"include_locations": True}).status_code == 404
        assert client.get("/api/v1/attendance", params={"start_date": DAY, "end_date": DAY,
                                                        "employee_id": alpha_employee}).json()["total"] == 0
        hours = client.get("/api/v1/attendance/hours", params={"start_date": DAY, "end_date": DAY,
                                                               "employee_id": alpha_employee}).json()
        assert hours["employees"] == [] and hours["totals"]["sessions"] == 0
        assert client.get(f"/api/v1/attendance/statistics?employee_id={alpha_employee}&month=10&year=2026").status_code == 404
    assert p["admin_b"].post(f"/api/v1/attendance/location-review/{alpha_sample}", json={"decision": "ACKNOWLEDGED"}).status_code == 404
    assert p["admin_b"].get("/api/v1/attendance/location-review").json() == []
    assert p["admin_b"].post("/api/v1/attendance/manual", json={
        "employee_id": alpha_employee, "date": DAY, "check_in_time": fake_clock.now.isoformat()}).status_code == 404

    # Uploading into someone else's shift: the same answer as an unknown id, and nothing is stored.
    before = owner_rows("SELECT count(*) AS n FROM location_samples")[0]["n"]
    fake_clock.advance(minutes=125)
    for attacker, target in (("b1", "e1"), ("e2", "e1"), ("c1", "b1")):
        r = sample(p[attacker], fake_clock, s[target], fix(fake_clock))
        assert r.status_code == 409 and r.json()["code"] == "SESSION_NOT_ACTIVE", (attacker, r.text)
    assert owner_rows("SELECT count(*) AS n FROM location_samples WHERE status = 'CAPTURED'")[0]["n"] == before
    # Over-posted ownership fields are rejected outright.
    forged = p["b1"].post("/api/v1/attendance/me/location-samples", json={
        "idempotency_key": "forged-key-123", "attendance_id": s["b1"], "location": fix(fake_clock),
        "company_id": p["admin_a"].company_id, "employee_id": alpha_employee})
    assert forged.status_code == 422
    forged_event = act(p["b1"], "LUNCH_IN", {**meta(fake_clock), "employee_id": alpha_employee})
    assert forged_event.status_code == 422

    # Reports, exports and dashboards aggregate only the caller's company.
    from openpyxl import load_workbook

    for admin, company_names in (("admin_a", {"EmployeeA1", "EmployeeA2", "EmployeeA3", "EmployeeA4", "LeadA1"}),
                                 ("admin_b", {"EmployeeB1", "EmployeeB2"}), ("admin_c", {"EmployeeC1", "EmployeeC2"})):
        report = p[admin].get("/api/v1/reports/attendance", params={"start_date": DAY, "end_date": DAY}).json()
        assert {r["employee_name"].split()[0] for r in report["records"]} == company_names
        book = load_workbook(BytesIO(p[admin].get("/api/v1/reports/attendance/export/excel",
                                                  params={"start_date": DAY, "end_date": DAY}).content))
        exported = {row[1].split()[0] for row in book.worksheets[0].iter_rows(min_row=2, values_only=True) if row[1]}
        assert exported == company_names
        assert p[admin].get("/api/v1/attendance/summary/today").json()["present_today"] == len(company_names)

    # Row-level security holds even for a raw query under another tenant's context.
    from app.database.db import engine

    with engine.begin() as conn:
        conn.execute(text("SELECT set_config('app.company_id', :c, true)"), {"c": str(p["admin_b"].company_id)})
        visible = conn.execute(text("SELECT DISTINCT company_id FROM location_samples")).scalars().all()
        events = conn.execute(text("SELECT DISTINCT company_id FROM attendance_events")).scalars().all()
    assert visible == [p["admin_b"].company_id] and events == [p["admin_b"].company_id]


def test_company_policies_are_isolated(world, fake_clock):
    p, s = world["p"], world["s"]
    # Beta changes its policy; Alpha's and Gamma's are untouched.
    set_policy(p["admin_b"], trackingIntervalMinutes=90, locationPunchPolicy="REQUIRED")
    assert p["admin_a"].get("/api/v1/attendance/policy").json()["settings"]["trackingIntervalMinutes"] == 120
    assert p["admin_c"].get("/api/v1/attendance/policy").json()["settings"]["trackingEnabled"] is False
    assert p["admin_a"].get("/api/v1/attendance/policy").json()["settings"]["locationPunchPolicy"] == "REQUIRED"
    # Only administrators of a company change its policy (HR and managers cannot, other companies cannot reach it).
    for name in ("hr_a", "m1", "e1"):
        assert p[name].patch("/api/v1/attendance/policy", json={"trackingIntervalMinutes": 240}).status_code == 403

    assert tracking(p["e1"])["interval_minutes"] == 120
    assert tracking(p["b1"])["interval_minutes"] == 90  # re-anchored to Beta's new interval
    gamma = tracking(p["c1"])
    assert gamma["state"] == "NOT_APPLICABLE"
    # Gamma collects nothing even though its browsers sent readings.
    assert owner_rows("SELECT count(*) AS n FROM location_samples WHERE company_id = :c",
                      c=p["admin_c"].company_id)[0]["n"] == 0
    fake_clock.advance(minutes=125)
    assert sample(p["c1"], fake_clock, s["c1"], fix(fake_clock)).json()["code"] == "TRACKING_NOT_ACTIVE"
    # Alpha's two-hour slot is due, Beta's 90-minute one too; each company follows its own schedule.
    assert tracking(p["e1"])["schedule_state"] == "DUE"
    assert tracking(p["b1"])["schedule_state"] in ("DUE", "OVERDUE")
